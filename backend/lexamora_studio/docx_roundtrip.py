import difflib
import hashlib
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from .models import DocxImport


def _document_snapshot(asset):
    with asset.file.open("rb") as stream:
        document = Document(stream)

    paragraphs = []
    headings = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = paragraph.style.name if paragraph.style else ""
        row = {"style": style, "text": text}
        paragraphs.append(row)
        if style.casefold().startswith(("heading", "title")):
            headings.append(row)

    tables = []
    for table in document.tables:
        tables.append([
            ["\n".join(part.strip() for part in cell.text.splitlines() if part.strip()) for cell in row.cells]
            for row in table.rows
        ])

    images = []
    for blip in document.element.body.iter(qn("a:blip")):
        relationship_id = blip.get(qn("r:embed"))
        part = document.part.related_parts.get(relationship_id)
        blob = getattr(part, "blob", b"")
        if not blob:
            continue
        images.append({
            "name": Path(str(getattr(part, "partname", ""))).name,
            "sha256": hashlib.sha256(blob).hexdigest(),
            "bytes": len(blob),
        })

    return {
        "paragraphs": paragraphs,
        "headings": headings,
        "tables": tables,
        "images": images,
    }


def _text_lines(snapshot):
    lines = [row["text"] for row in snapshot["paragraphs"]]
    for table in snapshot["tables"]:
        for row in table:
            lines.append(" | ".join(row))
    return lines


def _sequence_changes(source, generated, limit=200):
    return list(difflib.unified_diff(
        source,
        generated,
        fromfile="Original",
        tofile="Generated",
        lineterm="",
        n=1,
    ))[:limit]


def _ratio(source, generated):
    return round(difflib.SequenceMatcher(None, source, generated, autojunk=False).ratio() * 100, 1)


def compare_docx_export(job):
    source_import = (
        job.project.source_imports.filter(status=DocxImport.Status.ACCEPTED)
        .select_related("source_asset")
        .first()
    )
    if source_import is None:
        return {
            "available": False,
            "error": "This project has no accepted source DOCX. The generated document can still be downloaded.",
            "source_import": None,
        }
    if not job.output_asset_id:
        return {
            "available": False,
            "error": "The generated DOCX file is unavailable.",
            "source_import": source_import,
        }

    try:
        original = _document_snapshot(source_import.source_asset)
        generated = _document_snapshot(job.output_asset)
    except Exception as exc:
        return {
            "available": False,
            "error": f"DOCX comparison failed: {exc}",
            "source_import": source_import,
        }

    original_text = _text_lines(original)
    generated_text = _text_lines(generated)
    original_headings = [row["text"] for row in original["headings"]]
    generated_headings = [row["text"] for row in generated["headings"]]
    original_tables = [repr(table) for table in original["tables"]]
    generated_tables = [repr(table) for table in generated["tables"]]
    original_images = [item["sha256"] for item in original["images"]]
    generated_images = [item["sha256"] for item in generated["images"]]

    categories = {
        "structure": {
            "label": "Structure",
            "same": (
                len(original["paragraphs"]) == len(generated["paragraphs"])
                and len(original["tables"]) == len(generated["tables"])
                and len(original["headings"]) == len(generated["headings"])
                and len(original["images"]) == len(generated["images"])
            ),
            "original": {
                "paragraphs": len(original["paragraphs"]),
                "headings": len(original["headings"]),
                "tables": len(original["tables"]),
                "images": len(original["images"]),
            },
            "generated": {
                "paragraphs": len(generated["paragraphs"]),
                "headings": len(generated["headings"]),
                "tables": len(generated["tables"]),
                "images": len(generated["images"]),
            },
        },
        "texts": {
            "label": "Texts",
            "same": original_text == generated_text,
            "similarity": _ratio(original_text, generated_text),
            "diff": _sequence_changes(original_text, generated_text),
        },
        "tables": {
            "label": "Tables",
            "same": original["tables"] == generated["tables"],
            "similarity": _ratio(original_tables, generated_tables),
            "original_count": len(original["tables"]),
            "generated_count": len(generated["tables"]),
            "diff": _sequence_changes(original_tables, generated_tables, limit=80),
        },
        "images": {
            "label": "Images",
            "same": original_images == generated_images,
            "original_count": len(original_images),
            "generated_count": len(generated_images),
            "matching_hashes": sum(1 for left, right in zip(original_images, generated_images) if left == right),
            "original": original["images"],
            "generated": generated["images"],
        },
        "order": {
            "label": "Section order",
            "same": original_headings == generated_headings,
            "similarity": _ratio(original_headings, generated_headings),
            "original": original_headings,
            "generated": generated_headings,
            "diff": _sequence_changes(original_headings, generated_headings, limit=100),
        },
    }
    changed = sum(1 for category in categories.values() if not category["same"])
    return {
        "available": True,
        "source_import": source_import,
        "categories": categories,
        "changed_categories": changed,
        "exact_match": changed == 0,
    }