import json
from datetime import timedelta
from unittest.mock import patch

from cryptography.fernet import Fernet
from django.contrib.auth import get_user_model
from django.db import connection
from django.test import TestCase, override_settings
from django.test.utils import CaptureQueriesContext
from django.utils import timezone

from .models import Project, ProjectAccessExclusion, ProjectMembership, Prompt, WorkspaceMembership
from .permissions import accessible_projects, accessible_workspaces, has_capability, has_project_capability
from .services import create_workspace


class StudioPermissionsTests(TestCase):
    def setUp(self):
        users = get_user_model()
        self.owner = users.objects.create_user("owner", password="strong-pass")
        self.viewer = users.objects.create_user("viewer", password="strong-pass")
        self.outsider = users.objects.create_user("outsider", password="strong-pass")
        self.workspace = create_workspace(user=self.owner, name="Dom Studio", slug="dom-studio")
        WorkspaceMembership.objects.create(
            workspace=self.workspace,
            user=self.viewer,
            role=WorkspaceMembership.Role.VIEWER,
        )

    def test_user_without_membership_cannot_see_workspace(self):
        self.assertFalse(accessible_workspaces(self.outsider).filter(id=self.workspace.id).exists())

    def test_viewer_cannot_edit(self):
        self.assertTrue(has_capability(self.viewer, self.workspace, "view"))
        self.assertFalse(has_capability(self.viewer, self.workspace, "edit"))

    def test_owner_can_edit_and_manage_members(self):
        self.assertTrue(has_capability(self.owner, self.workspace, "edit"))
        self.assertTrue(has_capability(self.owner, self.workspace, "manage_members"))

    def test_dashboard_requires_login(self):
        response = self.client.get("/studio/")
        self.assertEqual(response.status_code, 302)
        self.assertIn("/login/", response.url)

    def test_cross_workspace_project_is_not_exposed(self):
        project = Project.objects.create(
            workspace=self.workspace,
            project_type=Project.Type.SERIES,
            title="Dom",
            created_by=self.owner,
            updated_by=self.owner,
        )
        self.client.force_login(self.outsider)
        self.assertEqual(self.client.get(f"/studio/projects/{project.id}/").status_code, 404)

    def test_workspace_creation_makes_owner_membership(self):
        created = create_workspace(user=self.outsider, name="Second", slug="second")
        membership = WorkspaceMembership.objects.get(workspace=created, user=self.outsider)
        self.assertEqual(membership.role, WorkspaceMembership.Role.OWNER)
        self.assertTrue(membership.can_use_ai)


class ProjectSharingTests(TestCase):
    def setUp(self):
        users = get_user_model()
        self.owner = users.objects.create_user("share-owner", email="owner@example.com", password="pass")
        self.viewer = users.objects.create_user("share-viewer", email="viewer@example.com", password="pass")
        self.editor = users.objects.create_user("share-editor", email="editor@example.com", password="pass")
        self.controller = users.objects.create_user("share-controller", email="controller@example.com", password="pass")
        self.outsider = users.objects.create_user("share-outsider", email="outsider@example.com", password="pass")
        self.workspace = create_workspace(user=self.owner, name="Private Studio", slug="private-studio")
        self.project = Project.objects.create(
            workspace=self.workspace, project_type=Project.Type.SERIES, title="Shared project",
            created_by=self.owner, updated_by=self.owner,
        )
        self.other_project = Project.objects.create(
            workspace=self.workspace, project_type=Project.Type.SERIES, title="Private project",
            created_by=self.owner, updated_by=self.owner,
        )
        ProjectMembership.objects.create(project=self.project, user=self.viewer, role=ProjectMembership.Role.VIEWER, invited_by=self.owner)
        ProjectMembership.objects.create(project=self.project, user=self.editor, role=ProjectMembership.Role.EDITOR, invited_by=self.owner)
        ProjectMembership.objects.create(project=self.project, user=self.controller, role=ProjectMembership.Role.CONTROLLER, invited_by=self.owner)

    def test_workspace_owner_can_open_project_sharing(self):
        self.assertTrue(has_project_capability(self.owner, self.project, "manage_project"))
        self.client.force_login(self.owner)
        detail = self.client.get(f"/studio/projects/{self.project.id}/")
        self.assertContains(detail, f"/studio/projects/{self.project.id}/access/")
        self.assertEqual(self.client.get(f"/studio/projects/{self.project.id}/access/").status_code, 200)

    def test_project_share_does_not_expose_workspace_or_sibling_project(self):
        self.assertTrue(accessible_projects(self.viewer).filter(id=self.project.id).exists())
        self.assertFalse(accessible_projects(self.viewer).filter(id=self.other_project.id).exists())
        self.assertFalse(accessible_workspaces(self.viewer).filter(id=self.workspace.id).exists())
        self.client.force_login(self.viewer)
        self.assertEqual(self.client.get(f"/studio/projects/{self.project.id}/").status_code, 200)
        self.assertEqual(self.client.get(f"/studio/projects/{self.other_project.id}/").status_code, 404)

    def test_viewer_cannot_edit_and_editor_can_edit(self):
        self.assertFalse(has_project_capability(self.viewer, self.project, "edit"))
        self.assertTrue(has_project_capability(self.editor, self.project, "edit"))
        self.client.force_login(self.viewer)
        self.assertEqual(self.client.get(f"/studio/projects/{self.project.id}/edit/").status_code, 403)
        self.client.force_login(self.editor)
        self.assertEqual(self.client.get(f"/studio/projects/{self.project.id}/edit/").status_code, 200)

    def test_full_control_can_manage_project_members(self):
        self.client.force_login(self.controller)
        response = self.client.post(
            f"/studio/projects/{self.project.id}/access/",
            {"email": self.outsider.email, "role": ProjectMembership.Role.VIEWER},
        )
        self.assertEqual(response.status_code, 302)
        self.assertTrue(ProjectMembership.objects.filter(project=self.project, user=self.outsider).exists())
        self.outsider.api_access.refresh_from_db()
        self.assertTrue(self.outsider.api_access.ai_api_enabled)

    def test_editor_cannot_manage_project_members(self):
        self.client.force_login(self.editor)
        self.assertEqual(self.client.get(f"/studio/projects/{self.project.id}/access/").status_code, 403)

    def test_full_control_can_manage_members_through_api(self):
        self.client.force_login(self.controller)
        response = self.client.post(
            f"/api/v1/studio/projects/{self.project.id}/members",
            data=json.dumps({"email": self.outsider.email, "role": ProjectMembership.Role.EDITOR}),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 201)
        self.outsider.api_access.refresh_from_db()
        self.assertTrue(self.outsider.api_access.ai_api_enabled)
        membership_id = response.json()["id"]
        self.assertEqual(self.client.get(f"/api/v1/studio/projects/{self.project.id}/members").status_code, 200)
        self.assertEqual(self.client.delete(f"/api/v1/studio/projects/{self.project.id}/members/{membership_id}").status_code, 200)
        self.outsider.api_access.refresh_from_db()
        self.assertTrue(self.outsider.api_access.ai_api_enabled)

    def test_viewer_cannot_manage_members_through_api(self):
        self.client.force_login(self.viewer)
        self.assertEqual(self.client.get(f"/api/v1/studio/projects/{self.project.id}/members").status_code, 403)

    def test_workspace_share_is_inherited_and_can_be_revoked_per_project(self):
        self.client.force_login(self.owner)
        shared = self.client.post(
            f"/studio/workspaces/{self.workspace.id}/access/",
            {"email": self.outsider.email, "role": WorkspaceMembership.Role.EDITOR, "can_use_ai": "on"},
        )
        self.assertEqual(shared.status_code, 302)
        self.assertTrue(accessible_projects(self.outsider).filter(id=self.project.id).exists())
        self.assertTrue(accessible_projects(self.outsider).filter(id=self.other_project.id).exists())

        revoked = self.client.post(
            f"/studio/projects/{self.other_project.id}/access/",
            {"action": "exclude", "user_id": str(self.outsider.id)},
        )
        self.assertEqual(revoked.status_code, 302)
        self.assertTrue(ProjectAccessExclusion.objects.filter(project=self.other_project, user=self.outsider).exists())
        self.assertFalse(accessible_projects(self.outsider).filter(id=self.other_project.id).exists())
        self.assertTrue(accessible_projects(self.outsider).filter(id=self.project.id).exists())

        direct = self.client.post(
            f"/studio/projects/{self.other_project.id}/access/",
            {"email": self.outsider.email, "role": ProjectMembership.Role.VIEWER},
        )
        self.assertEqual(direct.status_code, 302)
        self.assertFalse(ProjectAccessExclusion.objects.filter(project=self.other_project, user=self.outsider).exists())
        self.assertTrue(accessible_projects(self.outsider).filter(id=self.other_project.id).exists())

class StudioApiTests(TestCase):
    def setUp(self):
        users = get_user_model()
        self.owner = users.objects.create_user("api-owner", password="strong-pass")
        self.viewer = users.objects.create_user("api-viewer", password="strong-pass")
        self.workspace = create_workspace(user=self.owner, name="API Studio", slug="api-studio")
        WorkspaceMembership.objects.create(workspace=self.workspace, user=self.viewer, role=WorkspaceMembership.Role.VIEWER)

    def test_workspace_api_requires_login(self):
        self.assertEqual(self.client.get("/api/v1/studio/workspaces").status_code, 401)

    def test_owner_can_create_project(self):
        self.client.force_login(self.owner)
        response = self.client.post(
            "/api/v1/studio/projects",
            data=json.dumps({
                "workspaceId": str(self.workspace.id), "type": "SERIES", "title": "Dom",
                "promptTemplate": "No music. Keep character references.",
            }),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.json()["title"], "Dom")
        self.assertEqual(response.json()["promptTemplate"], "No music. Keep character references.")

    def test_viewer_cannot_create_project(self):
        self.client.force_login(self.viewer)
        response = self.client.post(
            "/api/v1/studio/projects",
            data=json.dumps({"workspaceId": str(self.workspace.id), "type": "SERIES", "title": "Denied"}),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 403)

class StudioSceneTests(TestCase):
    def setUp(self):
        from .models import AiModelProfile, DialogueLine, Episode, Prompt, PromptBlock, Scene
        users = get_user_model()
        self.editor = users.objects.create_user("scene-editor", password="strong-pass")
        self.translator = users.objects.create_user("scene-translator", password="strong-pass")
        self.workspace = create_workspace(user=self.editor, name="Scene Studio", slug="scene-studio")
        WorkspaceMembership.objects.create(workspace=self.workspace, user=self.translator, role=WorkspaceMembership.Role.TRANSLATOR)
        self.project = Project.objects.create(
            workspace=self.workspace, project_type=Project.Type.SERIES, title="Scene Project",
            created_by=self.editor, updated_by=self.editor,
        )
        self.episode = Episode.objects.create(
            project=self.project, number=1, title="Pilot", created_by=self.editor, updated_by=self.editor,
        )
        self.scene = Scene.objects.create(
            episode=self.episode, number=1, title="Opening", created_by=self.editor, updated_by=self.editor,
        )
        self.line = DialogueLine.objects.create(
            scene=self.scene, speaker="Hero", text="Original", status=DialogueLine.Status.APPROVED,
            created_by=self.editor, updated_by=self.editor,
        )
        self.model = AiModelProfile.objects.create(name="Test Video", provider="Test", model_id="test-video", media_type="VIDEO")
        self.prompt = Prompt.objects.create(
            scene=self.scene, ai_model=self.model, prompt_type=Prompt.Type.VIDEO,
            created_by=self.editor, updated_by=self.editor,
        )
        PromptBlock.objects.create(
            prompt=self.prompt, block_type=PromptBlock.Type.DIALOGUE_REFERENCE,
            content="Original", source_dialogue=self.line, created_by=self.editor, updated_by=self.editor,
        )

    def test_translator_cannot_change_approved_source_dialogue(self):
        self.client.force_login(self.translator)
        response = self.client.patch(
            f"/api/v1/studio/dialogue/{self.line.id}",
            data=json.dumps({"text": "Changed"}),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 403)
        self.line.refresh_from_db()
        self.assertEqual(self.line.text, "Original")

    def test_editor_changes_dialogue_and_marks_prompt_for_review(self):
        self.client.force_login(self.editor)
        response = self.client.patch(
            f"/api/v1/studio/dialogue/{self.line.id}",
            data=json.dumps({"text": "Changed"}),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 200)
        self.prompt.refresh_from_db()
        self.assertTrue(self.prompt.needs_review)

    def test_editor_can_create_structured_prompt(self):
        self.client.force_login(self.editor)
        response = self.client.post(
            f"/api/v1/studio/scenes/{self.scene.id}/prompts",
            data=json.dumps({
                "aiModelId": str(self.model.id),
                "type": "IMAGE",
                "blocks": [{"type": "NARRATIVE", "content": "Cinematic office"}, {"type": "NEGATIVE", "content": "No text"}],
            }),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 201)

class StudioAssetTests(TestCase):
    def setUp(self):
        import tempfile
        from lexamora_studio.models import Episode, Scene, private_storage

        self.temp_media = tempfile.TemporaryDirectory()
        self.private_storage = private_storage
        self.previous_location = private_storage._location
        private_storage._location = self.temp_media.name
        private_storage.__dict__.pop("base_location", None)
        private_storage.__dict__.pop("location", None)

        users = get_user_model()
        self.owner = users.objects.create_user("asset-owner", password="strong-pass")
        self.outsider = users.objects.create_user("asset-outsider", password="strong-pass")
        self.workspace = create_workspace(user=self.owner, name="Asset Studio", slug="asset-studio")
        self.other_workspace = create_workspace(user=self.outsider, name="Other Studio", slug="other-studio")
        self.project = Project.objects.create(
            workspace=self.workspace, project_type=Project.Type.SERIES, title="Asset Project",
            created_by=self.owner, updated_by=self.owner,
        )
        self.episode = Episode.objects.create(
            project=self.project, number=1, title="Pilot", created_by=self.owner, updated_by=self.owner,
        )
        self.scene = Scene.objects.create(
            episode=self.episode, number=1, title="Asset Scene", created_by=self.owner, updated_by=self.owner,
        )

    def tearDown(self):
        self.private_storage._location = self.previous_location
        self.private_storage.__dict__.pop("base_location", None)
        self.private_storage.__dict__.pop("location", None)
        self.temp_media.cleanup()

    def image_upload(self, name="scene.png"):
        import io
        from PIL import Image
        from django.core.files.uploadedfile import SimpleUploadedFile

        buffer = io.BytesIO()
        Image.new("RGB", (32, 24), "#f08a43").save(buffer, "PNG")
        return SimpleUploadedFile(name, buffer.getvalue(), content_type="image/png")

    def upload_asset(self):
        self.client.force_login(self.owner)
        return self.client.post(
            "/api/v1/studio/assets",
            data={
                "workspaceId": str(self.workspace.id),
                "projectId": str(self.project.id),
                "sceneId": str(self.scene.id),
                "kind": "SCENE_IMAGE",
                "file": self.image_upload(),
            },
        )

    def test_image_upload_generates_private_thumbnail(self):
        from .models import Asset

        response = self.upload_asset()
        self.assertEqual(response.status_code, 201, response.content)
        asset = Asset.objects.get(id=response.json()["id"])
        self.assertEqual((asset.width, asset.height), (32, 24))
        self.assertTrue(asset.thumbnail.name)
        self.assertTrue(asset.file.storage.exists(asset.file.name))
        self.assertTrue(asset.thumbnail.storage.exists(asset.thumbnail.name))

    def test_owner_can_view_original_inline_and_download_attachment(self):
        from .models import AccessEvent, AuditEvent

        response = self.upload_asset()
        asset_id = response.json()["id"]
        viewed = self.client.get(f"/api/v1/studio/assets/{asset_id}/view")
        self.assertEqual(viewed.status_code, 200)
        self.assertTrue(viewed["Content-Disposition"].startswith("inline;"))
        viewed.close()
        downloaded = self.client.get(f"/api/v1/studio/assets/{asset_id}/download")
        self.assertEqual(downloaded.status_code, 200)
        self.assertTrue(downloaded["Content-Disposition"].startswith("attachment;"))
        downloaded.close()
        self.assertTrue(AuditEvent.objects.filter(action="ASSET_VIEW", entity_id=asset_id).exists())
        self.assertTrue(AccessEvent.objects.filter(asset_id=asset_id, action="VIEW").exists())

    def test_video_view_supports_byte_ranges_for_fast_preview(self):
        from django.core.files.uploadedfile import SimpleUploadedFile

        self.client.force_login(self.owner)
        uploaded = self.client.post(
            "/api/v1/studio/assets",
            data={
                "workspaceId": str(self.workspace.id),
                "projectId": str(self.project.id),
                "kind": "GENERATION_OUTPUT",
                "file": SimpleUploadedFile("preview.mp4", b"0123456789abcdef", content_type="video/mp4"),
            },
        )
        self.assertEqual(uploaded.status_code, 201, uploaded.content)

        response = self.client.get(
            f"/api/v1/studio/assets/{uploaded.json()['id']}/view",
            HTTP_RANGE="bytes=4-9",
        )

        self.assertEqual(response.status_code, 206)
        self.assertEqual(response["Content-Range"], "bytes 4-9/16")
        self.assertEqual(response["Accept-Ranges"], "bytes")
        self.assertEqual(b"".join(response.streaming_content), b"456789")
    def test_user_from_another_workspace_cannot_download_asset(self):
        response = self.upload_asset()
        self.client.force_login(self.outsider)
        denied = self.client.get(f"/api/v1/studio/assets/{response.json()['id']}/download")
        self.assertEqual(denied.status_code, 404)

    def test_invalid_image_is_rejected(self):
        from django.core.files.uploadedfile import SimpleUploadedFile

        self.client.force_login(self.owner)
        response = self.client.post(
            "/api/v1/studio/assets",
            data={
                "workspaceId": str(self.workspace.id),
                "kind": "SCENE_IMAGE",
                "file": SimpleUploadedFile("bad.png", b"not an image", content_type="image/png"),
            },
        )
        self.assertEqual(response.status_code, 400)

    def test_generation_can_select_one_final_output(self):
        from .models import Asset, GenerationOutput

        first_response = self.upload_asset()
        second_response = self.upload_asset()
        self.client.force_login(self.owner)
        generation = self.client.post(
            f"/api/v1/studio/scenes/{self.scene.id}/generations",
            data=json.dumps({"reason": "Closer shot", "prompt": "Close-up of the glass"}),
            content_type="application/json",
        )
        self.assertEqual(generation.status_code, 201)
        output_ids = []
        for asset_id in (first_response.json()["id"], second_response.json()["id"]):
            output = self.client.post(
                f"/api/v1/studio/generations/{generation.json()['id']}/outputs",
                data=json.dumps({"assetId": asset_id, "modelMetadata": {"model": "Nano Banana 2"}}),
                content_type="application/json",
            )
            self.assertEqual(output.status_code, 201)
            output_ids.append(output.json()["id"])
        self.client.post(f"/api/v1/studio/generation-outputs/{output_ids[0]}/final")
        self.client.post(f"/api/v1/studio/generation-outputs/{output_ids[1]}/final")
        self.assertEqual(GenerationOutput.objects.filter(generation_id=generation.json()["id"], is_final=True).count(), 1)

    def test_upload_and_download_are_audited(self):
        from .models import AccessEvent, AuditEvent

        response = self.upload_asset()
        asset_id = response.json()["id"]
        self.assertTrue(AuditEvent.objects.filter(action="ASSET_UPLOAD", entity_id=asset_id).exists())
        download = self.client.get(f"/api/v1/studio/assets/{asset_id}/download")
        self.assertEqual(download.status_code, 200)
        download.close()
        self.assertTrue(AuditEvent.objects.filter(action="ASSET_DOWNLOAD", entity_id=asset_id).exists())
        self.assertTrue(AccessEvent.objects.filter(asset_id=asset_id, action="DOWNLOAD").exists())

class StudioRevisionTests(TestCase):
    def setUp(self):
        from .revisions import record_revision

        users = get_user_model()
        self.editor = users.objects.create_user("revision-editor", password="strong-pass")
        self.viewer = users.objects.create_user("revision-viewer", password="strong-pass")
        self.workspace = create_workspace(user=self.editor, name="Revision Studio", slug="revision-studio")
        WorkspaceMembership.objects.create(workspace=self.workspace, user=self.viewer, role=WorkspaceMembership.Role.VIEWER)
        self.project = Project.objects.create(
            workspace=self.workspace, project_type=Project.Type.SERIES, title="Original title",
            concept="First", created_by=self.editor, updated_by=self.editor,
        )
        self.initial = record_revision(instance=self.project, user=self.editor, operation="CREATE")

    def test_history_shows_author_and_timestamp(self):
        self.client.force_login(self.editor)
        self.client.patch(
            f"/api/v1/studio/projects/{self.project.id}",
            data=json.dumps({"title": "Changed title"}),
            content_type="application/json",
        )
        response = self.client.get(f"/api/v1/studio/entities/project/{self.project.id}/revisions")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.json()["results"]), 2)
        self.assertEqual(response.json()["results"][0]["author"], self.editor.username)
        self.assertTrue(response.json()["results"][0]["createdAt"])

    def test_restore_creates_new_revision(self):
        from .models import Revision

        self.client.force_login(self.editor)
        self.client.patch(
            f"/api/v1/studio/projects/{self.project.id}",
            data=json.dumps({"title": "Changed title"}),
            content_type="application/json",
        )
        response = self.client.post(
            f"/api/v1/studio/entities/project/{self.project.id}/restore/{self.initial.id}"
        )
        self.assertEqual(response.status_code, 200, response.content)
        self.project.refresh_from_db()
        self.assertEqual(self.project.title, "Original title")
        self.assertEqual(Revision.objects.filter(entity_id=self.project.id).count(), 3)
        self.assertEqual(Revision.objects.filter(entity_id=self.project.id).order_by("-sequence").first().operation, "RESTORE")

    def test_viewer_cannot_restore_revision(self):
        self.client.force_login(self.viewer)
        response = self.client.post(
            f"/api/v1/studio/entities/project/{self.project.id}/restore/{self.initial.id}"
        )
        self.assertEqual(response.status_code, 403)

    def test_revision_is_append_only(self):
        self.initial.operation = "TAMPERED"
        with self.assertRaises(ValueError):
            self.initial.save()


@override_settings(
    CREDENTIAL_ENCRYPTION_KEY=Fernet.generate_key().decode("ascii"),
    STUDIO_AI_RATE_PER_MINUTE=10,
)
class StudioAiSuggestionTests(TestCase):
    def setUp(self):
        from lessons.models import ProviderCredential, UserApiAccess
        from .models import AiModelProfile, DialogueLine, Episode, Prompt, PromptBlock, Scene

        users = get_user_model()
        self.editor = users.objects.create_user("ai-editor", password="strong-pass")
        self.viewer = users.objects.create_user("ai-viewer", password="strong-pass")
        access = UserApiAccess.objects.get(user=self.editor)
        access.ai_api_enabled = True
        access.save()
        credential = ProviderCredential(provider=ProviderCredential.Provider.OPENAI)
        credential.set_api_key("sk-test")
        credential.save()

        self.workspace = create_workspace(user=self.editor, name="AI Studio", slug="ai-studio")
        WorkspaceMembership.objects.create(workspace=self.workspace, user=self.viewer, role=WorkspaceMembership.Role.VIEWER)
        project = Project.objects.create(
            workspace=self.workspace, project_type=Project.Type.SERIES, title="AI Project",
            created_by=self.editor, updated_by=self.editor,
        )
        self.project = project
        episode = Episode.objects.create(project=project, number=1, title="Pilot", created_by=self.editor, updated_by=self.editor)
        scene = Scene.objects.create(episode=episode, number=1, title="Opening", created_by=self.editor, updated_by=self.editor)
        model = AiModelProfile.objects.create(name="AI Test Video", provider="Test", model_id="video-test", media_type="VIDEO")
        self.prompt = Prompt.objects.create(
            scene=scene, ai_model=model, prompt_type=Prompt.Type.VIDEO,
            created_by=self.editor, updated_by=self.editor,
        )
        self.narrative = PromptBlock.objects.create(
            prompt=self.prompt, block_type=PromptBlock.Type.NARRATIVE, content="Plain room",
            position=0, created_by=self.editor, updated_by=self.editor,
        )
        self.source_dialogue = DialogueLine.objects.create(
            scene=scene, speaker="Hero", text="Secret dialogue", language="en",
            position=0, created_by=self.editor, updated_by=self.editor,
        )
        self.dialogue = PromptBlock.objects.create(
            prompt=self.prompt, block_type=PromptBlock.Type.DIALOGUE_REFERENCE, content="Secret dialogue",
            source_dialogue=self.source_dialogue, position=1, created_by=self.editor, updated_by=self.editor,
        )

    def test_workspace_settings_show_image_model_and_token_usage(self):
        from .models import AiUsageLog

        AiUsageLog.objects.create(
            workspace=self.workspace, user=self.editor, prompt=self.prompt,
            action="GENERATE_IMAGE", model="gpt-image-1", status="SUCCESS",
            input_tokens=17, output_tokens=29, total_tokens=46,
        )
        self.client.force_login(self.editor)
        response = self.client.get(f"/studio/workspaces/{self.workspace.id}/edit/")
        self.assertEqual(response.status_code, 200, response.content)
        self.assertContains(response, "OpenAI image generation")
        self.assertContains(response, "gpt-image-1")
        self.assertContains(response, "AI token usage")
        self.assertContains(response, "46")

    def provider_response(self):
        return json.dumps({"blocks": [
            {"id": str(self.narrative.id), "content": "Cinematic room with precise lighting"},
            {"id": str(self.dialogue.id), "content": "Rewritten dialogue"},
        ]})

    @patch("lexamora_studio.views.run_text")
    def test_localized_field_translation_returns_both_dependent_languages(self, mocked_run_text):
        mocked_run_text.return_value = (
            json.dumps({"prompt": "A quiet office", "dialogue": "Una oficina tranquila"}),
            "gpt-5.4-mini",
        )
        self.client.force_login(self.editor)
        response = self.client.post(
            f"/studio/projects/{self.project.id}/localized-translate/",
            data=json.dumps({
                "source": "Spokojne biuro", "sourceLanguage": "PL",
                "targets": {"prompt": "EN", "dialogue": "ES"},
            }),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(response.json()["translations"]["prompt"], "A quiet office")
        self.assertEqual(response.json()["translations"]["dialogue"], "Una oficina tranquila")

    def test_inline_prompt_ui_is_compact_and_exposes_clear_ai_actions(self):
        self.client.force_login(self.editor)
        response = self.client.get(f"/studio/scenes/{self.prompt.scene_id}/")
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Added automatically")
        self.assertContains(response, "Copy")
        self.assertContains(response, "Paste")
        self.assertContains(response, "Clear")
        self.assertContains(response, "Suggest translation")
        self.assertContains(response, "Improve prompt")
        self.assertContains(response, "Entire prompt")
        self.assertContains(response, "Dialogue only")
        self.assertContains(response, "Selected text")
        self.assertContains(response, "Apply Translation")
        self.assertContains(response, "Improve Translation")
        self.assertContains(response, "data-prompt-toggle")
        self.assertContains(response, '<option value="BY">BY</option>', html=True)
        self.assertContains(response, "gpt-5.4-mini")
        self.assertNotContains(response, "Mandatory template")
        self.assertNotContains(response, "Improve action")

    @patch("lexamora_studio.ai.run_text")
    def test_prompt_translation_preview_and_apply_create_language_version(self, mocked_run_text):
        mocked_run_text.return_value = (json.dumps({"content": "Filmowy pokoj"}), "gpt-5.4-mini")
        self.client.force_login(self.editor)
        preview = self.client.post(
            f"/studio/prompts/{self.prompt.id}/ai-preview/",
            data=json.dumps({
                "action": "translate", "content": "Cinematic room", "targetLanguage": "PL",
                "scope": "FULL", "textModel": "gpt-5.4-mini",
            }),
            content_type="application/json",
        )
        self.assertEqual(preview.status_code, 200, preview.content)
        self.assertEqual(preview.json()["content"], "Filmowy pokoj")

        applied = self.client.post(
            f"/studio/prompts/{self.prompt.id}/apply-translation/",
            data=json.dumps({"content": preview.json()["content"], "targetLanguage": "PL", "scope": "FULL"}),
            content_type="application/json",
        )
        self.assertEqual(applied.status_code, 200, applied.content)
        translated = Prompt.objects.get(source_prompt=self.prompt, language="PL")
        self.assertEqual(translated.content, "Filmowy pokoj")
        self.assertEqual(translated.blocks.get().content, "Filmowy pokoj")
        self.assertEqual({row["language"] for row in applied.json()["versions"]}, {"EN", "PL"})

    @patch("lexamora_studio.ai.run_text")
    def test_prompt_improvement_preview_keeps_original_language(self, mocked_run_text):
        mocked_run_text.return_value = (json.dumps({"content": "A precise cinematic room."}), "gpt-5.4-mini")
        self.client.force_login(self.editor)
        response = self.client.post(
            f"/studio/prompts/{self.prompt.id}/ai-preview/",
            data=json.dumps({
                "action": "improve_prompt", "content": "Room cinematic", "targetLanguage": "PL",
                "scope": "FULL", "textModel": "gpt-5.4-mini",
            }),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(response.json()["language"], "EN")
        self.assertEqual(response.json()["content"], "A precise cinematic room.")
        provider_payload = json.loads(mocked_run_text.call_args.args[1])
        self.assertIn("without translating", provider_payload["task"])

    @patch("lexamora_studio.image_jobs.generate_image_with_usage")
    def test_photo_prompt_generation_returns_asset_json(self, mocked_generate_image):
        import io
        import tempfile
        from pathlib import Path
        from PIL import Image
        from django.core.files.uploadedfile import SimpleUploadedFile
        from .image_jobs import execute_image_generation_job
        from .models import Asset, ImageGenerationJob
        from .storage import create_asset

        output = io.BytesIO()
        Image.new("RGB", (48, 48), "#2a8b69").save(output, "PNG")
        mocked_generate_image.return_value = (output.getvalue(), "gpt-image-1", {"total_tokens": 42})
        self.prompt.prompt_type = Prompt.Type.IMAGE
        self.prompt.content = "A production still"
        self.prompt.save(update_fields=["prompt_type", "content", "updated_at"])
        self.client.force_login(self.editor)
        editor_page = self.client.get(f"/studio/scenes/{self.prompt.scene_id}/")
        self.assertContains(editor_page, "data-image-request-status")
        self.assertContains(editor_page, "data-generation-composer")
        self.assertContains(editor_page, "data-composer-media-type")
        self.assertContains(editor_page, "References")
        self.assertContains(editor_page, "1536 x 1024 - Landscape")
        self.assertContains(editor_page, "1024 x 1536 - Portrait")
        self.assertContains(editor_page, ">Close</button>", html=False)
        with tempfile.TemporaryDirectory() as directory:
            with self.settings(STUDIO_PRIVATE_MEDIA_ROOT=Path(directory)):
                reference_file = SimpleUploadedFile("reference.png", output.getvalue(), content_type="image/png")
                reference = create_asset(
                    user=self.editor, workspace=self.workspace, project=self.project,
                    prompt=self.prompt, uploaded=reference_file, kind=Asset.Kind.OTHER,
                )
                response = self.client.post(
                    f"/studio/prompts/{self.prompt.id}/images/generate/",
                    data=json.dumps({
                        "prompt": "An improved current prompt with exact lighting",
                        "size": "1536x1024",
                        "quality": "high",
                        "outputFormat": "png",
                        "outputCompression": 90,
                        "background": "opaque",
                        "moderation": "low",
                    }),
                    content_type="application/json",
                )
                self.assertEqual(response.status_code, 202, response.content)
                self.assertEqual(response.json()["status"], "QUEUED")
                jobs_response = self.client.get(f"/studio/prompts/{self.prompt.id}/images/jobs/")
                self.assertEqual(jobs_response.json()["jobs"][0]["jobId"], response.json()["jobId"])
                job = ImageGenerationJob.objects.get(id=response.json()["jobId"])
                self.assertEqual(job.request_prompt, "An improved current prompt with exact lighting")
                job.status = ImageGenerationJob.Status.RUNNING
                job.started_at = timezone.now()
                job.save(update_fields=["status", "started_at", "updated_at"])
                execute_image_generation_job(job.id)
                status_response = self.client.get(response.json()["statusUrl"])
                self.assertEqual(status_response.status_code, 200)
                self.assertEqual(status_response.json()["status"], "SUCCESS")
                self.assertTrue(status_response.json()["thumbnailUrl"])
        self.assertTrue(self.prompt.reference_assets.filter(id=reference.id).exists())
        generated = Asset.objects.get(id=status_response.json()["assetId"])
        self.assertEqual(generated.ai_metadata["model"], "gpt-image-1")
        self.assertEqual(generated.ai_metadata["referenceAssetIds"], [str(reference.id)])
        self.assertEqual(generated.ai_metadata["requestPrompt"], "An improved current prompt with exact lighting")
        self.assertEqual(generated.ai_metadata["settings"]["size"], "1536x1024")
        self.assertEqual(generated.ai_metadata["settings"]["quality"], "high")
        call_kwargs = mocked_generate_image.call_args.kwargs
        self.assertEqual(mocked_generate_image.call_args.args[0], "An improved current prompt with exact lighting")
        self.assertEqual(call_kwargs["model"], "gpt-image-1")
        self.assertEqual(call_kwargs["reference_images"][0][0], "reference.png")
        self.assertEqual(call_kwargs["size"], "1536x1024")
        self.assertEqual(call_kwargs["quality"], "high")
        self.assertEqual(call_kwargs["output_format"], "png")
        self.assertEqual(call_kwargs["background"], "opaque")
        self.assertEqual(call_kwargs["moderation"], "low")
        self.assertNotIn("composition_preset", call_kwargs)
        self.client.force_login(self.viewer)
        shared_status = self.client.get(response.json()["statusUrl"])
        self.assertEqual(shared_status.status_code, 200)
        self.assertEqual(shared_status.json()["status"], "SUCCESS")

    def test_image_worker_claims_five_distinct_jobs(self):
        from .management.commands.run_image_generation_worker import Command
        from .models import AiModelProfile, ImageGenerationJob

        image_model = AiModelProfile.objects.filter(media_type=AiModelProfile.MediaType.IMAGE).first()
        jobs = [
            ImageGenerationJob.objects.create(
                workspace=self.workspace,
                prompt=self.prompt,
                requested_by=self.editor,
                model_profile=image_model,
                request_prompt=f"Image request {index}",
                options={"size": "1024x1024", "quality": "low", "output_format": "png"},
            )
            for index in range(5)
        ]
        claimed = [Command._claim_next_job() for _ in range(5)]
        self.assertEqual(len(set(claimed)), 5)
        self.assertEqual(set(claimed), {job.id for job in jobs})
        self.assertEqual(
            ImageGenerationJob.objects.filter(status=ImageGenerationJob.Status.RUNNING).count(),
            5,
        )

    @patch("lexamora_studio.ai.run_text")
    def test_selected_text_preview_preserves_unselected_prompt(self, mocked_run_text):
        mocked_run_text.return_value = (json.dumps({"content": "pokoj"}), "gpt-5.4-mini")
        self.client.force_login(self.editor)
        source = "A room at night"
        response = self.client.post(
            f"/studio/prompts/{self.prompt.id}/ai-preview/",
            data=json.dumps({
                "action": "translate", "content": source, "targetLanguage": "PL",
                "scope": "SELECTED", "selectionStart": 2, "selectionEnd": 6,
                "textModel": "gpt-5.4-mini",
            }),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(response.json()["content"], "A pokoj at night")
        provider_payload = json.loads(mocked_run_text.call_args.args[1])
        self.assertEqual(provider_payload["content"], "room")

    @patch("lexamora_studio.ai.run_text")
    def test_dialogue_preview_updates_spoken_language_label(self, mocked_run_text):
        mocked_run_text.return_value = (
            json.dumps({"content": 'Hero says in Polish: "Dzien dobry". Camera stays wide.'}),
            "gpt-5.4-mini",
        )
        self.client.force_login(self.editor)
        response = self.client.post(
            f"/studio/prompts/{self.prompt.id}/ai-preview/",
            data=json.dumps({
                "action": "translate", "content": 'Hero says in Polish: "Dzien dobry". Camera stays wide.',
                "targetLanguage": "RU", "scope": "DIALOGUE", "textModel": "gpt-5.4-mini",
            }),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(response.json()["content"], 'Hero says in Russian: "Dzien dobry". Camera stays wide.')
        provider_payload = json.loads(mocked_run_text.call_args.args[1])
        self.assertIn("in Russian", " ".join(provider_payload["rules"]))

    @patch("lexamora_studio.ai.run_text")
    def test_improve_creates_suggestion_without_overwriting_or_sending_dialogue(self, mocked_run_text):
        from .models import AiSuggestion, AiUsageLog

        mocked_run_text.return_value = (self.provider_response(), "gpt-5.4-mini")
        self.client.force_login(self.editor)
        response = self.client.post(
            f"/api/v1/studio/prompts/{self.prompt.id}/improve",
            data=json.dumps({"mode": "cinematic", "textModel": "gpt-5.4-mini"}),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 201, response.content)
        self.narrative.refresh_from_db()
        self.assertEqual(self.narrative.content, "Plain room")
        self.assertEqual(AiSuggestion.objects.get().status, AiSuggestion.Status.PENDING)
        self.assertEqual(AiUsageLog.objects.get().status, "SUCCESS")
        provider_prompt = mocked_run_text.call_args.args[1]
        self.assertNotIn("Secret dialogue", provider_prompt)
        self.assertIn(self.prompt.template.content, provider_prompt)

    @patch("lexamora_studio.ai.run_text")
    def test_russian_prompt_can_be_improved_and_translated_to_english(self, mocked_run_text):
        self.narrative.content = "Простая комната и герой входит"
        self.narrative.save(update_fields=["content", "updated_at"])
        mocked_run_text.return_value = (
            json.dumps({"blocks": [{"id": str(self.narrative.id), "content": "A cinematic room as the hero enters."}]}),
            "gpt-5.4",
        )
        self.client.force_login(self.editor)
        response = self.client.post(
            f"/api/v1/studio/prompts/{self.prompt.id}/improve",
            data=json.dumps({"mode": "improve_translate_en", "textModel": "gpt-5.4"}),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 201, response.content)
        self.assertEqual(response.json()["model"], "gpt-5.4")
        provider_prompt = mocked_run_text.call_args.args[1]
        self.assertIn("Return every editable block in English", provider_prompt)

    @patch("lexamora_studio.ai.run_text")
    def test_dialogue_translation_creates_new_prompt_and_preserves_source(self, mocked_run_text):
        from .models import AiUsageLog, Prompt, TranslationUnit

        mocked_run_text.return_value = (
            json.dumps({"blocks": [{"id": str(self.dialogue.id), "content": "Tajny dialog"}]}),
            "gpt-5.4-mini",
        )
        self.client.force_login(self.editor)
        response = self.client.post(
            f"/api/v1/studio/prompts/{self.prompt.id}/translate-dialogue",
            data=json.dumps({"targetLanguage": "PL", "textModel": "gpt-5.4-mini"}),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 201, response.content)
        self.narrative.refresh_from_db()
        self.dialogue.refresh_from_db()
        self.assertEqual(self.narrative.content, "Plain room")
        self.assertEqual(self.dialogue.content, "Secret dialogue")
        translated_prompt = Prompt.objects.get(id=response.json()["promptId"])
        translated_blocks = list(translated_prompt.blocks.order_by("position"))
        self.assertEqual(translated_prompt.source_prompt, self.prompt)
        self.assertEqual(translated_prompt.language, "PL")
        self.assertEqual(translated_prompt.translation_scope, Prompt.TranslationScope.DIALOGUE)
        self.assertEqual(translated_blocks[0].content, "Plain room")
        self.assertEqual(translated_blocks[1].content, "Tajny dialog")
        self.assertTrue(TranslationUnit.objects.filter(
            dialogue_line=self.source_dialogue, target_language="PL", translated_text="Tajny dialog"
        ).exists())
        self.assertTrue(AiUsageLog.objects.filter(action="TRANSLATE_DIALOGUE", status="SUCCESS").exists())
        provider_prompt = mocked_run_text.call_args.args[1]
        self.assertNotIn("Plain room", provider_prompt)

    @patch("lexamora_studio.ai.run_text")
    def test_full_translation_creates_new_prompt_with_every_block_translated(self, mocked_run_text):
        from .models import Prompt

        mocked_run_text.return_value = (
            json.dumps({"blocks": [
                {"id": str(self.narrative.id), "content": "Pokoj"},
                {"id": str(self.dialogue.id), "content": "Tajny dialog"},
            ]}),
            "gpt-5.4-mini",
        )
        self.client.force_login(self.editor)
        response = self.client.post(
            f"/api/v1/studio/prompts/{self.prompt.id}/translate",
            data=json.dumps({"targetLanguage": "PL", "scope": "FULL", "textModel": "gpt-5.4-mini"}),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 201, response.content)
        translated_prompt = Prompt.objects.get(id=response.json()["promptId"])
        self.assertEqual(translated_prompt.translation_scope, Prompt.TranslationScope.FULL)
        self.assertEqual(list(translated_prompt.blocks.values_list("content", flat=True)), ["Pokoj", "Tajny dialog"])

    @patch("lexamora_studio.ai.run_text")
    def test_accept_updates_non_dialogue_and_preserves_dialogue(self, mocked_run_text):
        from .models import AiSuggestion, Revision

        mocked_run_text.return_value = (self.provider_response(), "gpt-5.4-mini")
        self.client.force_login(self.editor)
        created = self.client.post(
            f"/api/v1/studio/prompts/{self.prompt.id}/improve",
            data=json.dumps({"mode": "non_dialogue"}),
            content_type="application/json",
        )
        accepted = self.client.post(f"/api/v1/studio/suggestions/{created.json()['id']}/accept")
        self.assertEqual(accepted.status_code, 200, accepted.content)
        self.narrative.refresh_from_db()
        self.dialogue.refresh_from_db()
        self.assertEqual(self.narrative.content, "Cinematic room with precise lighting")
        self.assertEqual(self.dialogue.content, "Secret dialogue")
        self.assertEqual(AiSuggestion.objects.get().status, AiSuggestion.Status.ACCEPTED)
        self.assertTrue(Revision.objects.filter(entity_id=self.prompt.id, operation="AI_ACCEPT").exists())

        undone = self.client.post(f"/api/v1/studio/suggestions/{created.json()['id']}/undo")
        self.assertEqual(undone.status_code, 200, undone.content)
        self.narrative.refresh_from_db()
        self.assertEqual(self.narrative.content, "Plain room")
        self.assertEqual(AiSuggestion.objects.get().status, AiSuggestion.Status.UNDONE)
        self.assertTrue(Revision.objects.filter(entity_id=self.prompt.id, operation="AI_UNDO").exists())

    @patch("lexamora_studio.ai.run_text")
    def test_reject_keeps_prompt_unchanged(self, mocked_run_text):
        mocked_run_text.return_value = (self.provider_response(), "gpt-5.4-mini")
        self.client.force_login(self.editor)
        created = self.client.post(
            f"/api/v1/studio/prompts/{self.prompt.id}/improve",
            data=json.dumps({"mode": "shorter"}),
            content_type="application/json",
        )
        rejected = self.client.post(f"/api/v1/studio/suggestions/{created.json()['id']}/reject")
        self.assertEqual(rejected.status_code, 200)
        self.narrative.refresh_from_db()
        self.assertEqual(self.narrative.content, "Plain room")

    def test_user_without_ai_capability_cannot_improve(self):
        self.client.force_login(self.viewer)
        response = self.client.post(
            f"/api/v1/studio/prompts/{self.prompt.id}/improve",
            data=json.dumps({"mode": "non_dialogue"}),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 403)
    @override_settings(STUDIO_AI_RATE_PER_MINUTE=1)
    @patch("lexamora_studio.ai.run_text")
    def test_ai_rate_limit_returns_429(self, mocked_run_text):
        mocked_run_text.return_value = (self.provider_response(), "gpt-5.4-mini")
        self.client.force_login(self.editor)
        first = self.client.post(
            f"/api/v1/studio/prompts/{self.prompt.id}/improve",
            data=json.dumps({"mode": "non_dialogue"}),
            content_type="application/json",
        )
        second = self.client.post(
            f"/api/v1/studio/prompts/{self.prompt.id}/improve",
            data=json.dumps({"mode": "non_dialogue"}),
            content_type="application/json",
        )
        self.assertEqual(first.status_code, 201)
        self.assertEqual(second.status_code, 429)

    @patch("lexamora_studio.ai.run_text")
    def test_stale_suggestion_cannot_be_accepted(self, mocked_run_text):
        from .revisions import record_revision

        mocked_run_text.return_value = (self.provider_response(), "gpt-5.4-mini")
        self.client.force_login(self.editor)
        created = self.client.post(
            f"/api/v1/studio/prompts/{self.prompt.id}/improve",
            data=json.dumps({"mode": "non_dialogue"}),
            content_type="application/json",
        )
        self.prompt.title = "Changed after suggestion"
        self.prompt.updated_by = self.editor
        self.prompt.save()
        record_revision(instance=self.prompt, user=self.editor, operation="UPDATE")
        accepted = self.client.post(f"/api/v1/studio/suggestions/{created.json()['id']}/accept")
        self.assertEqual(accepted.status_code, 409)
        self.narrative.refresh_from_db()
        self.assertEqual(self.narrative.content, "Plain room")


class StudioSubtitleTranslationTests(TestCase):
    def setUp(self):
        from .models import AiModelProfile, DialogueLine, Episode, Prompt, PromptBlock, Scene

        users = get_user_model()
        self.editor = users.objects.create_user("subtitle-editor", password="strong-pass")
        self.translator = users.objects.create_user("subtitle-translator", password="strong-pass")
        self.viewer = users.objects.create_user("subtitle-viewer", password="strong-pass")
        self.workspace = create_workspace(user=self.editor, name="Subtitle Studio", slug="subtitle-studio")
        WorkspaceMembership.objects.create(workspace=self.workspace, user=self.translator, role=WorkspaceMembership.Role.TRANSLATOR)
        WorkspaceMembership.objects.create(workspace=self.workspace, user=self.viewer, role=WorkspaceMembership.Role.VIEWER)
        self.project = Project.objects.create(
            workspace=self.workspace, project_type=Project.Type.SERIES, title="Subtitle Project",
            created_by=self.editor, updated_by=self.editor,
        )
        self.episode = Episode.objects.create(
            project=self.project, number=1, title="Pilot", created_by=self.editor, updated_by=self.editor,
        )
        self.scene = Scene.objects.create(
            episode=self.episode, number=1, title="Opening", created_by=self.editor, updated_by=self.editor,
        )
        self.line = DialogueLine.objects.create(
            scene=self.scene, speaker="Hero", text="Original source", status=DialogueLine.Status.APPROVED,
            created_by=self.editor, updated_by=self.editor,
        )
        model = AiModelProfile.objects.create(name="Subtitle Test", provider="Test", model_id="test", media_type="VIDEO")
        self.prompt = Prompt.objects.create(
            scene=self.scene, ai_model=model, prompt_type=Prompt.Type.VIDEO,
            created_by=self.editor, updated_by=self.editor,
        )
        PromptBlock.objects.create(
            prompt=self.prompt, block_type=PromptBlock.Type.DIALOGUE_REFERENCE,
            content="Original source", source_dialogue=self.line, created_by=self.editor, updated_by=self.editor,
        )

    def test_translator_can_translate_but_cannot_edit_source(self):
        from .models import TranslationUnit

        self.client.force_login(self.translator)
        translated = self.client.patch(
            f"/api/v1/studio/dialogue/{self.line.id}/translations/pl",
            data=json.dumps({"translation": "Tekst docelowy", "status": "APPROVED"}),
            content_type="application/json",
        )
        self.assertEqual(translated.status_code, 200, translated.content)
        self.assertEqual(TranslationUnit.objects.get().status, TranslationUnit.Status.APPROVED)
        source_edit = self.client.patch(
            f"/api/v1/studio/dialogue/{self.line.id}",
            data=json.dumps({"text": "Forbidden"}),
            content_type="application/json",
        )
        self.assertEqual(source_edit.status_code, 403)

    def test_source_change_marks_translation_stale_and_prompt_for_review(self):
        from .models import TranslationUnit
        from .services import save_translation

        save_translation(
            dialogue_line=self.line, user=self.translator, target_language="pl",
            translated_text="Tekst docelowy", status=TranslationUnit.Status.APPROVED,
        )
        self.client.force_login(self.editor)
        changed = self.client.patch(
            f"/api/v1/studio/dialogue/{self.line.id}",
            data=json.dumps({"text": "Updated source"}),
            content_type="application/json",
        )
        self.assertEqual(changed.status_code, 200)
        self.assertEqual(TranslationUnit.objects.get().status, TranslationUnit.Status.STALE)
        self.prompt.refresh_from_db()
        self.assertTrue(self.prompt.needs_review)

    def test_bulk_subtitles_and_reorder(self):
        from .models import SubtitleLine

        self.client.force_login(self.translator)
        created = self.client.post(
            f"/api/v1/studio/episodes/{self.episode.id}/subtitle-tracks",
            data=json.dumps({"language": "pl", "kind": "WORKING"}),
            content_type="application/json",
        )
        self.assertEqual(created.status_code, 201)
        track_id = created.json()["id"]
        pasted = self.client.post(
            f"/api/v1/studio/subtitle-tracks/{track_id}/lines",
            data=json.dumps({"lines": ["FIRST", "SECOND", "THIRD"]}),
            content_type="application/json",
        )
        self.assertEqual(pasted.status_code, 200, pasted.content)
        ids = [row["id"] for row in pasted.json()["results"]]
        reordered = self.client.post(
            f"/api/v1/studio/subtitle-tracks/{track_id}/lines/reorder",
            data=json.dumps({"lineIds": list(reversed(ids))}),
            content_type="application/json",
        )
        self.assertEqual(reordered.status_code, 200)
        values = list(SubtitleLine.objects.filter(track_id=track_id).values_list("text", flat=True))
        self.assertEqual(values, ["THIRD", "SECOND", "FIRST"])

    def test_viewer_cannot_replace_subtitles(self):
        from .models import SubtitleTrack

        track = SubtitleTrack.objects.create(
            episode=self.episode, language="pl", kind=SubtitleTrack.Kind.WORKING,
            created_by=self.editor, updated_by=self.editor,
        )
        self.client.force_login(self.viewer)
        response = self.client.post(
            f"/api/v1/studio/subtitle-tracks/{track.id}/lines",
            data=json.dumps({"lines": ["Denied"]}),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 403)


@override_settings(
    STUDIO_PDF_FONT_PATH="C:/Windows/Fonts/arial.ttf",
    STUDIO_EXPORT_RATE_PER_HOUR=10,
)
class StudioExportTests(TestCase):
    def setUp(self):
        import tempfile
        from .models import DialogueLine, Episode, Scene, SubtitleLine, SubtitleTrack, private_storage

        self.temp_media = tempfile.TemporaryDirectory()
        self.private_storage = private_storage
        self.previous_location = private_storage._location
        private_storage._location = self.temp_media.name
        private_storage.__dict__.pop("base_location", None)
        private_storage.__dict__.pop("location", None)

        users = get_user_model()
        self.owner = users.objects.create_user("export-owner", password="strong-pass")
        self.viewer = users.objects.create_user("export-viewer", password="strong-pass")
        self.workspace = create_workspace(user=self.owner, name="Export Studio", slug="export-studio")
        WorkspaceMembership.objects.create(workspace=self.workspace, user=self.viewer, role=WorkspaceMembership.Role.VIEWER)
        self.project = Project.objects.create(
            workspace=self.workspace, project_type=Project.Type.SERIES,
            title="Dom serial", concept="Historia domu", original_language="pl",
            translation_languages=["ru"], rights_holder="Leksa Programs",
            created_by=self.owner, updated_by=self.owner,
        )
        self.episode = Episode.objects.create(
            project=self.project, number=1, title="Ostatnie zadanie",
            summary="Pilot episode", created_by=self.owner, updated_by=self.owner,
        )
        scene = Scene.objects.create(
            episode=self.episode, number=1, title="Office",
            description="A cinematic office scene", created_by=self.owner, updated_by=self.owner,
        )
        DialogueLine.objects.create(
            scene=scene, speaker="Hero", text="Dzien dobry",
            created_by=self.owner, updated_by=self.owner,
        )
        track = SubtitleTrack.objects.create(
            episode=self.episode, language="pl", kind=SubtitleTrack.Kind.WORKING,
            created_by=self.owner, updated_by=self.owner,
        )
        SubtitleLine.objects.create(
            track=track, position=0, text="DZIEN DOBRY",
            created_by=self.owner, updated_by=self.owner,
        )

    def tearDown(self):
        self.private_storage._location = self.previous_location
        self.private_storage.__dict__.pop("base_location", None)
        self.private_storage.__dict__.pop("location", None)
        self.temp_media.cleanup()

    def test_pdf_export_creates_private_asset_and_audit_history(self):
        from .models import AccessEvent, AuditEvent, ExportJob

        self.client.force_login(self.owner)
        response = self.client.post(
            f"/api/v1/studio/projects/{self.project.id}/export/pdf",
            data=json.dumps({"sections": ["metadata", "episodes", "scenes", "dialogue", "subtitles"]}),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 201, response.content)
        job = ExportJob.objects.get(id=response.json()["id"])
        self.assertEqual(job.status, ExportJob.Status.SUCCESS)
        self.assertTrue(job.output_asset.file.storage.exists(job.output_asset.file.name))
        with job.output_asset.file.open("rb") as stream:
            self.assertEqual(stream.read(4), b"%PDF")
        self.assertTrue(AuditEvent.objects.filter(action="EXPORT_GENERATED").exists())

        download = self.client.get(f"/api/v1/studio/exports/{job.id}/download")
        self.assertEqual(download.status_code, 200)
        download.close()
        self.assertTrue(AuditEvent.objects.filter(action="EXPORT_DOWNLOAD").exists())
        self.assertTrue(AccessEvent.objects.filter(asset=job.output_asset, action="DOWNLOAD").exists())

    def test_viewer_cannot_generate_export(self):
        self.client.force_login(self.viewer)
        response = self.client.post(
            f"/api/v1/studio/projects/{self.project.id}/export/pdf",
            data=json.dumps({"sections": []}),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 403)

    @override_settings(STUDIO_EXPORT_RATE_PER_HOUR=1)
    def test_export_rate_limit(self):
        self.client.force_login(self.owner)
        first = self.client.post(
            f"/api/v1/studio/projects/{self.project.id}/export/pdf",
            data=json.dumps({"sections": ["metadata"]}),
            content_type="application/json",
        )
        second = self.client.post(
            f"/api/v1/studio/projects/{self.project.id}/export/pdf",
            data=json.dumps({"sections": ["metadata"]}),
            content_type="application/json",
        )
        self.assertEqual(first.status_code, 201)
        self.assertEqual(second.status_code, 429)

class StudioWebCreationTests(TestCase):
    def setUp(self):
        users = get_user_model()
        self.owner = users.objects.create_user("web-owner", password="strong-pass")
        self.viewer = users.objects.create_user("web-viewer", password="strong-pass")
        self.workspace = create_workspace(user=self.owner, name="Web Studio", slug="web-studio")
        WorkspaceMembership.objects.create(workspace=self.workspace, user=self.viewer, role=WorkspaceMembership.Role.VIEWER)

    def test_owner_can_create_project_episode_and_scene_in_web_ui(self):
        from .models import Episode, Scene

        self.client.force_login(self.owner)
        response = self.client.post(
            f"/studio/workspaces/{self.workspace.id}/projects/new/",
            {"project_type": "SERIES", "title": "Browser Project", "original_language": "ru", "translation_languages": "en, pl", "status": "DRAFT"},
        )
        project = Project.objects.get(title="Browser Project")
        self.assertRedirects(response, f"/studio/projects/{project.id}/")
        self.assertEqual(project.translation_languages, [])

        response = self.client.post(
            f"/studio/projects/{project.id}/episodes/new/",
            {"number": 1, "title": "Pilot"},
        )
        episode = Episode.objects.get(project=project)
        self.assertRedirects(response, f"/studio/projects/{project.id}/")

        response = self.client.post(
            f"/studio/episodes/{episode.id}/scenes/new/",
            {"number": 1, "title": "Opening", "status": "DRAFT"},
        )
        scene = Scene.objects.get(episode=episode)
        self.assertRedirects(response, f"/studio/scenes/{scene.id}/")

    def test_viewer_cannot_open_project_creation_form(self):
        self.client.force_login(self.viewer)
        response = self.client.get(f"/studio/workspaces/{self.workspace.id}/projects/new/")
        self.assertEqual(response.status_code, 403)

class StudioWebEditingAndImagesTests(TestCase):
    def setUp(self):
        from .models import Episode, Scene

        users = get_user_model()
        self.owner = users.objects.create_user("web-edit-owner", password="strong-pass")
        self.viewer = users.objects.create_user("web-edit-viewer", password="strong-pass")
        self.outsider = users.objects.create_user("web-edit-outsider", password="strong-pass")
        self.workspace = create_workspace(user=self.owner, name="Editable Studio", slug="editable-studio")
        WorkspaceMembership.objects.create(workspace=self.workspace, user=self.viewer, role=WorkspaceMembership.Role.VIEWER)
        self.project = Project.objects.create(workspace=self.workspace, project_type=Project.Type.SERIES, title="Before", original_language="ru", translation_languages=["en"], created_by=self.owner, updated_by=self.owner)
        self.episode = Episode.objects.create(project=self.project, number=1, title="Pilot", created_by=self.owner, updated_by=self.owner)
        self.scene = Scene.objects.create(episode=self.episode, number=1, title="Opening", created_by=self.owner, updated_by=self.owner)

    @staticmethod
    def image_file(name="reference.png", color="#36a889"):
        import io
        from PIL import Image
        from django.core.files.uploadedfile import SimpleUploadedFile

        buffer = io.BytesIO()
        Image.new("RGB", (48, 36), color).save(buffer, "PNG")
        return SimpleUploadedFile(name, buffer.getvalue(), content_type="image/png")

    def test_owner_can_edit_project_and_revision_is_recorded(self):
        from .models import Revision

        self.client.force_login(self.owner)
        response = self.client.post(
            f"/studio/projects/{self.project.id}/edit/",
            {
                "project_type": "SERIES", "title": "After", "concept": "Updated", "status": "DRAFT",
                "tracks-TOTAL_FORMS": "0", "tracks-INITIAL_FORMS": "0",
                "tracks-MIN_NUM_FORMS": "0", "tracks-MAX_NUM_FORMS": "1000",
            },
        )
        self.assertRedirects(response, f"/studio/projects/{self.project.id}/edit/")
        self.project.refresh_from_db()
        self.assertEqual(self.project.title, "After")
        self.assertEqual(self.project.translation_languages, ["en"])
        self.assertTrue(Revision.objects.filter(entity_id=self.project.id, operation="UPDATE").exists())

    def test_project_settings_save_languages_and_template(self):
        from .models import RecommendedTrack, StudioTextModel

        self.client.force_login(self.owner)
        translation_model = StudioTextModel.objects.get(model_id="gpt-5.4-nano")
        response = self.client.post(
            f"/studio/projects/{self.project.id}/settings/",
            {
                "original_language": "RU", "documentation_language": "RU",
                "dialogue_language": "PL", "prompt_language": "EN",
                "translation_languages": "pl, en", "prompt_template": "No music.",
                "default_translation_model": translation_model.id,
            },
        )
        self.assertRedirects(response, f"/studio/projects/{self.project.id}/settings/")
        self.project.refresh_from_db()
        self.assertEqual((self.project.documentation_language, self.project.dialogue_language, self.project.prompt_language), ("RU", "PL", "EN"))
        self.assertEqual(self.project.prompt_template, "No music.")
        self.assertEqual(self.project.default_translation_model, translation_model)
        edit = self.client.post(
            f"/studio/projects/{self.project.id}/edit/",
            {
                "project_type": "SERIES", "title": self.project.title, "description": "",
                "concept": "", "rights_holder": "", "publication_info": "",
                "status": "DRAFT", "status_comment": "",
                "tracks-TOTAL_FORMS": "1", "tracks-INITIAL_FORMS": "0",
                "tracks-MIN_NUM_FORMS": "0", "tracks-MAX_NUM_FORMS": "1000",
                "tracks-0-is_primary": "on", "tracks-0-platform": "Spotify",
                "tracks-0-artist": "Artist", "tracks-0-title": "Track",
                "tracks-0-url": "https://example.com/track", "tracks-0-position": "0",
            },
        )
        self.assertRedirects(edit, f"/studio/projects/{self.project.id}/edit/")
        track = RecommendedTrack.objects.get(project=self.project)
        self.assertTrue(track.is_primary)
        self.assertEqual((track.platform, track.artist, track.title), ("Spotify", "Artist", "Track"))
        detail = self.client.get(f"/studio/projects/{self.project.id}/")
        self.assertContains(detail, "Recommended tracks")
        self.assertContains(detail, "No description yet.")

    def test_project_sections_are_collapsible_and_can_be_hidden(self):
        self.client.force_login(self.owner)
        self.project.description = "Description"
        self.project.concept = "Concept"
        self.project.publication_info = "Publication details"
        self.project.save(update_fields=["description", "concept", "publication_info", "updated_at"])
        detail = self.client.get(f"/studio/projects/{self.project.id}/")
        content = detail.content.decode()
        self.assertContains(detail, 'data-project-section="service"')
        self.assertContains(detail, 'data-project-section="music"')
        self.assertContains(detail, 'data-project-section="legal"')
        self.assertContains(detail, 'data-project-section="characters" open')
        self.assertContains(detail, 'data-project-section="episodes" open')
        self.assertContains(detail, "Add and edit tracks in Edit Project")
        self.assertNotContains(detail, "Publication details")
        self.assertLess(content.index("Concept"), content.index("Project information"))
        self.assertLess(content.index("Recommended tracks"), content.index('data-media-library'))
        settings = self.client.get(f"/studio/projects/{self.project.id}/settings/")
        self.assertContains(settings, "Do not show these blocks")
        response = self.client.post(
            f"/studio/projects/{self.project.id}/settings/",
            {
                "original_language": self.project.original_language,
                "documentation_language": self.project.documentation_language,
                "dialogue_language": self.project.dialogue_language,
                "prompt_language": self.project.prompt_language,
                "translation_languages": "", "prompt_template": "",
                "hidden_sections": ["music", "legal"],
            },
        )
        self.assertRedirects(response, f"/studio/projects/{self.project.id}/settings/")
        self.project.refresh_from_db()
        self.assertEqual(self.project.hidden_sections, ["music", "legal"])
        hidden_detail = self.client.get(f"/studio/projects/{self.project.id}/")
        self.assertNotContains(hidden_detail, 'data-project-section="music"')
        self.assertNotContains(hidden_detail, 'data-project-section="legal"')

    def test_episode_supports_multiple_covers_and_one_avatar(self):
        import tempfile
        from pathlib import Path
        from .models import Asset
        from .storage import create_asset

        self.client.force_login(self.owner)
        with tempfile.TemporaryDirectory() as directory:
            with self.settings(STUDIO_PRIVATE_MEDIA_ROOT=Path(directory)):
                first = create_asset(user=self.owner, workspace=self.workspace, project=self.project, uploaded=self.image_file("episode-a.png"), kind=Asset.Kind.OTHER)
                second = create_asset(user=self.owner, workspace=self.workspace, project=self.project, uploaded=self.image_file("episode-b.png", "#934fc4"), kind=Asset.Kind.OTHER)
                for asset in (first, second):
                    response = self.client.post(f"/studio/assets/attach/episode/{self.episode.id}/", {"asset_id": asset.id})
                    self.assertEqual(response.status_code, 302)
                avatar = self.client.post(f"/studio/assets/attach/episode_avatar/{self.episode.id}/", {"asset_id": second.id})
                self.assertEqual(avatar.status_code, 302)
                self.episode.refresh_from_db()
                self.assertEqual(self.episode.cover_assets.count(), 2)
                self.assertEqual(self.episode.avatar_asset_id, second.id)
                page = self.client.get(f"/studio/projects/{self.project.id}/")
                self.assertContains(page, "episode-cover-strip")
                self.assertContains(page, "episode-b.png")
                self.assertContains(page, "episode-cover-add-tile")
                self.assertContains(page, "data-image-preview")
                self.assertContains(page, f'data-dialog-open="{self.episode.id}-cover-picker"')

    def test_episode_cover_metadata_can_be_updated_inline(self):
        import tempfile
        from pathlib import Path
        from .models import Asset, EpisodeCover
        from .storage import create_asset

        self.client.force_login(self.owner)
        with tempfile.TemporaryDirectory() as directory:
            with self.settings(STUDIO_PRIVATE_MEDIA_ROOT=Path(directory)):
                asset = create_asset(
                    user=self.owner, workspace=self.workspace, project=self.project,
                    uploaded=self.image_file("episode-social.png"), kind=Asset.Kind.OTHER,
                )
                self.client.post(f"/studio/assets/attach/episode/{self.episode.id}/", {"asset_id": asset.id})
                cover = EpisodeCover.objects.get(episode=self.episode, asset=asset)
                response = self.client.post(
                    f"/studio/episode-covers/{cover.id}/update/",
                    {"language_code": "PL", "platform": "OTHER", "custom_platform": "Vimeo"},
                    HTTP_X_REQUESTED_WITH="XMLHttpRequest",
                )
                self.assertEqual(response.status_code, 200)
                cover.refresh_from_db()
                self.assertEqual((cover.language_code, cover.platform, cover.custom_platform), ("PL", "OTHER", "Vimeo"))
                page = self.client.get(f"/studio/projects/{self.project.id}/")
                self.assertContains(page, "episode-cover-strip")
                self.assertNotContains(page, "data-episode-cover-form")
                editor = self.client.get(f"/studio/episodes/{self.episode.id}/edit/")
                self.assertContains(editor, "data-episode-cover-form")
                self.assertContains(editor, "Vimeo")
                self.assertContains(editor, 'data-live-preview-target="#episode-settings-cover-display"')

    def test_project_music_uses_refreshing_save_and_minus_controls(self):
        from .models import RecommendedTrack

        track = RecommendedTrack.objects.create(
            project=self.project, artist="Artist", title="Old track",
            created_by=self.owner, updated_by=self.owner,
        )
        self.client.force_login(self.owner)
        editor = self.client.get(f"/studio/projects/{self.project.id}/edit/")
        self.assertContains(editor, "data-preserve-position")
        self.assertContains(editor, 'class="panel project-edit-form" data-preserve-position data-stay-on-save')
        self.assertContains(editor, "data-remove-track")
        self.assertContains(editor, "Track removed - save to apply")
        response = self.client.post(
            f"/studio/projects/{self.project.id}/edit/",
            {
                "project_type": "SERIES", "title": self.project.title, "description": "",
                "concept": "", "rights_holder": "", "publication_info": "",
                "status": "DRAFT", "status_comment": "",
                "tracks-TOTAL_FORMS": "1", "tracks-INITIAL_FORMS": "1",
                "tracks-MIN_NUM_FORMS": "0", "tracks-MAX_NUM_FORMS": "1000",
                "tracks-0-id": str(track.id), "tracks-0-is_primary": "",
                "tracks-0-platform": "", "tracks-0-artist": "Artist",
                "tracks-0-title": "Old track", "tracks-0-url": "",
                "tracks-0-position": "0", "tracks-0-DELETE": "on",
            },
        )
        self.assertRedirects(response, f"/studio/projects/{self.project.id}/edit/")
        self.assertFalse(RecommendedTrack.objects.filter(id=track.id).exists())

    def test_language_names_only_appear_for_scene_title_and_dialogue_speaker(self):
        from .models import DialogueLine

        line = DialogueLine.objects.create(
            scene=self.scene, speaker="Hero", text="Hello",
            created_by=self.owner, updated_by=self.owner,
        )
        self.client.force_login(self.owner)
        scene_page = self.client.get(f"/studio/scenes/{self.scene.id}/")
        self.assertEqual(scene_page.content.decode().count("Documentation Language <b"), 1)
        self.assertContains(scene_page, f'/studio/dialogue/{line.id}/edit/')
        dialogue_page = self.client.get(f"/studio/dialogue/{line.id}/edit/")
        self.assertContains(dialogue_page, '<select name="speaker_documentation"', html=False)
        self.assertTrue(dialogue_page.context["form"].fields["speaker_prompt"].widget.attrs["readonly"])
        self.assertTrue(dialogue_page.context["form"].fields["speaker"].widget.attrs["readonly"])
        self.assertEqual(dialogue_page.content.decode().count("Documentation Language <b"), 1)
        self.assertContains(dialogue_page, "Status Comment")

    def test_project_language_change_requires_confirmation_and_propagates(self):
        from .models import AiModelProfile, DialogueLine, Prompt

        model = AiModelProfile.objects.create(
            name="Language image model", provider="test", model_id="language-image",
            media_type=AiModelProfile.MediaType.IMAGE,
        )
        original = Prompt.objects.create(
            scene=self.scene, ai_model=model, prompt_type=Prompt.Type.IMAGE,
            content="Original prompt", created_by=self.owner, updated_by=self.owner,
        )
        translated = Prompt.objects.create(
            scene=self.scene, ai_model=model, source_prompt=original, language="EN",
            prompt_type=Prompt.Type.IMAGE, position=1, content="Translated prompt",
            created_by=self.owner, updated_by=self.owner,
        )
        dialogue = DialogueLine.objects.create(
            scene=self.scene, speaker="Hero", text="Czesc", language="RU",
            created_by=self.owner, updated_by=self.owner,
        )
        payload = {
            "original_language": "PL", "documentation_language": "EN", "dialogue_language": "EN",
            "prompt_language": "EN", "translation_languages": "en", "prompt_template": "",
            "tracks-TOTAL_FORMS": "0", "tracks-INITIAL_FORMS": "0",
            "tracks-MIN_NUM_FORMS": "0", "tracks-MAX_NUM_FORMS": "1000",
        }
        self.client.force_login(self.owner)
        form_page = self.client.get(f"/studio/projects/{self.project.id}/settings/")
        self.assertContains(form_page, "data-language-propagation-dialog")
        self.assertNotContains(form_page, '<aside class="form-warning">')

        rejected = self.client.post(f"/studio/projects/{self.project.id}/settings/", payload)
        self.assertEqual(rejected.status_code, 200)
        self.assertContains(rejected, "Confirm that the new language")
        self.project.refresh_from_db()
        self.assertEqual(self.project.original_language, "ru")

        payload["confirm_language_propagation"] = "on"
        saved = self.client.post(f"/studio/projects/{self.project.id}/settings/", payload, follow=True)
        self.assertRedirects(saved, f"/studio/projects/{self.project.id}/settings/")
        self.assertContains(saved, "Updated 1 original prompts, 1 translations, and 1 dialogue lines")
        self.project.refresh_from_db()
        original.refresh_from_db()
        translated.refresh_from_db()
        dialogue.refresh_from_db()
        self.assertEqual(self.project.original_language, "PL")
        self.assertEqual((original.original_language, original.language), ("PL", "PL"))
        self.assertEqual((translated.original_language, translated.language), ("PL", "EN"))
        self.assertEqual(dialogue.language, "PL")

    def test_viewer_cannot_edit_project(self):
        self.client.force_login(self.viewer)
        response = self.client.get(f"/studio/projects/{self.project.id}/edit/")
        self.assertEqual(response.status_code, 403)

    def test_api_language_change_also_requires_propagation_confirmation(self):
        self.client.force_login(self.owner)
        denied = self.client.patch(
            f"/api/v1/studio/projects/{self.project.id}",
            data=json.dumps({"originalLanguage": "PL"}),
            content_type="application/json",
        )
        self.assertEqual(denied.status_code, 409, denied.content)
        self.assertEqual(denied.json()["error"]["code"], "language_propagation_confirmation_required")
        self.project.refresh_from_db()
        self.assertEqual(self.project.original_language, "ru")

        saved = self.client.patch(
            f"/api/v1/studio/projects/{self.project.id}",
            data=json.dumps({"originalLanguage": "PL", "confirmLanguagePropagation": True}),
            content_type="application/json",
        )
        self.assertEqual(saved.status_code, 200, saved.content)
        self.assertEqual(saved.json()["originalLanguage"], "PL")

    def test_owner_can_upload_private_scene_image_and_open_thumbnail(self):
        import tempfile
        from pathlib import Path
        from .models import Asset

        self.client.force_login(self.owner)
        with tempfile.TemporaryDirectory() as directory:
            with self.settings(STUDIO_PRIVATE_MEDIA_ROOT=Path(directory)):
                response = self.client.post(f"/studio/scenes/{self.scene.id}/images/new/", {"file": self.image_file()})
                asset = Asset.objects.get(scene=self.scene)
                self.assertRedirects(response, f"/studio/scenes/{self.scene.id}/")
                self.assertEqual(asset.kind, Asset.Kind.SCENE_IMAGE)
                self.assertEqual((asset.width, asset.height), (48, 36))
                thumbnail = self.client.get(f"/api/v1/studio/assets/{asset.id}/thumbnail")
                self.assertEqual(thumbnail.status_code, 200)

    def test_scene_accepts_multiple_images_and_rejects_workspace_duplicate(self):
        import tempfile
        from pathlib import Path
        from .models import Asset

        self.client.force_login(self.owner)
        with tempfile.TemporaryDirectory() as directory:
            with self.settings(STUDIO_PRIVATE_MEDIA_ROOT=Path(directory)):
                response = self.client.post(
                    f"/studio/scenes/{self.scene.id}/images/new/",
                    {"file": [self.image_file("one.png", "#113355"), self.image_file("two.png", "#557799")]},
                )
                self.assertEqual(response.status_code, 302)
                self.assertEqual(Asset.objects.filter(scene=self.scene).count(), 2)
                duplicate = self.client.post(
                    f"/studio/scenes/{self.scene.id}/images/new/",
                    {"file": self.image_file("one.png", "#113355")},
                    follow=True,
                )
                self.assertContains(duplicate, "already uploaded in this workspace")
                self.assertEqual(Asset.objects.filter(scene=self.scene).count(), 2)

    def test_outsider_cannot_open_scene_image_form(self):
        self.client.force_login(self.outsider)
        self.assertEqual(self.client.get(f"/studio/scenes/{self.scene.id}/images/new/").status_code, 404)

class StudioDocxImportTests(TestCase):
    def setUp(self):
        import tempfile
        from .models import private_storage

        self.temp_media = tempfile.TemporaryDirectory()
        self.private_storage = private_storage
        self.previous_location = private_storage._location
        private_storage._location = self.temp_media.name
        private_storage.__dict__.pop("base_location", None)
        private_storage.__dict__.pop("location", None)
        users = get_user_model()
        self.owner = users.objects.create_user("docx-owner", password="strong-pass")
        self.viewer = users.objects.create_user("docx-viewer", password="strong-pass")
        self.workspace = create_workspace(user=self.owner, name="Import Studio", slug="import-studio")
        WorkspaceMembership.objects.create(workspace=self.workspace, user=self.viewer, role=WorkspaceMembership.Role.VIEWER)

    def tearDown(self):
        self.private_storage._location = self.previous_location
        self.private_storage.__dict__.pop("base_location", None)
        self.private_storage.__dict__.pop("location", None)
        self.temp_media.cleanup()

    @staticmethod
    def docx_file():
        import io
        import zipfile
        from PIL import Image
        from django.core.files.uploadedfile import SimpleUploadedFile

        namespace = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        paragraphs = [
            ("", "DOM TEST"), ("Heading1", "СЕРИЯ 1: Pilot"), ("Heading3", "Краткое описание серии"), ("", "Summary"),
            ("Heading1", "Сцена 1. Opening"), ("Heading3", "Общее описание сцены"), ("", "A room"),
            ("Heading3", "Диалоги"), ("", 'Hero: “Hello!”'),
        ]
        body = []
        for style, text in paragraphs:
            ppr = f'<w:pPr><w:pStyle w:val="{style}"/></w:pPr>' if style else ""
            body.append(f'<w:p>{ppr}<w:r><w:t>{text}</w:t></w:r></w:p>')
        xml = f'<?xml version="1.0" encoding="UTF-8"?><w:document xmlns:w="{namespace}"><w:body>{"".join(body)}</w:body></w:document>'
        image = io.BytesIO()
        Image.new("RGB", (20, 12), "#227755").save(image, "PNG")
        output = io.BytesIO()
        with zipfile.ZipFile(output, "w") as archive:
            archive.writestr("[Content_Types].xml", "<Types/>")
            archive.writestr("word/document.xml", xml)
            archive.writestr("word/media/image1.png", image.getvalue())
        return SimpleUploadedFile("master.docx", output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")

    def test_owner_can_preview_and_accept_docx_with_embedded_image(self):
        from .models import Asset, DialogueLine, DocxImport, Episode, Scene

        self.client.force_login(self.owner)
        response = self.client.post(f"/studio/workspaces/{self.workspace.id}/imports/docx/new/", {"file": self.docx_file()})
        draft = DocxImport.objects.get()
        self.assertRedirects(response, f"/studio/imports/{draft.id}/")
        self.assertEqual(draft.parsed_data["image_count"], 1)
        self.assertEqual(len(draft.parsed_data["episodes"]), 1)
        accepted = self.client.post(f"/studio/imports/{draft.id}/accept/")
        draft.refresh_from_db()
        self.assertRedirects(accepted, f"/studio/projects/{draft.project_id}/")
        self.assertEqual(Episode.objects.filter(project=draft.project).count(), 1)
        self.assertEqual(Scene.objects.filter(episode__project=draft.project).count(), 1)
        self.assertEqual(DialogueLine.objects.filter(scene__episode__project=draft.project).count(), 1)
        self.assertEqual(Asset.objects.filter(project=draft.project, content_type="image/png").count(), 1)

    def test_workspace_document_can_be_archived_and_moved_to_recycle_bin(self):
        from .models import Asset, DocxImport

        self.client.force_login(self.owner)
        self.client.post(f"/studio/workspaces/{self.workspace.id}/imports/docx/new/", {"file": self.docx_file()})
        draft = DocxImport.objects.get()
        self.client.post(f"/studio/imports/{draft.id}/accept/")
        draft.refresh_from_db()

        workspace_page = self.client.get(f"/studio/workspaces/{self.workspace.id}/")
        self.assertContains(workspace_page, "master.docx")
        archived = self.client.post(f"/studio/imports/{draft.id}/archive/")
        self.assertRedirects(archived, f"/studio/workspaces/{self.workspace.id}/#documents")
        draft.refresh_from_db()
        self.assertIsNotNone(draft.archived_at)

        restored = self.client.post(
            f"/studio/imports/{draft.id}/archive/",
            {"action": "restore"},
        )
        self.assertRedirects(restored, f"/studio/workspaces/{self.workspace.id}/#documents")
        draft.refresh_from_db()
        self.assertIsNone(draft.archived_at)

        trashed = self.client.post(f"/studio/imports/{draft.id}/trash/")
        self.assertRedirects(trashed, f"/studio/workspaces/{self.workspace.id}/#documents")
        self.assertIsNotNone(Asset.all_objects.get(id=draft.source_asset_id).deleted_at)
        recycle = self.client.get(f"/studio/workspaces/{self.workspace.id}/recycle-bin/")
        self.assertContains(recycle, "master.docx")

    def test_import_edit_export_docx_round_trip_uses_current_project_data(self):
        from docx import Document
        from .docx_exports import generate_docx_export
        from .models import DocxImport

        self.client.force_login(self.owner)
        self.client.post(f"/studio/workspaces/{self.workspace.id}/imports/docx/new/", {"file": self.docx_file()})
        draft = DocxImport.objects.get()
        self.client.post(f"/studio/imports/{draft.id}/accept/")
        draft.refresh_from_db()
        scene = draft.project.episodes.first().scenes.first()
        scene.title = "Edited opening"
        scene.updated_by = self.owner
        scene.save()
        job = generate_docx_export(project=draft.project, user=self.owner)
        with job.output_asset.file.open("rb") as stream:
            document = Document(stream)
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("Edited opening", text)
        self.assertGreaterEqual(len(document.inline_shapes), 1)
        self.assertEqual(job.sections, ["DOCX"])
    def test_owner_can_edit_preview_before_accepting(self):
        from .models import AuditEvent, DocxImport

        self.client.force_login(self.owner)
        self.client.post(f"/studio/workspaces/{self.workspace.id}/imports/docx/new/", {"file": self.docx_file()})
        draft = DocxImport.objects.get()
        saved = self.client.post(
            f"/studio/imports/{draft.id}/",
            {"title": "Edited project", "episode_0_include": "on", "episode_0_title": "Edited pilot", "episode_0_scene_0_include": "on", "episode_0_scene_0_title": "Edited scene"},
        )
        self.assertRedirects(saved, f"/studio/imports/{draft.id}/")
        draft.refresh_from_db()
        self.assertEqual(draft.parsed_data["title"], "Edited project")
        self.assertEqual(draft.parsed_data["episodes"][0]["scenes"][0]["title"], "Edited scene")
        self.assertTrue(AuditEvent.objects.filter(action="DOCX_IMPORT_PREVIEW_EDITED").exists())
        self.client.post(f"/studio/imports/{draft.id}/accept/")
        draft.refresh_from_db()
        self.assertEqual(draft.project.title, "Edited project")
        self.assertEqual(draft.project.episodes.first().scenes.first().title, "Edited scene")
    def test_docx_export_redirects_to_roundtrip_report_before_download(self):
        import io
        from PIL import Image
        from django.core.files.uploadedfile import SimpleUploadedFile
        from docx import Document
        from docx.shared import Inches
        from .models import AuditEvent, DocxImport, ExportJob

        image = io.BytesIO()
        Image.new("RGB", (40, 30), "#557799").save(image, "PNG")
        image.seek(0)
        document = Document()
        document.add_heading("Round trip source", 0)
        document.add_heading("EPISODE 1: Pilot", 1)
        document.add_heading("\u0421\u0446\u0435\u043d\u0430 1. Opening", 1)
        document.add_heading("\u041e\u0431\u0449\u0435\u0435 \u043e\u043f\u0438\u0441\u0430\u043d\u0438\u0435 \u0441\u0446\u0435\u043d\u044b", 3)
        document.add_paragraph("A room with a window.")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Field"
        table.cell(0, 1).text = "Value"
        table.cell(1, 0).text = "Mood"
        table.cell(1, 1).text = "Quiet"
        document.add_picture(image, width=Inches(1))
        payload = io.BytesIO()
        document.save(payload)

        self.client.force_login(self.owner)
        self.client.post(
            f"/studio/workspaces/{self.workspace.id}/imports/docx/new/",
            {"file": SimpleUploadedFile(
                "round-trip.docx",
                payload.getvalue(),
                content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )},
        )
        draft = DocxImport.objects.get()
        self.client.post(f"/studio/imports/{draft.id}/accept/")
        draft.refresh_from_db()

        generated = self.client.post(
            f"/studio/projects/{draft.project_id}/exports/",
            {"format": "docx", "sections": ["metadata"]},
        )
        job = ExportJob.objects.get(project=draft.project, sections=["DOCX"])
        self.assertRedirects(generated, f"/studio/exports/{job.id}/round-trip/")
        report = self.client.get(f"/studio/exports/{job.id}/round-trip/")
        self.assertEqual(report.status_code, 200)
        self.assertContains(report, "DOCX round-trip report")
        self.assertContains(report, "Structure")
        self.assertContains(report, "Texts")
        self.assertContains(report, "Tables")
        self.assertContains(report, "Images")
        self.assertContains(report, "Section order")
        self.assertContains(report, f"/api/v1/studio/exports/{job.id}/download")
        self.assertTrue(AuditEvent.objects.filter(action="DOCX_ROUNDTRIP_REVIEWED").exists())

        history = self.client.get(f"/studio/projects/{draft.project_id}/exports/")
        self.assertContains(history, "Review &amp; download", html=True)
    def test_viewer_cannot_upload_docx(self):
        self.client.force_login(self.viewer)
        response = self.client.get(f"/studio/workspaces/{self.workspace.id}/imports/docx/new/")
        self.assertEqual(response.status_code, 403)

class StudioMasterDocumentTests(TestCase):
    def setUp(self):
        from .models import DialogueLine, Episode, Scene

        users = get_user_model()
        self.owner = users.objects.create_user("master-owner", password="strong-pass")
        self.viewer = users.objects.create_user("master-viewer", password="strong-pass")
        self.workspace = create_workspace(user=self.owner, name="Master Studio", slug="master-studio")
        WorkspaceMembership.objects.create(workspace=self.workspace, user=self.viewer, role=WorkspaceMembership.Role.VIEWER)
        self.project = Project.objects.create(
            workspace=self.workspace,
            project_type=Project.Type.SERIES,
            title="Master Project",
            concept="A complete document",
            original_language="ru",
            created_by=self.owner,
            updated_by=self.owner,
        )
        self.episode = Episode.objects.create(
            project=self.project, number=1, title="Pilot", position=0,
            created_by=self.owner, updated_by=self.owner,
        )
        self.first_scene = Scene.objects.create(
            episode=self.episode, number=1, title="First scene", position=0,
            created_by=self.owner, updated_by=self.owner,
        )
        self.second_scene = Scene.objects.create(
            episode=self.episode, number=2, title="Second scene", position=1,
            created_by=self.owner, updated_by=self.owner,
        )
        self.first_line = DialogueLine.objects.create(
            scene=self.first_scene, speaker="One", text="First line", position=0,
            created_by=self.owner, updated_by=self.owner,
        )
        self.second_line = DialogueLine.objects.create(
            scene=self.first_scene, speaker="Two", text="Second line", position=1,
            created_by=self.owner, updated_by=self.owner,
        )

    def test_member_can_open_master_document(self):
        self.client.force_login(self.viewer)
        response = self.client.get(f"/studio/projects/{self.project.id}/master/")
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Master Project")
        self.assertContains(response, "First scene")
        self.assertContains(response, "Second line")
        self.assertNotContains(response, f"/studio/move/scene/{self.first_scene.id}/up/")

    def test_owner_can_reorder_scenes_and_records_history(self):
        from .models import AuditEvent, Revision, Scene

        self.client.force_login(self.owner)
        response = self.client.post(f"/studio/move/scene/{self.second_scene.id}/up/")
        self.assertRedirects(
            response,
            f"/studio/projects/{self.project.id}/master/#item-{self.second_scene.id}",
        )
        self.assertEqual(
            list(Scene.objects.filter(episode=self.episode).order_by("position").values_list("id", flat=True)),
            [self.second_scene.id, self.first_scene.id],
        )
        self.assertTrue(Revision.objects.filter(entity_id=self.second_scene.id, operation="REORDER").exists())
        self.assertTrue(AuditEvent.objects.filter(entity_id=self.second_scene.id, action="ENTITY_REORDERED").exists())

    def test_dialogue_reorder_respects_unique_positions(self):
        from .models import DialogueLine

        self.client.force_login(self.owner)
        response = self.client.post(f"/studio/move/dialogue/{self.second_line.id}/up/")
        self.assertEqual(response.status_code, 302)
        self.assertEqual(
            list(DialogueLine.objects.filter(scene=self.first_scene).order_by("position").values_list("id", flat=True)),
            [self.second_line.id, self.first_line.id],
        )

    def test_dialogue_editor_saves_all_language_fields_without_server_error(self):
        self.client.force_login(self.owner)
        response = self.client.post(
            f"/studio/dialogue/{self.first_line.id}/edit/",
            {
                "speaker_documentation": "Narrator", "speaker_prompt": "Narrator EN", "speaker": "Narrator PL",
                "text_documentation": "Documentation copy", "text_prompt": "Prompt copy", "text": "Dialogue copy",
                "delivery_documentation": "Calm", "delivery_prompt": "Calmly", "delivery": "Spokojnie",
                "language": "PL", "status": "IN_REVIEW", "status_comment": "Translation checked",
            },
        )
        self.assertRedirects(
            response,
            f"/studio/scenes/{self.first_scene.id}/#dialogue-{self.first_line.id}",
        )
        self.first_line.refresh_from_db()
        self.assertEqual(self.first_line.text_prompt, "Prompt copy")
        self.assertEqual(self.first_line.text, "Dialogue copy")
        self.assertEqual(self.first_line.status_comment, "Translation checked")

        editor = self.client.get(f"/studio/dialogue/{self.first_line.id}/edit/")
        self.assertContains(editor, "Documentation Language")
        self.assertContains(editor, "Prompt Language")
        self.assertContains(editor, "Dialogue Language")

    def test_viewer_cannot_reorder(self):
        self.client.force_login(self.viewer)
        response = self.client.post(f"/studio/move/scene/{self.second_scene.id}/up/")
        self.assertEqual(response.status_code, 403)
        self.first_scene.refresh_from_db()
        self.second_scene.refresh_from_db()
        self.assertEqual((self.first_scene.position, self.second_scene.position), (0, 1))

class StudioImageGenerationWorkflowTests(TestCase):
    def setUp(self):
        from .models import Episode, Scene

        users = get_user_model()
        self.owner = users.objects.create_user("generation-owner", password="strong-pass")
        self.viewer = users.objects.create_user("generation-viewer", password="strong-pass")
        self.workspace = create_workspace(user=self.owner, name="Generation Studio", slug="generation-studio")
        WorkspaceMembership.objects.create(
            workspace=self.workspace, user=self.viewer, role=WorkspaceMembership.Role.VIEWER
        )
        self.project = Project.objects.create(
            workspace=self.workspace,
            project_type=Project.Type.SERIES,
            title="Visual Project",
            created_by=self.owner,
            updated_by=self.owner,
        )
        self.episode = Episode.objects.create(
            project=self.project, number=1, title="Pilot",
            created_by=self.owner, updated_by=self.owner,
        )
        self.scene = Scene.objects.create(
            episode=self.episode, number=1, title="Visual scene",
            created_by=self.owner, updated_by=self.owner,
        )

    @staticmethod
    def image_file(name="result.png"):
        import io
        from PIL import Image
        from django.core.files.uploadedfile import SimpleUploadedFile

        buffer = io.BytesIO()
        Image.new("RGB", (64, 48), "#54a778").save(buffer, "PNG")
        return SimpleUploadedFile(name, buffer.getvalue(), content_type="image/png")

    def test_owner_can_create_generation_upload_result_and_select_final(self):
        import tempfile
        from pathlib import Path
        from .models import AdditionalGeneration, AuditEvent, GenerationOutput, Revision

        self.client.force_login(self.owner)
        created = self.client.post(
            f"/studio/scenes/{self.scene.id}/generations/new/",
            {"reason": "Closer composition", "prompt": "A close-up with softer light", "status": "DRAFT"},
        )
        generation = AdditionalGeneration.objects.exclude(scene=self.scene).get()
        generated_scene = generation.scene
        self.assertEqual(generated_scene.title, "Additional generation 1")
        self.assertEqual(generated_scene.scene_type, "ADDITIONAL_GENERATION")
        self.assertGreater(generated_scene.position, self.scene.position)
        self.assertRedirects(created, f"/studio/scenes/{generated_scene.id}/media/")
        self.assertTrue(Revision.objects.filter(entity_id=generation.id, operation="CREATE").exists())

        with tempfile.TemporaryDirectory() as directory:
            with self.settings(STUDIO_PRIVATE_MEDIA_ROOT=Path(directory)):
                uploaded = self.client.post(
                    f"/studio/generations/{generation.id}/outputs/new/",
                    {"file": self.image_file(), "model_name": "Test image model"},
                )
                output = GenerationOutput.objects.get(generation=generation)
                self.assertRedirects(uploaded, f"/studio/scenes/{generated_scene.id}/media/")
                self.assertEqual(output.model_metadata["model"], "Test image model")
                edited = self.client.post(
                    f"/studio/assets/{output.asset_id}/edit/",
                    {"original_filename": "approved-result.png", "kind": "GENERATION_OUTPUT"},
                )
                self.assertRedirects(edited, f"/studio/scenes/{generated_scene.id}/media/")
                output.asset.refresh_from_db()
                self.assertEqual(output.asset.original_filename, "approved-result.png")
                selected = self.client.post(f"/studio/generation-outputs/{output.id}/final/")
                self.assertRedirects(
                    selected,
                    f"/studio/scenes/{generated_scene.id}/media/#generation-{generation.id}",
                )

        output.refresh_from_db()
        generation.refresh_from_db()
        self.assertTrue(output.is_final)
        self.assertEqual(generation.status, AdditionalGeneration.Status.FINAL)
        self.assertTrue(Revision.objects.filter(entity_id=generation.id, operation="FINAL_OUTPUT").exists())
        self.assertTrue(AuditEvent.objects.filter(action="GENERATION_OUTPUT_SELECTED").exists())

    def test_viewer_can_browse_but_cannot_create_generation(self):
        self.client.force_login(self.viewer)
        page = self.client.get(f"/studio/scenes/{self.scene.id}/media/")
        self.assertEqual(page.status_code, 200)
        self.assertContains(page, "Images &amp; additional generations", html=True)
        self.assertNotContains(page, "New generation")
        denied = self.client.post(
            f"/studio/scenes/{self.scene.id}/generations/new/",
            {"reason": "Change", "prompt": "Change it", "status": "DRAFT"},
        )
        self.assertEqual(denied.status_code, 403)

class StudioInlineEditingWorkflowTests(TestCase):
    def setUp(self):
        from .models import AiModelProfile, Episode, Scene

        users = get_user_model()
        self.owner = users.objects.create_user("inline-owner", password="strong-pass")
        self.viewer = users.objects.create_user("inline-viewer", password="strong-pass")
        self.workspace = create_workspace(user=self.owner, name="Inline Studio", slug="inline-studio")
        WorkspaceMembership.objects.create(
            workspace=self.workspace, user=self.viewer, role=WorkspaceMembership.Role.VIEWER
        )
        self.project = Project.objects.create(
            workspace=self.workspace, project_type=Project.Type.SERIES, title="Inline Project",
            created_by=self.owner, updated_by=self.owner,
        )
        self.episode = Episode.objects.create(
            project=self.project, number=1, title="Pilot",
            created_by=self.owner, updated_by=self.owner,
        )
        self.first_scene = Scene.objects.create(
            episode=self.episode, number=1, title="First", position=0,
            created_by=self.owner, updated_by=self.owner,
        )
        self.second_scene = Scene.objects.create(
            episode=self.episode, number=2, title="Second", position=1,
            created_by=self.owner, updated_by=self.owner,
        )
        self.ai_model = AiModelProfile.objects.create(
            name="Inline Image Model", provider="test", model_id="inline-image",
            media_type=AiModelProfile.MediaType.IMAGE,
        )

    @staticmethod
    def image_file(name="prompt.png"):
        import io
        from PIL import Image
        from django.core.files.uploadedfile import SimpleUploadedFile

        output = io.BytesIO()
        Image.new("RGB", (72, 54), "#286c91").save(output, "PNG")
        return SimpleUploadedFile(name, output.getvalue(), content_type="image/png")

    def test_inline_prompt_create_edit_and_image_upload(self):
        import tempfile
        from pathlib import Path
        from .models import Asset, AuditEvent, Prompt, PromptBlock

        self.client.force_login(self.owner)
        self.project.prompt_template = "Camera: locked. No music."
        self.project.save(update_fields=["prompt_template", "updated_at"])
        created = self.client.post(
            f"/studio/scenes/{self.first_scene.id}/prompts/quick-create/",
            {
                "title": "Opening frame",
                "ai_model": str(self.ai_model.id),
                "prompt_type": "IMAGE",
                "status": "DRAFT",
                "content": "Wide establishing shot",
            },
        )
        prompt = Prompt.objects.get(scene=self.first_scene)
        block = prompt.blocks.get()
        self.assertRedirects(created, f"/studio/scenes/{self.first_scene.id}/#prompt-{prompt.id}")
        self.assertEqual(prompt.content, "Wide establishing shot\n\nCamera: locked. No music.")
        self.assertEqual(block.content, prompt.content)
        prompt_page = self.client.get(f"/studio/scenes/{self.first_scene.id}/")
        self.assertContains(prompt_page, "Add Template")

        saved = self.client.post(
            f"/studio/prompts/{prompt.id}/quick-save/",
            {
                "title": "Opening frame revised",
                "ai_model": str(self.ai_model.id),
                "prompt_type": "IMAGE",
                "status": "IN_REVIEW",
                "content": "Closer establishing shot\n\nNo text overlays",
            },
        )
        self.assertRedirects(saved, f"/studio/scenes/{self.first_scene.id}/#prompt-{prompt.id}")
        prompt.refresh_from_db()
        block.refresh_from_db()
        self.assertEqual(prompt.title, "Opening frame revised")
        self.assertEqual(prompt.status, "IN_REVIEW")
        self.assertEqual(prompt.content, "Closer establishing shot\n\nNo text overlays")
        self.assertEqual(block.content, "Closer establishing shot\n\nNo text overlays")
        self.assertEqual(PromptBlock.objects.filter(prompt=prompt).count(), 1)

        ajax_saved = self.client.post(
            f"/studio/prompts/{prompt.id}/quick-save/",
            {
                "title": "Opening frame revised",
                "ai_model": str(self.ai_model.id),
                "prompt_type": "IMAGE",
                "status": "IN_REVIEW",
                "content": "Saved without reloading the scene",
            },
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )
        self.assertEqual(ajax_saved.status_code, 200, ajax_saved.content)
        self.assertEqual(ajax_saved.json()["content"], "Saved without reloading the scene")
        self.assertEqual(ajax_saved.json()["id"], str(prompt.id))
        self.assertIn("previewUrl", ajax_saved.json())

        with tempfile.TemporaryDirectory() as directory:
            with self.settings(STUDIO_PRIVATE_MEDIA_ROOT=Path(directory)):
                uploaded = self.client.post(
                    f"/studio/prompts/{prompt.id}/images/new/",
                    {"file": self.image_file()},
                )
                self.assertRedirects(uploaded, f"/studio/scenes/{self.first_scene.id}/#prompt-{prompt.id}")
                asset = Asset.objects.get(prompt=prompt)
                self.assertEqual(asset.scene, self.first_scene)
                page = self.client.get(f"/studio/scenes/{self.first_scene.id}/")
                self.assertContains(page, asset.original_filename)
                self.assertContains(page, "data-image-modal")
        self.assertTrue(AuditEvent.objects.filter(action="PROMPT_INLINE_UPDATED").exists())

    def test_original_prompt_language_is_editable_and_inherited_by_translations(self):
        from .models import Prompt

        self.client.force_login(self.owner)
        original = Prompt.objects.create(
            scene=self.first_scene, ai_model=self.ai_model, prompt_type=Prompt.Type.IMAGE,
            content="Source", created_by=self.owner, updated_by=self.owner,
        )
        translated = Prompt.objects.create(
            scene=self.first_scene, ai_model=self.ai_model, source_prompt=original,
            language="DE", prompt_type=Prompt.Type.IMAGE, position=1, content="Ubersetzung",
            created_by=self.owner, updated_by=self.owner,
        )
        page = self.client.get(f"/studio/scenes/{self.first_scene.id}/")
        self.assertContains(page, "Original &middot; EN")
        self.assertContains(page, "Translation &middot; DE")
        self.assertNotContains(page, "<label>Original language")
        response = self.client.post(
            f"/studio/prompts/{original.id}/quick-save/",
            {
                "title": "Source prompt", "ai_model": str(self.ai_model.id),
                "original_language": "PL", "prompt_type": "IMAGE", "status": "DRAFT",
                "content": "Source",
            },
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(response.json()["originalLanguage"], "PL")
        self.assertEqual({row["label"] for row in response.json()["versions"]}, {"Original · PL", "Translation · DE"})
        original.refresh_from_db()
        translated.refresh_from_db()
        self.assertEqual((original.original_language, original.language), ("PL", "PL"))
        self.assertEqual((translated.original_language, translated.language), ("PL", "DE"))

    def test_scene_chain_quick_edit_and_drag_reorder_endpoint(self):
        from .models import AuditEvent, Scene

        self.client.force_login(self.owner)
        page = self.client.get(f"/studio/projects/{self.project.id}/scene-chain/")
        self.assertEqual(page.status_code, 200)
        self.assertContains(page, "Scene chain")
        self.assertContains(page, "data-scene-list")
        self.assertContains(page, "drag-handle")
        self.assertContains(page, "data-theme-toggle")
        self.assertContains(page, 'data-theme="business"', html=False)
        self.assertContains(page, 'themeOrder=["dark","light","business"]', html=False)
        self.assertContains(page, "studio/business_icons.css", html=False)
        self.assertContains(page, "studio/business_icons.js", html=False)
        self.assertContains(page, "timeline-track")
        self.assertContains(page, "scene-editor-deck")
        self.assertContains(page, "data-scroll-direction")
        self.assertContains(page, 'class="scene-editor-panel"')
        self.assertContains(page, "Open prompts and images")
        self.assertNotContains(page, 'class="prompt-editor"')
        self.assertNotContains(page, "asset-picker image-picker-dialog")

        saved = self.client.post(
            f"/studio/scenes/{self.first_scene.id}/quick-save/",
            {
                "number": 1,
                "title": "First revised",
                "hook": "Hook",
                "description": "Updated in chain",
                "location": "Room",
                "actions": "Walks",
                "performance_notes": "Quiet",
                "scene_type": "ORIGINAL",
                "status": "IN_REVIEW",
            },
        )
        self.assertRedirects(saved, f"/studio/projects/{self.project.id}/scene-chain/#scene-{self.first_scene.id}")
        self.first_scene.refresh_from_db()
        self.assertEqual(self.first_scene.title, "First revised")

        reordered = self.client.post(
            f"/studio/episodes/{self.episode.id}/scenes/reorder/",
            data=json.dumps({"sceneIds": [str(self.second_scene.id), str(self.first_scene.id)]}),
            content_type="application/json",
        )
        self.assertEqual(reordered.status_code, 200)
        self.assertEqual(
            list(Scene.objects.filter(episode=self.episode).order_by("position").values_list("id", flat=True)),
            [self.second_scene.id, self.first_scene.id],
        )
        self.assertTrue(AuditEvent.objects.filter(action="SCENES_DRAG_REORDERED").exists())

    def test_scene_editor_updates_type_and_production_status(self):
        self.client.force_login(self.owner)
        response = self.client.post(
            f"/studio/scenes/{self.first_scene.id}/",
            {
                "title": self.first_scene.title,
                "hook": "",
                "description": "Alternative production cut",
                "location": "",
                "actions": "",
                "performance_notes": "",
                "scene_type": "ALTERNATIVE",
                "status": "PRODUCTION",
            },
        )
        self.assertRedirects(response, f"/studio/scenes/{self.first_scene.id}/?focus=top#scene-navigation-top")
        self.first_scene.refresh_from_db()
        self.assertEqual(self.first_scene.scene_type, "ALTERNATIVE")
        self.assertEqual(self.first_scene.status, "PRODUCTION")

    def test_scene_editor_saves_three_language_fields_and_status_comment(self):
        self.client.force_login(self.owner)
        response = self.client.post(
            f"/studio/scenes/{self.first_scene.id}/",
            {
                "title": "Opis sceny", "title_prompt": "Scene description", "title_dialogue": "Opis dialogu",
                "hook": "Dokumentacja", "hook_prompt": "Prompt copy", "hook_dialogue": "Dialog copy",
                "description": "Główny opis", "description_prompt": "Main description",
                "description_dialogue": "Opis docelowy", "location": "Biuro", "location_prompt": "Office",
                "location_dialogue": "Biuro dialog", "actions": "Idzie", "actions_prompt": "Walks",
                "actions_dialogue": "Idzie dialog", "performance_notes": "Spokojnie",
                "performance_notes_prompt": "Calmly", "performance_notes_dialogue": "Spokojnie dialog",
                "scene_type": "ORIGINAL", "status": "PRODUCTION", "status_comment": "Ready for shooting",
            },
        )
        self.assertEqual(response.status_code, 302)
        self.first_scene.refresh_from_db()
        self.assertEqual(self.first_scene.title_prompt, "Scene description")
        self.assertEqual(self.first_scene.description_dialogue, "Opis docelowy")
        self.assertEqual(self.first_scene.status_comment, "Ready for shooting")

    def test_new_character_can_include_multilingual_fields_and_avatar(self):
        import tempfile
        from pathlib import Path
        from .models import Character

        self.client.force_login(self.owner)
        with tempfile.TemporaryDirectory() as directory:
            with self.settings(STUDIO_PRIVATE_MEDIA_ROOT=Path(directory)):
                response = self.client.post(
                    f"/studio/projects/{self.project.id}/characters/new/",
                    {
                        "name": "Kot", "name_prompt": "Cat", "name_dialogue": "Gato",
                        "description": "Opis", "description_prompt": "Description",
                        "description_dialogue": "Descripcion", "visual_description": "Rudy",
                        "visual_description_prompt": "Ginger", "visual_description_dialogue": "Naranja",
                        "avatar_file": self.image_file("cat-avatar.png"),
                    },
                )
                self.assertEqual(response.status_code, 302)
                character = Character.objects.get(project=self.project, name="Kot")
                self.assertEqual(character.name_prompt, "Cat")
                self.assertEqual(character.name_dialogue, "Gato")
                self.assertIsNotNone(character.avatar_asset_id)
                self.assertTrue(character.reference_assets.filter(id=character.avatar_asset_id).exists())

    def test_scene_copy_appends_deep_copy_and_delete_renumbers(self):
        from .models import DialogueLine, Prompt, Scene

        DialogueLine.objects.create(
            scene=self.first_scene, speaker="Cat", text="Hello", position=0,
            created_by=self.owner, updated_by=self.owner,
        )
        Prompt.objects.create(
            scene=self.first_scene, ai_model=self.ai_model, prompt_type="IMAGE",
            title="Frame", content="A ginger cat", position=0, needs_review=True,
            created_by=self.owner, updated_by=self.owner,
        )
        self.first_scene.scene_type = Scene.Type.ALTERNATIVE
        self.first_scene.save(update_fields=["scene_type", "updated_at"])
        self.client.force_login(self.owner)

        copied_response = self.client.post(f"/studio/scenes/{self.first_scene.id}/copy/")
        copied = Scene.objects.get(episode=self.episode, number=3)
        self.assertRedirects(copied_response, f"/studio/scenes/{copied.id}/")
        self.assertEqual((copied.position, copied.scene_type, copied.status), (2, "ALTERNATIVE", "DRAFT"))
        self.assertEqual(copied.dialogue_lines.get().text, "Hello")
        self.assertTrue(copied.prompts.get().needs_review)

        deleted_response = self.client.post(f"/studio/scenes/{self.second_scene.id}/delete/")
        self.assertRedirects(deleted_response, f"/studio/scenes/{copied.id}/")
        ordered = list(Scene.objects.filter(episode=self.episode).order_by("position"))
        self.assertEqual([(scene.id, scene.number, scene.position) for scene in ordered], [
            (self.first_scene.id, 1, 0), (copied.id, 2, 1),
        ])

    def test_scene_pages_show_full_path_actions_and_ordered_project_navigation(self):
        self.client.force_login(self.owner)
        page = self.client.get(f"/studio/scenes/{self.first_scene.id}/")
        self.assertContains(page, 'aria-label="Scene path"')
        self.assertContains(page, "Lexamora Studio")
        self.assertContains(page, f"Project {self.project.title}")
        self.assertContains(page, f'/studio/scenes/{self.first_scene.id}/copy/')
        self.assertContains(page, f'/studio/scenes/{self.first_scene.id}/delete/')
        content = page.content.decode()
        self.assertLess(content.index("General View"), content.index("Master Document"))
        self.assertLess(content.index("Master Document"), content.index("Scene Chain"))
        self.assertLess(content.index("Scene Chain"), content.index("Translations"))

        chain = self.client.get(f"/studio/projects/{self.project.id}/scene-chain/")
        self.assertContains(chain, 'name="scene_type"')
        self.assertContains(chain, "Production")
        self.assertContains(chain, f'/studio/scenes/{self.first_scene.id}/copy/')
        self.assertContains(chain, f'/studio/scenes/{self.first_scene.id}/delete/')

    def test_asset_api_can_attach_image_to_prompt(self):
        import tempfile
        from pathlib import Path
        from .models import Asset, Prompt

        prompt = Prompt.objects.create(
            scene=self.first_scene,
            ai_model=self.ai_model,
            prompt_type="IMAGE",
            title="API prompt",
            created_by=self.owner,
            updated_by=self.owner,
        )
        self.client.force_login(self.owner)
        with tempfile.TemporaryDirectory() as directory:
            with self.settings(STUDIO_PRIVATE_MEDIA_ROOT=Path(directory)):
                response = self.client.post(
                    "/api/v1/studio/assets",
                    {
                        "workspaceId": str(self.workspace.id),
                        "promptId": str(prompt.id),
                        "kind": "OTHER",
                        "file": self.image_file("api-prompt.png"),
                    },
                )
        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.json()["promptId"], str(prompt.id))
        self.assertTrue(Asset.objects.filter(prompt=prompt, scene=self.first_scene).exists())

    def test_scene_detail_is_inline_editor_with_two_way_navigation(self):
        self.client.force_login(self.owner)
        page = self.client.get(f"/studio/scenes/{self.first_scene.id}/")
        self.assertEqual(page.status_code, 200)
        self.assertContains(page, 'id="scene-editor-form"')
        self.assertContains(page, "Create Scene", count=1)
        self.assertContains(page, f'/studio/scenes/{self.second_scene.id}/')
        self.assertContains(page, "data-scene-cancel disabled", count=1)
        self.assertNotContains(page, 'name="number"')
        self.assertContains(page, 'id="scene-navigation-bottom"')
        self.assertContains(page, "+ Add model...")
        self.assertContains(page, "scene-editor-command-bar")
        self.assertContains(page, "scene-identity-meta")
        self.assertContains(page, "Created")
        self.assertNotContains(page, '<div class="project-shell-identity">')

    def test_workspace_avatar_and_manual_generation_model(self):
        import tempfile
        from pathlib import Path
        from .models import AiModelProfile

        self.client.force_login(self.owner)
        form_page = self.client.get(f"/studio/workspaces/{self.workspace.id}/avatar/")
        self.assertContains(form_page, 'enctype="multipart/form-data"')
        with tempfile.TemporaryDirectory() as directory:
            with self.settings(STUDIO_PRIVATE_MEDIA_ROOT=Path(directory)):
                uploaded = self.client.post(f"/studio/workspaces/{self.workspace.id}/avatar/", {"file": self.image_file("workspace.png")})
                self.assertRedirects(uploaded, f"/studio/workspaces/{self.workspace.id}/")
                self.workspace.refresh_from_db()
                self.assertIsNotNone(self.workspace.avatar_asset_id)
        created = self.client.post(
            f"/studio/workspaces/{self.workspace.id}/models/",
            {"name": "Manual Motion", "provider": "Custom", "model_id": "custom/motion-v1", "media_type": "VIDEO"},
        )
        self.assertRedirects(created, f"/studio/workspaces/{self.workspace.id}/models/")
        self.assertTrue(AiModelProfile.objects.filter(name="Manual Motion", model_id="custom/motion-v1").exists())

    def test_workspace_description_saves_revision_without_server_error(self):
        from .models import Revision

        self.client.force_login(self.owner)
        response = self.client.post(
            f"/studio/workspaces/{self.workspace.id}/edit/",
            {"name": self.workspace.name, "description": "A fuller production workspace description."},
        )
        self.assertRedirects(response, f"/studio/workspaces/{self.workspace.id}/edit/")
        self.workspace.refresh_from_db()
        self.assertEqual(self.workspace.description, "A fuller production workspace description.")
        self.assertTrue(Revision.objects.filter(entity_type="lexamora_studio.workspace", entity_id=self.workspace.id).exists())

    def test_owner_is_visible_and_cannot_be_excluded_from_project(self):
        self.client.force_login(self.owner)
        page = self.client.get(f"/studio/projects/{self.project.id}/access/")
        self.assertContains(page, "Workspace owner / full access")
        owner_membership = WorkspaceMembership.objects.get(workspace=self.workspace, user=self.owner)
        denied = self.client.post(
            f"/studio/projects/{self.project.id}/access/",
            {"action": "exclude", "user_id": str(owner_membership.user_id)},
        )
        self.assertEqual(denied.status_code, 403)
        self.assertTrue(accessible_projects(self.owner).filter(id=self.project.id).exists())

    def test_workspace_admin_cannot_remove_access(self):
        users = get_user_model()
        admin = users.objects.create_user("workspace-admin", password="strong-pass")
        WorkspaceMembership.objects.create(
            workspace=self.workspace, user=admin, role=WorkspaceMembership.Role.ADMIN,
        )
        viewer_membership = WorkspaceMembership.objects.get(workspace=self.workspace, user=self.viewer)
        self.client.force_login(admin)
        workspace_response = self.client.post(
            f"/studio/workspaces/{self.workspace.id}/access/",
            {"action": "remove", "membership_id": str(viewer_membership.id)},
        )
        project_response = self.client.post(
            f"/studio/projects/{self.project.id}/access/",
            {"action": "exclude", "user_id": str(self.viewer.id)},
        )
        self.assertEqual(workspace_response.status_code, 403)
        self.assertEqual(project_response.status_code, 403)

    def test_only_workspace_owner_or_admin_can_copy_or_archive(self):
        users = get_user_model()
        editor = users.objects.create_user("archive-editor", email="archive-editor@example.com", password="strong-pass")
        WorkspaceMembership.objects.create(workspace=self.workspace, user=editor, role=WorkspaceMembership.Role.EDITOR)
        self.client.force_login(editor)
        self.assertEqual(self.client.post(f"/studio/workspaces/{self.workspace.id}/copy/").status_code, 403)
        self.assertEqual(self.client.post(f"/studio/workspaces/{self.workspace.id}/trash/").status_code, 403)
        self.assertEqual(self.client.post(f"/studio/projects/{self.project.id}/copy/").status_code, 403)
        self.assertEqual(self.client.post(f"/studio/projects/{self.project.id}/trash/").status_code, 403)

    def test_archive_has_no_auto_expiry_and_purge_requires_30_seconds(self):
        from django.core import signing
        from django.utils import timezone
        from .models import Project

        self.client.force_login(self.owner)
        archived = self.client.post(f"/studio/projects/{self.project.id}/trash/")
        self.assertRedirects(archived, f"/studio/workspaces/{self.workspace.id}/")
        Project.all_objects.filter(id=self.project.id).update(deleted_at=timezone.now() - timedelta(days=90))
        archive_page = self.client.get(f"/studio/workspaces/{self.workspace.id}/projects/trash/")
        self.assertContains(archive_page, "Project archive")
        self.assertContains(archive_page, "data-purge-delay=\"30\"")
        self.assertIsNone(Project.all_objects.get(id=self.project.id).purged_at)
        immediate_token = archive_page.context["projects"][0].purge_token
        self.assertEqual(self.client.post(f"/studio/projects/{self.project.id}/purge/", {"purge_token": immediate_token}).status_code, 403)
        ready_token = signing.dumps(
            {"kind": "project", "id": str(self.project.id), "user": str(self.owner.id), "issued": timezone.now().timestamp() - 31},
            salt="studio-archive-purge",
        )
        purged = self.client.post(f"/studio/projects/{self.project.id}/purge/", {"purge_token": ready_token})
        self.assertRedirects(purged, f"/studio/workspaces/{self.workspace.id}/projects/trash/")
        self.assertIsNotNone(Project.all_objects.get(id=self.project.id).purged_at)

    def test_workspace_archive_purge_uses_the_same_delay(self):
        from django.core import signing
        from django.utils import timezone
        from .models import Workspace

        self.client.force_login(self.owner)
        archived = self.client.post(f"/studio/workspaces/{self.workspace.id}/trash/")
        self.assertRedirects(archived, "/studio/")
        archive_page = self.client.get("/studio/")
        self.assertContains(archive_page, "Workspace archive")
        immediate_token = archive_page.context["archived_workspaces"][0].purge_token
        self.assertEqual(self.client.post(f"/studio/workspaces/{self.workspace.id}/purge/", {"purge_token": immediate_token}).status_code, 403)
        ready_token = signing.dumps(
            {"kind": "workspace", "id": str(self.workspace.id), "user": str(self.owner.id), "issued": timezone.now().timestamp() - 31},
            salt="studio-archive-purge",
        )
        purged = self.client.post(f"/studio/workspaces/{self.workspace.id}/purge/", {"purge_token": ready_token})
        self.assertRedirects(purged, "/studio/")
        self.assertIsNotNone(Workspace.all_objects.get(id=self.workspace.id).purged_at)

    def test_new_workspace_and_project_shares_send_email(self):
        from django.core import mail

        users = get_user_model()
        workspace_guest = users.objects.create_user("workspace-guest", email="workspace-guest@example.com", password="strong-pass")
        project_guest = users.objects.create_user("project-guest", email="project-guest@example.com", password="strong-pass")
        self.client.force_login(self.owner)
        with self.settings(EMAIL_BACKEND="django.core.mail.backends.locmem.EmailBackend"):
            with self.captureOnCommitCallbacks(execute=True):
                workspace_response = self.client.post(
                    f"/studio/workspaces/{self.workspace.id}/access/",
                    {"email": workspace_guest.email, "role": WorkspaceMembership.Role.VIEWER, "can_use_ai": "on"},
                )
            with self.captureOnCommitCallbacks(execute=True):
                project_response = self.client.post(
                    f"/studio/projects/{self.project.id}/access/",
                    {"email": project_guest.email, "role": "VIEWER"},
                )
        self.assertRedirects(workspace_response, f"/studio/workspaces/{self.workspace.id}/access/")
        self.assertRedirects(project_response, f"/studio/projects/{self.project.id}/access/")
        self.assertEqual([message.to for message in mail.outbox[-2:]], [[workspace_guest.email], [project_guest.email]])

    def test_translations_show_prompt_versions_as_separate_cards_and_filters(self):
        translated = Prompt.objects.create(
            scene=self.first_scene, ai_model=self.ai_model, source_prompt=None,
            prompt_type="IMAGE", title="Original prompt", language="EN", original_language="EN",
            content="Original", created_by=self.owner, updated_by=self.owner,
        )
        Prompt.objects.create(
            scene=self.first_scene, ai_model=self.ai_model, source_prompt=translated,
            prompt_type="IMAGE", title="Polish prompt", language="PL", original_language="EN",
            content="Polski tekst", position=1, created_by=self.owner, updated_by=self.owner,
        )
        self.client.force_login(self.owner)
        page = self.client.get(f"/studio/projects/{self.project.id}/translations/?scene={self.first_scene.id}&prompt_language=PL&sort=desc")
        self.assertEqual(page.status_code, 200)
        self.assertContains(page, "Original prompt")
        self.assertContains(page, "Polski tekst")
        self.assertContains(page, "Saved translation")

    def test_scene_reorder_recomputes_display_numbers(self):
        from .models import Scene

        self.client.force_login(self.owner)
        response = self.client.post(
            f"/studio/episodes/{self.episode.id}/scenes/reorder/",
            data=json.dumps({"sceneIds": [str(self.second_scene.id), str(self.first_scene.id)]}),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 200)
        ordered = list(Scene.objects.filter(episode=self.episode).order_by("position"))
        self.assertEqual([(item.id, item.number) for item in ordered], [(self.second_scene.id, 1), (self.first_scene.id, 2)])

    def test_crop_creates_second_project_image_and_detach_keeps_file(self):
        import tempfile
        from pathlib import Path
        from .models import Asset
        from .storage import create_asset

        self.client.force_login(self.owner)
        with tempfile.TemporaryDirectory() as directory:
            with self.settings(STUDIO_PRIVATE_MEDIA_ROOT=Path(directory)):
                source = create_asset(
                    user=self.owner, workspace=self.workspace, project=self.project,
                    scene=self.first_scene, uploaded=self.image_file("source.png"), kind=Asset.Kind.SCENE_IMAGE,
                )
                cropped = self.client.post(
                    f"/studio/assets/{source.id}/crop/",
                    data=json.dumps({"x": 4, "y": 3, "width": 40, "height": 30}),
                    content_type="application/json",
                )
                self.assertEqual(cropped.status_code, 201, cropped.content)
                crop = Asset.objects.get(id=cropped.json()["id"])
                self.assertEqual((crop.width, crop.height), (40, 30))
                self.assertEqual(crop.project, self.project)
                self.assertIsNone(crop.scene)
                self.assertTrue(Asset.objects.filter(id=source.id).exists())

                detached = self.client.post(
                    f"/studio/assets/{source.id}/detach/scene/{self.first_scene.id}/",
                    {"next": f"/studio/scenes/{self.first_scene.id}/#images"},
                )
                self.assertRedirects(detached, f"/studio/scenes/{self.first_scene.id}/#images")
                source.refresh_from_db()
                self.assertIsNone(source.scene)
                self.assertTrue(Asset.objects.filter(id=source.id, deleted_at__isnull=True).exists())

    def test_workspace_media_library_uploads_multiple_images(self):
        import tempfile
        from pathlib import Path
        from .models import Asset

        self.client.force_login(self.owner)
        with tempfile.TemporaryDirectory() as directory:
            with self.settings(STUDIO_PRIVATE_MEDIA_ROOT=Path(directory)):
                response = self.client.post(
                    f"/studio/workspaces/{self.workspace.id}/images/new/",
                    {
                        "file": [self.image_file("workspace-a.png"), self.image_file("workspace-b.png")],
                        "next": f"/studio/workspaces/{self.workspace.id}/#workspace-images",
                    },
                )
                self.assertRedirects(response, f"/studio/workspaces/{self.workspace.id}/#workspace-images")
                self.assertEqual(Asset.objects.filter(workspace=self.workspace, project__isnull=True).count(), 2)
                page = self.client.get(f"/studio/workspaces/{self.workspace.id}/")
                self.assertContains(page, "Workspace Images")
                self.assertContains(page, "workspace-a.png")
                self.assertContains(page, "data-media-library")

    def test_workspace_image_is_available_in_project_picker_and_can_be_cropped(self):
        import tempfile
        from pathlib import Path
        from .models import Asset
        from .storage import create_asset

        self.client.force_login(self.owner)
        with tempfile.TemporaryDirectory() as directory:
            with self.settings(STUDIO_PRIVATE_MEDIA_ROOT=Path(directory)):
                source = create_asset(
                    user=self.owner,
                    workspace=self.workspace,
                    uploaded=self.image_file("workspace-reference.png"),
                    kind=Asset.Kind.OTHER,
                )
                page = self.client.get(f"/studio/projects/{self.project.id}/")
                self.assertContains(page, "workspace-reference.png")
                self.assertContains(page, "Workspace / 72x54")
                cropped = self.client.post(
                    f"/studio/assets/{source.id}/crop/",
                    data=json.dumps({"x": 4, "y": 3, "width": 40, "height": 30}),
                    content_type="application/json",
                )
                self.assertEqual(cropped.status_code, 201, cropped.content)
                result = Asset.objects.get(id=cropped.json()["id"])
                self.assertIsNone(result.project)
                self.assertEqual((result.width, result.height), (40, 30))

    def test_workspace_gallery_deduplicates_identical_files_and_keeps_project_links(self):
        import tempfile
        from pathlib import Path
        from .models import Asset, Project
        from .storage import create_asset

        second_project = Project.objects.create(
            workspace=self.workspace,
            project_type=Project.Type.SERIES,
            title="Second image project",
            created_by=self.owner,
            updated_by=self.owner,
        )
        self.client.force_login(self.owner)
        with tempfile.TemporaryDirectory() as directory:
            with self.settings(STUDIO_PRIVATE_MEDIA_ROOT=Path(directory)):
                original = create_asset(
                    user=self.owner,
                    workspace=self.workspace,
                    project=self.project,
                    uploaded=self.image_file("original.png"),
                    kind=Asset.Kind.OTHER,
                )
                duplicate = create_asset(
                    user=self.owner,
                    workspace=self.workspace,
                    project=second_project,
                    uploaded=self.image_file("duplicate.png"),
                    kind=Asset.Kind.OTHER,
                )
                self.assertEqual(original.checksum_sha256, duplicate.checksum_sha256)

                response = self.client.get(f"/studio/workspaces/{self.workspace.id}/")
                gallery_assets = response.context["gallery_assets"]
                self.assertEqual(len(gallery_assets), 1)
                self.assertEqual(gallery_assets[0].id, original.id)
                self.assertEqual(
                    set(gallery_assets[0].gallery_project_ids.split(",")),
                    {str(self.project.id), str(second_project.id)},
                )

    def test_workspace_gallery_query_count_does_not_grow_per_image(self):
        import io
        import tempfile
        from pathlib import Path
        from PIL import Image
        from django.core.files.uploadedfile import SimpleUploadedFile
        from .models import Asset
        from .storage import create_asset

        self.client.force_login(self.owner)
        with tempfile.TemporaryDirectory() as directory:
            with self.settings(STUDIO_PRIVATE_MEDIA_ROOT=Path(directory)):
                first_output = io.BytesIO()
                Image.new("RGB", (72, 54), (0, 100, 180)).save(first_output, "PNG")
                create_asset(
                    user=self.owner,
                    workspace=self.workspace,
                    project=self.project,
                    uploaded=SimpleUploadedFile(
                        "performance-0.png",
                        first_output.getvalue(),
                        content_type="image/png",
                    ),
                    kind=Asset.Kind.OTHER,
                )
                with CaptureQueriesContext(connection) as baseline_queries:
                    baseline_response = self.client.get(f"/studio/workspaces/{self.workspace.id}/")
                self.assertEqual(baseline_response.status_code, 200)

                for index in range(1, 8):
                    output = io.BytesIO()
                    Image.new("RGB", (72, 54), (index * 20, 100, 180)).save(output, "PNG")
                    create_asset(
                        user=self.owner,
                        workspace=self.workspace,
                        project=self.project,
                        uploaded=SimpleUploadedFile(
                            f"performance-{index}.png",
                            output.getvalue(),
                            content_type="image/png",
                        ),
                        kind=Asset.Kind.OTHER,
                    )

                with CaptureQueriesContext(connection) as populated_queries:
                    populated_response = self.client.get(f"/studio/workspaces/{self.workspace.id}/")
                self.assertEqual(populated_response.status_code, 200)
                self.assertLessEqual(len(populated_queries), len(baseline_queries) + 1)

    def test_project_image_detaches_but_only_workspace_gallery_offers_trash(self):
        import tempfile
        from pathlib import Path
        from .models import Asset
        from .storage import create_asset

        self.client.force_login(self.owner)
        with tempfile.TemporaryDirectory() as directory:
            with self.settings(STUDIO_PRIVATE_MEDIA_ROOT=Path(directory)):
                asset = create_asset(
                    user=self.owner,
                    workspace=self.workspace,
                    project=self.project,
                    scene=self.first_scene,
                    uploaded=self.image_file("detachable.png"),
                    kind=Asset.Kind.SCENE_IMAGE,
                )
                project_page = self.client.get(f"/studio/projects/{self.project.id}/")
                self.assertContains(project_page, f"/studio/assets/{asset.id}/detach/project/{self.project.id}/")
                self.assertNotContains(project_page, f"/studio/assets/{asset.id}/trash/")
                workspace_page = self.client.get(f"/studio/workspaces/{self.workspace.id}/")
                self.assertContains(workspace_page, f"/studio/assets/{asset.id}/trash/")

                detached = self.client.post(
                    f"/studio/assets/{asset.id}/detach/project/{self.project.id}/",
                    {"next": f"/studio/projects/{self.project.id}/#images"},
                )
                self.assertRedirects(detached, f"/studio/projects/{self.project.id}/#images")
                asset.refresh_from_db()
                self.assertIsNone(asset.project)
                self.assertIsNone(asset.scene)
                self.assertIsNone(asset.deleted_at)

    def test_shared_workspace_image_is_not_duplicated_and_requires_complete_access_to_delete(self):
        import tempfile
        from pathlib import Path
        from .models import Asset
        from .storage import create_asset

        editor = get_user_model().objects.create_user("shared-image-editor", password="strong-pass")
        WorkspaceMembership.objects.create(
            workspace=self.workspace,
            user=editor,
            role=WorkspaceMembership.Role.EDITOR,
        )
        hidden_project = Project.objects.create(
            workspace=self.workspace,
            project_type=Project.Type.SERIES,
            title="Private linked project",
            created_by=self.owner,
            updated_by=self.owner,
        )
        ProjectAccessExclusion.objects.create(
            project=hidden_project,
            user=editor,
            revoked_by=self.owner,
        )

        with tempfile.TemporaryDirectory() as directory:
            with self.settings(STUDIO_PRIVATE_MEDIA_ROOT=Path(directory)):
                asset = create_asset(
                    user=self.owner,
                    workspace=self.workspace,
                    project=self.project,
                    uploaded=self.image_file("shared-once.png"),
                    kind=Asset.Kind.OTHER,
                )
                asset.projects.add(hidden_project)
                self.assertEqual(Asset.objects.count(), 1)

                self.client.force_login(editor)
                page = self.client.get(f"/studio/workspaces/{self.workspace.id}/")
                self.assertEqual(page.status_code, 200)
                self.assertNotContains(page, hidden_project.title)
                denied = self.client.post(
                    f"/studio/assets/{asset.id}/trash/",
                    {"workspace_delete": "1", "confirm_usage": "1"},
                )
                self.assertEqual(denied.status_code, 403)
                asset.refresh_from_db()
                self.assertIsNone(asset.deleted_at)

                self.client.force_login(self.owner)
                needs_confirmation = self.client.post(
                    f"/studio/assets/{asset.id}/trash/",
                    {"workspace_delete": "1"},
                )
                self.assertEqual(needs_confirmation.status_code, 302)
                asset.refresh_from_db()
                self.assertIsNone(asset.deleted_at)
                confirmed = self.client.post(
                    f"/studio/assets/{asset.id}/trash/",
                    {"workspace_delete": "1", "confirm_usage": "1"},
                )
                self.assertEqual(confirmed.status_code, 302)
                asset.refresh_from_db()
                self.assertIsNotNone(asset.deleted_at)

    def test_project_cover_can_be_selected_from_workspace(self):
        import tempfile
        from pathlib import Path
        from .models import Asset
        from .storage import create_asset

        self.client.force_login(self.owner)
        with tempfile.TemporaryDirectory() as directory:
            with self.settings(STUDIO_PRIVATE_MEDIA_ROOT=Path(directory)):
                asset = create_asset(
                    user=self.owner,
                    workspace=self.workspace,
                    project=self.project,
                    uploaded=self.image_file("project-cover.png"),
                    kind=Asset.Kind.OTHER,
                )
                page = self.client.get(f"/studio/workspaces/{self.workspace.id}/")
                self.assertContains(page, f'id="{self.project.id}-cover-picker"')
                self.assertContains(page, "Project Cover")
                self.assertContains(page, "data-preserve-position")

                selected = self.client.post(
                    f"/studio/assets/attach/project_cover/{self.project.id}/",
                    {
                        "asset_id": str(asset.id),
                        "next": f"/studio/workspaces/{self.workspace.id}/#project-{self.project.id}",
                    },
                )
                self.assertRedirects(
                    selected,
                    f"/studio/workspaces/{self.workspace.id}/#project-{self.project.id}",
                )
                self.project.refresh_from_db()
                self.assertEqual(self.project.cover_asset_id, asset.id)

                self.client.force_login(self.viewer)
                forbidden = self.client.post(
                    f"/studio/assets/attach/project_cover/{self.project.id}/",
                    {"asset_id": str(asset.id)},
                )
                self.assertEqual(forbidden.status_code, 403)

    def test_workspace_and_character_avatars_use_shared_image_picker(self):
        import tempfile
        from pathlib import Path
        from .models import Asset, Character
        from .storage import create_asset

        character = Character.objects.create(
            project=self.project,
            name="Murr",
            description="Studio character",
            created_by=self.owner,
            updated_by=self.owner,
        )
        self.client.force_login(self.owner)
        with tempfile.TemporaryDirectory() as directory:
            with self.settings(STUDIO_PRIVATE_MEDIA_ROOT=Path(directory)):
                workspace_asset = create_asset(
                    user=self.owner,
                    workspace=self.workspace,
                    project=None,
                    uploaded=self.image_file("workspace-avatar.png"),
                    kind=Asset.Kind.OTHER,
                )
                character_asset = create_asset(
                    user=self.owner,
                    workspace=self.workspace,
                    project=self.project,
                    uploaded=self.image_file("character-avatar.png"),
                    kind=Asset.Kind.CHARACTER_REFERENCE,
                )

                workspace_page = self.client.get(f"/studio/workspaces/{self.workspace.id}/")
                self.assertContains(workspace_page, 'id="workspace-avatar-picker"')
                self.assertContains(workspace_page, "Workspace Avatar")

                project_page = self.client.get(f"/studio/projects/{self.project.id}/")
                self.assertContains(project_page, f'id="{character.id}-avatar-picker"')
                self.assertContains(project_page, "Character Avatar")
                self.assertNotContains(project_page, "image-modal-footer")

                workspace_selected = self.client.post(
                    f"/studio/assets/attach/workspace_avatar/{self.workspace.id}/",
                    {
                        "asset_id": str(workspace_asset.id),
                        "next": f"/studio/workspaces/{self.workspace.id}/",
                    },
                )
                self.assertRedirects(workspace_selected, f"/studio/workspaces/{self.workspace.id}/")
                self.workspace.refresh_from_db()
                self.assertEqual(self.workspace.avatar_asset_id, workspace_asset.id)

                character_selected = self.client.post(
                    f"/studio/assets/attach/character_avatar/{character.id}/",
                    {
                        "asset_id": str(character_asset.id),
                        "next": f"/studio/projects/{self.project.id}/#characters",
                    },
                )
                self.assertRedirects(
                    character_selected,
                    f"/studio/projects/{self.project.id}/#characters",
                )
                character.refresh_from_db()
                character_asset.refresh_from_db()
                self.assertEqual(character.avatar_asset_id, character_asset.id)
                self.assertTrue(character.reference_assets.filter(id=character_asset.id).exists())
                self.assertTrue(character_asset.projects.filter(id=self.project.id).exists())

    def test_project_header_is_shared_by_all_four_views(self):
        self.client.force_login(self.owner)
        urls = (
            f"/studio/projects/{self.project.id}/",
            f"/studio/projects/{self.project.id}/scene-chain/",
            f"/studio/projects/{self.project.id}/master/",
            f"/studio/projects/{self.project.id}/translations/",
        )
        for url in urls:
            with self.subTest(url=url):
                page = self.client.get(url)
                self.assertEqual(page.status_code, 200)
                self.assertContains(page, f"Project {self.project.title}")
                self.assertContains(page, "General View")
                self.assertContains(page, "Scene Chain")
                self.assertContains(page, "Master Document")
                self.assertContains(page, "Translations")

    def test_workspace_tiles_copy_trash_and_restore(self):
        self.client.force_login(self.owner)
        dashboard = self.client.get("/studio/")
        self.assertContains(dashboard, f"/studio/workspaces/{self.workspace.id}/copy/")
        self.assertContains(dashboard, f"/studio/workspaces/{self.workspace.id}/trash/")
        copied_response = self.client.post(f"/studio/workspaces/{self.workspace.id}/copy/")
        copied = accessible_workspaces(self.owner).exclude(id=self.workspace.id).get()
        self.assertRedirects(copied_response, f"/studio/workspaces/{copied.id}/")
        trashed = self.client.post(f"/studio/workspaces/{copied.id}/trash/")
        self.assertRedirects(trashed, "/studio/")
        self.assertFalse(accessible_workspaces(self.owner).filter(id=copied.id).exists())
        restored = self.client.post(f"/studio/workspaces/{copied.id}/restore/")
        self.assertRedirects(restored, f"/studio/workspaces/{copied.id}/")
        self.assertTrue(accessible_workspaces(self.owner).filter(id=copied.id).exists())

    def test_prompt_can_attach_existing_project_image(self):
        import tempfile
        from pathlib import Path
        from .models import Asset, Prompt
        from .storage import create_asset

        prompt = Prompt.objects.create(scene=self.first_scene, ai_model=self.ai_model, prompt_type="IMAGE", title="References", created_by=self.owner, updated_by=self.owner)
        self.client.force_login(self.owner)
        with tempfile.TemporaryDirectory() as directory:
            with self.settings(STUDIO_PRIVATE_MEDIA_ROOT=Path(directory)):
                asset = create_asset(user=self.owner, workspace=self.workspace, project=self.project, uploaded=self.image_file("library.png"), kind=Asset.Kind.OTHER)
                other_project = Project.objects.create(
                    workspace=self.workspace,
                    title="Second accessible project",
                    created_by=self.owner,
                    updated_by=self.owner,
                )
                asset.projects.add(other_project)
                response = self.client.post(f"/studio/prompts/{prompt.id}/images/attach/", {"asset_id": str(asset.id)})
                self.assertRedirects(response, f"/studio/scenes/{self.first_scene.id}/#prompt-{prompt.id}")
                self.assertTrue(prompt.reference_assets.filter(id=asset.id).exists())

    def test_project_copy_and_soft_delete_restore(self):
        import tempfile
        from pathlib import Path
        from .models import Asset
        from .storage import create_asset

        self.client.force_login(self.owner)
        with tempfile.TemporaryDirectory() as directory:
            with self.settings(STUDIO_PRIVATE_MEDIA_ROOT=Path(directory)):
                source_asset = create_asset(user=self.owner, workspace=self.workspace, project=self.project, uploaded=self.image_file("cover.png"), kind=Asset.Kind.OTHER)
                copied_response = self.client.post(f"/studio/projects/{self.project.id}/copy/")
                copied = Project.objects.exclude(id=self.project.id).get()
                self.assertRedirects(copied_response, f"/studio/projects/{copied.id}/")
                self.assertEqual(copied.episodes.count(), 1)
                self.assertEqual(copied.episodes.get().scenes.count(), 2)
                self.assertTrue(copied.media_assets.filter(id=source_asset.id).exists())
                self.assertEqual(Asset.objects.count(), 1)
                changed_cover = self.client.post(
                    f"/studio/assets/attach/project_cover/{copied.id}/",
                    {
                        "asset_id": str(source_asset.id),
                        "next": f"/studio/workspaces/{self.workspace.id}/#project-{copied.id}",
                    },
                )
                self.assertRedirects(
                    changed_cover,
                    f"/studio/workspaces/{self.workspace.id}/#project-{copied.id}",
                )
                copied.refresh_from_db()
                self.assertEqual(copied.cover_asset_id, source_asset.id)
                deleted = self.client.post(f"/studio/projects/{copied.id}/trash/")
                self.assertRedirects(deleted, f"/studio/workspaces/{self.workspace.id}/")
                self.assertFalse(Project.objects.filter(id=copied.id).exists())
                restored = self.client.post(f"/studio/projects/{copied.id}/restore/")
                self.assertRedirects(restored, f"/studio/projects/{copied.id}/")
                self.assertTrue(Project.objects.filter(id=copied.id).exists())

    def test_viewer_sees_inline_prompts_but_cannot_change_them(self):
        from .models import Prompt

        prompt = Prompt.objects.create(
            scene=self.first_scene, ai_model=self.ai_model, prompt_type="IMAGE",
            title="Read only", created_by=self.owner, updated_by=self.owner,
        )
        self.client.force_login(self.viewer)
        page = self.client.get(f"/studio/scenes/{self.first_scene.id}/")
        self.assertEqual(page.status_code, 200)
        self.assertContains(page, "prompt-fieldset")
        self.assertContains(page, "disabled")
        denied = self.client.post(
            f"/studio/prompts/{prompt.id}/quick-save/",
            {"title": "Changed", "ai_model": str(self.ai_model.id), "prompt_type": "IMAGE", "status": "DRAFT"},
        )
        self.assertEqual(denied.status_code, 403)
        reorder = self.client.post(
            f"/studio/episodes/{self.episode.id}/scenes/reorder/",
            data=json.dumps({"sceneIds": [str(self.second_scene.id), str(self.first_scene.id)]}),
            content_type="application/json",
        )
        self.assertEqual(reorder.status_code, 403)


class StudioAssetLifecycleTests(TestCase):
    def setUp(self):
        from .models import Project

        users = get_user_model()
        self.owner = users.objects.create_user("asset-owner", password="strong-pass")
        self.viewer = users.objects.create_user("asset-viewer", password="strong-pass")
        self.workspace = create_workspace(user=self.owner, name="Asset Studio", slug="asset-studio")
        WorkspaceMembership.objects.create(
            workspace=self.workspace, user=self.viewer, role=WorkspaceMembership.Role.VIEWER
        )
        self.project = Project.objects.create(
            workspace=self.workspace, project_type=Project.Type.SERIES, title="Asset Project",
            created_by=self.owner, updated_by=self.owner,
        )

    @staticmethod
    def image_file(name="lifecycle.png"):
        import io
        from PIL import Image
        from django.core.files.uploadedfile import SimpleUploadedFile

        output = io.BytesIO()
        Image.new("RGB", (80, 60), "#d87a31").save(output, "PNG")
        return SimpleUploadedFile(name, output.getvalue(), content_type="image/png")

    def test_image_can_be_trashed_restored_and_permanently_purged(self):
        import tempfile
        from pathlib import Path
        from .models import Asset, AuditEvent
        from .storage import create_asset

        self.client.force_login(self.owner)
        with tempfile.TemporaryDirectory() as directory:
            with self.settings(STUDIO_PRIVATE_MEDIA_ROOT=Path(directory)):
                asset = create_asset(
                    user=self.owner, workspace=self.workspace, project=self.project,
                    uploaded=self.image_file(), kind=Asset.Kind.OTHER,
                )
                original_path = Path(asset.file.path)
                thumbnail_path = Path(asset.thumbnail.path)
                trashed = self.client.post(
                    f"/studio/assets/{asset.id}/trash/",
                    {
                        "next": f"/studio/projects/{self.project.id}/images/trash/",
                        "workspace_delete": "1",
                        "confirm_usage": "1",
                    },
                )
                self.assertEqual(trashed.status_code, 302)
                self.assertFalse(Asset.objects.filter(id=asset.id).exists())
                self.assertEqual(self.client.get(f"/api/v1/studio/assets/{asset.id}/view").status_code, 404)
                trash_page = self.client.get(f"/studio/workspaces/{self.workspace.id}/recycle-bin/")
                self.assertContains(trash_page, asset.original_filename)
                trash_preview = self.client.get(f"/studio/assets/{asset.id}/trash-thumbnail/")
                self.assertEqual(trash_preview.status_code, 200)
                trash_preview.close()

                restored = self.client.post(f"/studio/assets/{asset.id}/restore/")
                self.assertEqual(restored.status_code, 302)
                self.assertTrue(Asset.objects.filter(id=asset.id).exists())

                with self.captureOnCommitCallbacks(execute=True):
                    purged = self.client.post(f"/studio/assets/{asset.id}/purge/")
                self.assertEqual(purged.status_code, 302)
                asset = Asset.all_objects.get(id=asset.id)
                self.assertIsNotNone(asset.purged_at)
                self.assertEqual(asset.file.name, "")
                self.assertFalse(original_path.exists())
                self.assertFalse(thumbnail_path.exists())
        self.assertTrue(AuditEvent.objects.filter(action="ASSET_TRASHED").exists())
        self.assertTrue(AuditEvent.objects.filter(action="ASSET_RESTORED").exists())
        self.assertTrue(AuditEvent.objects.filter(action="ASSET_PURGED").exists())

    def test_recycle_bin_removes_stale_database_record_when_file_is_missing(self):
        import tempfile
        from pathlib import Path
        from .models import Asset
        from .storage import create_asset, trash_asset

        self.client.force_login(self.owner)
        with tempfile.TemporaryDirectory() as directory:
            with self.settings(STUDIO_PRIVATE_MEDIA_ROOT=Path(directory)):
                asset = create_asset(
                    user=self.owner,
                    workspace=self.workspace,
                    uploaded=self.image_file("missing.png"),
                    kind=Asset.Kind.OTHER,
                )
                trash_asset(asset=asset, user=self.owner)
                Path(asset.file.path).unlink()
                with self.captureOnCommitCallbacks(execute=True):
                    response = self.client.get(
                        f"/studio/workspaces/{self.workspace.id}/recycle-bin/"
                    )
                self.assertEqual(response.status_code, 200)
                self.assertNotContains(response, "missing.png")
                stale = Asset.all_objects.get(id=asset.id)
                self.assertIsNotNone(stale.purged_at)
                self.assertEqual(stale.file.name, "")

    def test_avatar_crop_can_be_attached_in_one_request(self):
        import tempfile
        from pathlib import Path
        from .models import Asset
        from .storage import create_asset

        self.client.force_login(self.owner)
        with tempfile.TemporaryDirectory() as directory:
            with self.settings(STUDIO_PRIVATE_MEDIA_ROOT=Path(directory)):
                source = create_asset(
                    user=self.owner,
                    workspace=self.workspace,
                    uploaded=self.image_file("avatar-source.png"),
                    kind=Asset.Kind.OTHER,
                )
                response = self.client.post(
                    f"/studio/assets/{source.id}/crop/",
                    data=json.dumps({
                        "x": 10, "y": 5, "width": 50, "height": 50,
                        "attachScope": "workspace_avatar",
                        "attachOwnerId": str(self.workspace.id),
                    }),
                    content_type="application/json",
                )
                self.assertEqual(response.status_code, 201, response.content)
                self.assertTrue(response.json()["attached"])
                self.workspace.refresh_from_db()
                self.assertEqual(str(self.workspace.avatar_asset_id), response.json()["id"])

    def test_avatar_only_image_can_be_restored_and_purge_clears_avatar_link(self):
        import tempfile
        from pathlib import Path
        from .models import Asset
        from .storage import create_asset, trash_asset

        self.client.force_login(self.owner)
        with tempfile.TemporaryDirectory() as directory:
            with self.settings(STUDIO_PRIVATE_MEDIA_ROOT=Path(directory)):
                asset = create_asset(
                    user=self.owner,
                    workspace=self.workspace,
                    uploaded=self.image_file("workspace-avatar.png"),
                    kind=Asset.Kind.OTHER,
                )
                self.workspace.avatar_asset = asset
                self.workspace.save(update_fields=["avatar_asset", "updated_at"])
                trash_asset(asset=asset, user=self.owner)

                preview = self.client.get(f"/studio/assets/{asset.id}/trash-thumbnail/")
                self.assertEqual(preview.status_code, 200)
                preview.close()
                restored = self.client.post(
                    f"/studio/assets/{asset.id}/restore/",
                    {"next": f"/studio/workspaces/{self.workspace.id}/recycle-bin/"},
                )
                self.assertRedirects(
                    restored,
                    f"/studio/workspaces/{self.workspace.id}/recycle-bin/",
                )
                self.assertTrue(Asset.objects.filter(id=asset.id).exists())

                trash_asset(asset=asset, user=self.owner)
                with self.captureOnCommitCallbacks(execute=True):
                    purged = self.client.post(
                        f"/studio/assets/{asset.id}/purge/",
                        {"next": f"/studio/workspaces/{self.workspace.id}/recycle-bin/"},
                    )
                self.assertRedirects(
                    purged,
                    f"/studio/workspaces/{self.workspace.id}/recycle-bin/",
                )
                self.workspace.refresh_from_db()
                self.assertIsNone(self.workspace.avatar_asset_id)
                self.assertIsNotNone(Asset.all_objects.get(id=asset.id).purged_at)

    def test_empty_recycle_bin_purges_all_workspace_files(self):
        import tempfile
        from pathlib import Path
        from .models import Asset, Character
        from .storage import create_asset, trash_asset

        self.client.force_login(self.owner)
        character = Character.objects.create(
            project=self.project,
            name="Recycle avatar",
            created_by=self.owner,
            updated_by=self.owner,
        )
        with tempfile.TemporaryDirectory() as directory:
            with self.settings(STUDIO_PRIVATE_MEDIA_ROOT=Path(directory)):
                first = create_asset(
                    user=self.owner,
                    workspace=self.workspace,
                    uploaded=self.image_file("first.png"),
                    kind=Asset.Kind.OTHER,
                )
                second = create_asset(
                    user=self.owner,
                    workspace=self.workspace,
                    uploaded=self.image_file("second.png"),
                    kind=Asset.Kind.OTHER,
                )
                character.avatar_asset = second
                character.save(update_fields=["avatar_asset", "updated_at"])
                trash_asset(asset=first, user=self.owner)
                trash_asset(asset=second, user=self.owner)

                with self.captureOnCommitCallbacks(execute=True):
                    response = self.client.post(
                        f"/studio/workspaces/{self.workspace.id}/recycle-bin/clear/",
                    )
                self.assertRedirects(
                    response,
                    f"/studio/workspaces/{self.workspace.id}/recycle-bin/",
                )
                self.assertEqual(
                    Asset.all_objects.filter(
                        workspace=self.workspace,
                        deleted_at__isnull=False,
                        purged_at__isnull=True,
                    ).count(),
                    0,
                )
                character.refresh_from_db()
                self.assertIsNone(character.avatar_asset_id)

    def test_project_copy_to_another_workspace_copies_image_files(self):
        import tempfile
        from pathlib import Path
        from .models import Asset, Character
        from .storage import create_asset

        target_workspace = create_workspace(
            user=self.owner,
            name="Target Studio",
            slug="target-studio",
        )
        character = Character.objects.create(
            project=self.project,
            name="Source character",
            created_by=self.owner,
            updated_by=self.owner,
        )
        self.client.force_login(self.owner)
        with tempfile.TemporaryDirectory() as directory:
            with self.settings(STUDIO_PRIVATE_MEDIA_ROOT=Path(directory)):
                source_asset = create_asset(
                    user=self.owner,
                    workspace=self.workspace,
                    project=self.project,
                    uploaded=self.image_file("cross-workspace-cover.png"),
                    kind=Asset.Kind.CHARACTER_REFERENCE,
                )
                self.project.cover_asset = source_asset
                self.project.save(update_fields=["cover_asset", "updated_at"])
                character.avatar_asset = source_asset
                character.reference_assets.add(source_asset)
                character.save(update_fields=["avatar_asset", "updated_at"])

                response = self.client.post(
                    f"/studio/projects/{self.project.id}/copy/",
                    {"target_workspace": str(target_workspace.id)},
                )
                copied = target_workspace.projects.get()
                self.assertRedirects(response, f"/studio/projects/{copied.id}/")
                copied.refresh_from_db()
                copied_character = copied.characters.get(name=character.name)
                self.assertIsNotNone(copied.cover_asset)
                self.assertNotEqual(copied.cover_asset_id, source_asset.id)
                self.assertEqual(copied.cover_asset.workspace_id, target_workspace.id)
                self.assertEqual(
                    copied.cover_asset.checksum_sha256,
                    source_asset.checksum_sha256,
                )
                self.assertTrue(copied.cover_asset.file.name)
                self.assertEqual(
                    copied_character.avatar_asset_id,
                    copied.cover_asset_id,
                )
                self.assertTrue(
                    copied_character.reference_assets.filter(
                        id=copied.cover_asset_id,
                    ).exists()
                )
                self.assertTrue(Asset.objects.filter(id=source_asset.id).exists())

    def test_viewer_cannot_delete_image(self):
        import tempfile
        from pathlib import Path
        from .models import Asset
        from .storage import create_asset

        with tempfile.TemporaryDirectory() as directory:
            with self.settings(STUDIO_PRIVATE_MEDIA_ROOT=Path(directory)):
                asset = create_asset(
                    user=self.owner, workspace=self.workspace, project=self.project,
                    uploaded=self.image_file(), kind=Asset.Kind.OTHER,
                )
                self.client.force_login(self.viewer)
                self.assertEqual(self.client.post(f"/studio/assets/{asset.id}/trash/").status_code, 403)
                self.assertTrue(Asset.objects.filter(id=asset.id).exists())

    def test_project_detach_keeps_workspace_file_and_attach_restores_link(self):
        import tempfile
        from pathlib import Path
        from .models import Asset
        from .storage import create_asset

        self.client.force_login(self.owner)
        with tempfile.TemporaryDirectory() as directory:
            with self.settings(STUDIO_PRIVATE_MEDIA_ROOT=Path(directory)):
                asset = create_asset(
                    user=self.owner, workspace=self.workspace, project=self.project,
                    uploaded=self.image_file("shared.png"), kind=Asset.Kind.OTHER,
                )
                detached = self.client.post(
                    f"/studio/assets/{asset.id}/detach/project/{self.project.id}/",
                    {"next": f"/studio/projects/{self.project.id}/#images"},
                )
                self.assertEqual(detached.status_code, 302)
                asset.refresh_from_db()
                self.assertTrue(Asset.objects.filter(id=asset.id, workspace=self.workspace).exists())
                self.assertFalse(asset.projects.filter(id=self.project.id).exists())
                self.assertIsNone(asset.project_id)

                repeated = self.client.post(
                    f"/studio/assets/{asset.id}/detach/project/{self.project.id}/",
                    {"next": f"/studio/projects/{self.project.id}/#images"},
                )
                self.assertEqual(repeated.status_code, 302)

                attached = self.client.post(
                    f"/studio/assets/attach/project/{self.project.id}/",
                    {"asset_id": str(asset.id), "next": f"/studio/projects/{self.project.id}/#images"},
                )
                self.assertEqual(attached.status_code, 302)
                self.assertTrue(asset.projects.filter(id=self.project.id).exists())


class StudioGeneralSettingsTests(TestCase):
    def setUp(self):
        users = get_user_model()
        self.admin = users.objects.create_superuser("studio-admin", "admin@example.com", "strong-pass")
        self.user = users.objects.create_user("studio-user", password="strong-pass")

    def test_settings_are_admin_only_and_can_change_default_model(self):
        from .models import StudioTextModel

        self.client.force_login(self.user)
        self.assertEqual(self.client.get("/studio/settings/").status_code, 403)
        self.client.force_login(self.admin)
        model = StudioTextModel.objects.get(model_id="gpt-5.4")
        response = self.client.post("/studio/settings/", {
            "action": "update_model",
            "id": str(model.id),
            "model-name": "GPT 5.4 production",
            "model-model_id": "gpt-5.4",
            "model-is_active": "on",
            "model-is_default": "on",
        })
        self.assertRedirects(response, "/studio/settings/")
        model.refresh_from_db()
        self.assertTrue(model.is_default)
        self.assertFalse(StudioTextModel.objects.exclude(id=model.id).filter(is_default=True).exists())

    def test_admin_can_update_default_prompt_addition(self):
        from .models import PromptTemplate

        self.client.force_login(self.admin)
        response = self.client.post("/studio/settings/", {
            "action": "update_prompt_addition",
            "prompt_addition": "Negative Prompt: No Music",
        })
        self.assertRedirects(response, "/studio/settings/")
        template = PromptTemplate.objects.get(is_default=True)
        self.assertEqual(template.content, "Negative Prompt: No Music")
        page = self.client.get("/studio/settings/")
        self.assertContains(page, "Default prompt addition")
        self.assertNotContains(page, "Create template")

    def test_admin_can_add_or_activate_model_from_prompt_dropdown(self):
        from .models import StudioTextModel

        self.client.force_login(self.admin)
        response = self.client.post(
            "/studio/settings/text-models/quick-create/",
            {"name": "Fast prompt model", "model_id": "gpt-5.4-nano"},
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(response.json()["modelId"], "gpt-5.4-nano")
        model = StudioTextModel.objects.get(model_id="gpt-5.4-nano")
        self.assertEqual(model.name, "Fast prompt model")
        self.assertTrue(model.is_active)

        self.client.force_login(self.user)
        denied = self.client.post(
            "/studio/settings/text-models/quick-create/",
            {"name": "Forbidden", "model_id": "gpt-5.4-mini"},
        )
        self.assertEqual(denied.status_code, 403)


class StudioProductionPilotFeaturesTests(TestCase):
    @staticmethod
    def image_file(name="pilot-avatar.png"):
        import io
        from PIL import Image
        from django.core.files.uploadedfile import SimpleUploadedFile

        output = io.BytesIO()
        Image.new("RGB", (96, 96), "#d87a31").save(output, "PNG")
        return SimpleUploadedFile(name, output.getvalue(), content_type="image/png")

    def setUp(self):
        from .models import AiModelProfile, Character

        users = get_user_model()
        self.owner = users.objects.create_user("pilot-owner", email="owner@example.com", password="strong-pass")
        self.new_owner = users.objects.create_user("pilot-new-owner", email="new-owner@example.com", password="strong-pass")
        self.admin = users.objects.create_superuser("pilot-admin", email="admin@example.com", password="strong-pass")
        self.workspace = create_workspace(user=self.owner, name="Pilot Studio", slug="pilot-studio")
        self.project = Project.objects.create(
            workspace=self.workspace,
            project_type=Project.Type.SERIES,
            title="Pilot Project",
            created_by=self.owner,
            updated_by=self.owner,
        )
        self.character = Character.objects.create(
            project=self.project,
            name="Hero",
            description="Original hero",
            created_by=self.owner,
            updated_by=self.owner,
        )
        self.image_model = AiModelProfile.objects.create(
            name="GPT Image Test",
            provider="OpenAI",
            model_id="gpt-image-test",
            media_type=AiModelProfile.MediaType.IMAGE,
            is_active=True,
        )
        self.workspace.default_image_model = self.image_model
        self.workspace.save(update_fields=["default_image_model", "updated_at"])

    def test_superuser_can_transfer_workspace_owner(self):
        self.client.force_login(self.admin)
        response = self.client.post(
            f"/studio/workspaces/{self.workspace.id}/owner/",
            {"email": self.new_owner.email},
        )
        self.assertRedirects(response, f"/studio/workspaces/{self.workspace.id}/access/")
        self.workspace.refresh_from_db()
        self.assertEqual(self.workspace.owner, self.new_owner)
        self.assertEqual(
            WorkspaceMembership.objects.get(workspace=self.workspace, user=self.owner).role,
            WorkspaceMembership.Role.ADMIN,
        )
        self.assertEqual(
            WorkspaceMembership.objects.get(workspace=self.workspace, user=self.new_owner).role,
            WorkspaceMembership.Role.OWNER,
        )

    def test_superadmin_group_member_can_transfer_workspace_owner(self):
        from django.contrib.auth.models import Group

        delegated_admin = get_user_model().objects.create_user(
            "delegated-admin", email="delegated@example.com", password="strong-pass"
        )
        delegated_admin.groups.add(Group.objects.get_or_create(name="Superadmin")[0])
        self.client.force_login(delegated_admin)
        access_page = self.client.get(f"/studio/workspaces/{self.workspace.id}/access/")
        self.assertEqual(access_page.status_code, 200)
        self.assertContains(access_page, "Transfer workspace ownership")
        response = self.client.post(
            f"/studio/workspaces/{self.workspace.id}/owner/",
            {"email": self.new_owner.email},
        )
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.url, f"/studio/workspaces/{self.workspace.id}/access/")
        self.workspace.refresh_from_db()
        self.assertEqual(self.workspace.owner, self.new_owner)

    def test_project_users_can_be_added_in_bulk(self):
        first = get_user_model().objects.create_user(
            "bulk-one", email="bulk-one@example.com", password="strong-pass"
        )
        second = get_user_model().objects.create_user(
            "bulk-two", email="bulk-two@example.com", password="strong-pass"
        )
        self.client.force_login(self.owner)
        response = self.client.post(
            f"/studio/projects/{self.project.id}/access/",
            {
                "action": "bulk_add",
                "emails": "bulk-one@example.com; bulk-two@example.com\nmissing@example.com",
                "role": ProjectMembership.Role.EDITOR,
            },
        )
        self.assertRedirects(response, f"/studio/projects/{self.project.id}/access/")
        self.assertTrue(ProjectMembership.objects.filter(project=self.project, user=first, role=ProjectMembership.Role.EDITOR).exists())
        self.assertTrue(ProjectMembership.objects.filter(project=self.project, user=second, role=ProjectMembership.Role.EDITOR).exists())

    def test_superuser_can_select_registered_project_users(self):
        first = get_user_model().objects.create_user(
            "selected-one", email="selected-one@example.com", password="strong-pass"
        )
        second = get_user_model().objects.create_user(
            "selected-two", email="selected-two@example.com", password="strong-pass"
        )
        self.client.force_login(self.admin)
        page = self.client.get(f"/studio/projects/{self.project.id}/access/")
        self.assertContains(page, "Select registered users")
        response = self.client.post(
            f"/studio/projects/{self.project.id}/access/",
            {
                "action": "select_users",
                "users": [str(first.id), str(second.id)],
                "role": ProjectMembership.Role.EDITOR,
            },
        )
        self.assertRedirects(response, f"/studio/projects/{self.project.id}/access/")
        self.assertEqual(
            ProjectMembership.objects.filter(project=self.project, user__in=[first, second]).count(),
            2,
        )

    def test_superuser_can_select_registered_workspace_users(self):
        first = get_user_model().objects.create_user(
            "workspace-one", email="workspace-one@example.com", password="strong-pass"
        )
        second = get_user_model().objects.create_user(
            "workspace-two", email="workspace-two@example.com", password="strong-pass"
        )
        self.client.force_login(self.admin)
        page = self.client.get(f"/studio/workspaces/{self.workspace.id}/access/")
        self.assertContains(page, "Select registered users")
        response = self.client.post(
            f"/studio/workspaces/{self.workspace.id}/access/",
            {
                "action": "select_users",
                "users": [str(first.id), str(second.id)],
                "role": WorkspaceMembership.Role.EDITOR,
                "can_use_ai": "on",
            },
        )
        self.assertRedirects(response, f"/studio/workspaces/{self.workspace.id}/access/")
        memberships = WorkspaceMembership.objects.filter(workspace=self.workspace, user__in=[first, second])
        self.assertEqual(memberships.count(), 2)
        self.assertFalse(memberships.filter(can_use_ai=False).exists())

    def test_character_can_be_copied_and_soft_deleted(self):
        from .models import Character

        self.client.force_login(self.owner)
        copied_response = self.client.post(f"/studio/characters/{self.character.id}/copy/")
        copied = Character.objects.exclude(id=self.character.id).get()
        self.assertRedirects(copied_response, f"/studio/characters/{copied.id}/")
        self.assertEqual(copied.description, self.character.description)
        deleted_response = self.client.post(f"/studio/characters/{copied.id}/delete/")
        self.assertRedirects(deleted_response, f"/studio/projects/{self.project.id}/")
        self.assertFalse(Character.objects.filter(id=copied.id).exists())
        self.assertTrue(Character.all_objects.filter(id=copied.id, deleted_at__isnull=False).exists())

    def test_uploaded_character_avatar_is_cropped_before_assignment(self):
        import tempfile
        from pathlib import Path

        self.client.force_login(self.owner)
        with tempfile.TemporaryDirectory() as directory:
            with self.settings(STUDIO_PRIVATE_MEDIA_ROOT=Path(directory)):
                uploaded = self.client.post(
                    f"/studio/characters/{self.character.id}/images/new/",
                    {"file": self.image_file(), "image_role": "character_avatar"},
                    HTTP_X_REQUESTED_WITH="XMLHttpRequest",
                )
                self.assertEqual(uploaded.status_code, 201, uploaded.content)
                payload = uploaded.json()
                self.character.refresh_from_db()
                self.assertIsNone(self.character.avatar_asset_id)
                cropped = self.client.post(
                    payload["cropUrl"],
                    data=json.dumps({
                        "x": 0, "y": 0, "width": 80, "height": 80,
                        "attachScope": "character_avatar",
                        "attachOwnerId": str(self.character.id),
                    }),
                    content_type="application/json",
                )
                self.assertEqual(cropped.status_code, 201, cropped.content)
                self.character.refresh_from_db()
                self.assertEqual(str(self.character.avatar_asset_id), cropped.json()["id"])

    def test_character_avatar_can_be_removed_without_deleting_workspace_image(self):
        import tempfile
        from pathlib import Path
        from .models import Asset
        from .storage import create_asset

        self.client.force_login(self.owner)
        with tempfile.TemporaryDirectory() as directory:
            with self.settings(STUDIO_PRIVATE_MEDIA_ROOT=Path(directory)):
                asset = create_asset(
                    user=self.owner,
                    workspace=self.workspace,
                    uploaded=self.image_file("removable-avatar.png"),
                    kind=Asset.Kind.OTHER,
                )
                asset.projects.add(self.project)
                self.character.avatar_asset = asset
                self.character.save(update_fields=["avatar_asset", "updated_at"])
                self.project.cover_asset = asset
                self.project.save(update_fields=["cover_asset", "updated_at"])

                page = self.client.get(f"/studio/projects/{self.project.id}/")
                self.assertContains(page, f"/studio/characters/{self.character.id}/avatar/remove/")
                self.assertContains(page, f"/api/v1/studio/assets/{asset.id}/thumbnail")
                self.assertContains(page, "project-shell-avatar")
                self.assertContains(page, "getElementById(selector.slice(1))")

                removed = self.client.post(
                    f"/studio/characters/{self.character.id}/avatar/remove/",
                    HTTP_X_REQUESTED_WITH="XMLHttpRequest",
                )
                self.assertEqual(removed.status_code, 200, removed.content)
                self.character.refresh_from_db()
                self.project.refresh_from_db()
                self.assertIsNone(self.character.avatar_asset_id)
                self.assertEqual(self.project.cover_asset_id, asset.id)
                self.assertTrue(Asset.objects.filter(id=asset.id).exists())

    def test_speech_preferences_are_saved_per_user(self):
        from .models import StudioUserPreference

        self.client.force_login(self.owner)
        response = self.client.post(
            "/studio/preferences/speech/",
            data=json.dumps({"language": "pl-PL", "translationLanguage": "BY", "continuous": True, "interim": False}),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 200, response.content)
        preference = StudioUserPreference.objects.get(user=self.owner)
        self.assertEqual(preference.speech_language, "pl-PL")
        self.assertTrue(preference.speech_continuous)
        self.assertFalse(preference.speech_interim)
        self.assertEqual(preference.translation_language, "BY")

    def test_user_can_pause_and_resume_ai_usage(self):
        from .models import StudioUserPreference

        self.client.force_login(self.owner)
        paused = self.client.post("/studio/preferences/ai-toggle/")
        self.assertEqual(paused.status_code, 200, paused.content)
        self.assertFalse(paused.json()["enabled"])
        self.assertFalse(StudioUserPreference.objects.get(user=self.owner).ai_enabled)
        resumed = self.client.post("/studio/preferences/ai-toggle/")
        self.assertTrue(resumed.json()["enabled"])

    @patch("lexamora_studio.views.user_has_ai_access", return_value=True)
    def test_video_generation_ignores_image_only_options(self, _has_access):
        from .models import AiModelProfile, ImageGenerationJob

        video_model = AiModelProfile.objects.create(
            name="Veo capability test",
            provider="Google",
            model_id="veo-capability-test",
            media_type=AiModelProfile.MediaType.VIDEO,
            defaults={
                "max_references": 3,
                "max_outputs": 2,
                "ui_fields": ["video-ratio", "video-size", "quantity"],
            },
        )
        self.client.force_login(self.owner)
        response = self.client.post(
            f"/studio/projects/{self.project.id}/image-generation/queue/",
            data=json.dumps({
                "prompt": "A short vertical motion test",
                "modelProfileId": str(video_model.id),
                "size": "720p",
                "aspectRatio": "9:16",
                "quality": "high",
                "outputFormat": "jpeg",
                "referenceAssetIds": [],
            }),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 202, response.content)
        options = ImageGenerationJob.objects.get(id=response.json()["jobId"]).options
        self.assertEqual(options["size"], "720p")
        self.assertEqual(options["aspect_ratio"], "9:16")
        self.assertNotIn("quality", options)
        self.assertNotIn("output_format", options)

    @patch("lexamora_studio.views.user_has_ai_access", return_value=True)
    def test_video_generation_supports_ingredients_and_duration(self, _has_access):
        from .models import AiModelProfile, ImageGenerationJob

        self.client.force_login(self.owner)
        video_model = AiModelProfile.objects.create(
            name="Veo ingredients test",
            provider="Google",
            model_id="veo-ingredients-test",
            media_type=AiModelProfile.MediaType.VIDEO,
            defaults={
                "max_references": 2,
                "max_ingredient_references": 3,
                "reference_modes": ["FRAMES", "INGREDIENTS"],
                "durations": [4, 8],
                "duration": 8,
                "ui_fields": ["video-reference-mode", "video-ratio", "video-size", "video-duration"],
            },
        )
        response = self.client.post(
            f"/studio/projects/{self.project.id}/image-generation/queue/",
            data=json.dumps({
                "prompt": "Use the selected characters as visual ingredients",
                "modelProfileId": str(video_model.id),
                "referenceMode": "INGREDIENTS",
                "durationSeconds": 4,
                "size": "720p",
                "aspectRatio": "16:9",
                "referenceAssetIds": [],
            }),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 202, response.content)
        options = ImageGenerationJob.objects.get(id=response.json()["jobId"]).options
        self.assertEqual(options["reference_mode"], "INGREDIENTS")
        self.assertEqual(options["duration_seconds"], 4)
        self.assertNotIn("first_frame_asset_id", options)

    def test_movie_editor_saves_project_timeline(self):
        from .models import MovieTimeline, MovieTimelineRevision

        self.client.force_login(self.owner)
        page = self.client.get(f"/studio/projects/{self.project.id}/movie-editor/")
        self.assertEqual(page.status_code, 200)
        self.assertContains(page, "Auto rough cut")
        self.assertContains(page, "Draft Editor")
        self.assertNotContains(page, "Movie Editor 2")
        self.assertContains(page, "data-timeline-split")
        self.assertContains(page, "data-timeline-zoom")
        self.assertContains(page, "data-playhead")
        self.assertContains(page, "data-marquee")
        self.assertContains(page, "data-render-start")
        self.assertContains(page, "Rough-cut exports")
        self.assertContains(page, "data-editor-undo")
        self.assertContains(page, "data-preview-scrub")
        self.assertContains(page, "data-preview-zoom")
        self.assertContains(page, "data-media-view-cycle")
        self.assertContains(page, "data-track-width-restore")
        self.assertContains(page, "data-editor-layout")
        self.assertContains(page, 'data-layout-mode="columns"')
        self.assertContains(page, 'data-layout-mode="stacked"')
        self.assertContains(page, "data-layout-save")
        self.assertContains(page, "data-layout-load")
        self.assertContains(page, "data-layout-reset")
        self.assertContains(page, "data-timeline-height-resizer")
        self.assertContains(page, 'data-timeline-edge-resizer="left"')
        self.assertContains(page, 'data-timeline-edge-resizer="right"')
        self.assertContains(page, 'data-panel-width-resizer="media"')
        self.assertContains(page, 'data-panel-width-resizer="inspector"')
        self.assertNotContains(page, "data-panel-resizer")
        self.assertContains(page, "data-media-context-cut")
        self.assertContains(page, "data-media-context-root")
        self.assertContains(page, "data-stage-resizer")
        self.assertContains(page, "data-inspector-resizer")
        self.assertContains(page, "data-panel-height-resizer", count=4)
        self.assertContains(page, "data-preview-dock")
        self.assertContains(page, "data-track-sidebar-resizer")
        self.assertContains(page, "data-clip-render-toolbar")
        self.assertContains(page, 'data-inspector-tab="VIDEO"')
        self.assertContains(page, "data-media-library-dialog")
        self.assertContains(page, "Export As")
        self.assertContains(page, "data-delete-track")
        self.assertContains(page, "data-clip-context")
        self.assertContains(page, "movie-timeline-position-pad")
        self.assertContains(page, "studio/movie_editor.js")
        response = self.client.post(
            f"/studio/projects/{self.project.id}/movie-editor/",
            data=json.dumps({
                "title": "Episode draft",
                "aspectRatio": "9:16",
                "resolution": "1080x1920",
                "fps": 25,
                "timeline": {"tracks": [{"id": "video-1", "kind": "VIDEO", "clips": []}]},
            }),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 200, response.content)
        timeline = MovieTimeline.objects.get(project=self.project)
        self.assertEqual(timeline.title, "Episode draft")
        self.assertEqual(timeline.aspect_ratio, "9:16")
        self.assertEqual(len(timeline.timeline["tracks"]), 1)
        self.assertEqual(timeline.timeline["schemaVersion"], 1)
        self.assertEqual(MovieTimelineRevision.objects.filter(timeline=timeline).count(), 1)

        from .movie_timeline import normalize_movie_timeline
        framed = normalize_movie_timeline({
            "schemaVersion": 1,
            "mediaAssetIds": ["asset-1", "asset-2"],
            "mediaFolders": [{
                "id": "folder-references", "name": "References",
                "assetIds": ["asset-1", "missing-asset", "asset-1"],
            }],
            "tracks": [{
            "id": "framed-video", "kind": "VIDEO", "height": 132, "clips": [{
                "id": "framed-clip", "assetId": "asset-1", "duration": 1000,
                "scale": 0.25, "positionX": 240, "positionY": -130,
            }],
        }]})
        framed_clip = framed["tracks"][0]["clips"][0]
        self.assertEqual(framed["tracks"][0]["height"], 132)
        self.assertEqual(framed_clip["scale"], 0.25)
        self.assertEqual(framed_clip["positionX"], 240)
        self.assertEqual(framed_clip["positionY"], -130)
        self.assertEqual(framed["mediaFolders"], [{
            "id": "folder-references", "name": "References", "assetIds": ["asset-1"],
        }])

        response = self.client.post(
            f"/studio/projects/{self.project.id}/movie-editor/",
            data=json.dumps({
                "title": "Second draft",
                "aspectRatio": "16:9",
                "resolution": "1920x1080",
                "fps": 25,
                "timeline": {"schemaVersion": 1, "tracks": []},
            }),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 200, response.content)
        timeline.refresh_from_db()
        revision = timeline.revisions.first()
        self.assertEqual(revision.title, "Episode draft")

        history = self.client.get(f"/studio/projects/{self.project.id}/movie-editor/revisions/")
        self.assertEqual(history.status_code, 200)
        self.assertGreaterEqual(len(history.json()["items"]), 2)
        restored = self.client.post(
            f"/studio/projects/{self.project.id}/movie-editor/revisions/{revision.id}/restore/",
        )
        self.assertEqual(restored.status_code, 200, restored.content)
        self.assertEqual(restored.json()["title"], "Episode draft")
        self.assertEqual(len(restored.json()["timeline"]["tracks"]), 1)

    def test_movie_editor_rejects_unknown_timeline_schema(self):
        self.client.force_login(self.owner)
        response = self.client.post(
            f"/studio/projects/{self.project.id}/movie-editor/",
            data=json.dumps({"timeline": {"schemaVersion": 99, "tracks": []}}),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn("not supported", response.json()["error"])

    def test_movie_editor_manages_workspace_montage_projects(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        from .models import MovieTimeline

        self.client.force_login(self.owner)
        self.client.get(f"/studio/projects/{self.project.id}/movie-editor/")
        original = MovieTimeline.objects.get(project=self.project)
        self.assertTrue(original.projects.filter(id=self.project.id).exists())

        created = self.client.post(
            f"/studio/projects/{self.project.id}/movie-editor/edits/",
            data=json.dumps({"action": "create", "title": "Second montage"}),
            content_type="application/json",
        )
        self.assertEqual(created.status_code, 201, created.content)
        second = MovieTimeline.objects.get(id=created.json()["id"])
        self.assertEqual(second.workspace, self.workspace)
        self.assertTrue(second.projects.filter(id=self.project.id).exists())

        copied = self.client.post(
            f"/studio/projects/{self.project.id}/movie-editor/edits/",
            data=json.dumps({"action": "copy", "timelineId": str(second.id)}),
            content_type="application/json",
        )
        self.assertEqual(copied.status_code, 201, copied.content)
        copy = MovieTimeline.objects.get(id=copied.json()["id"])
        detached = self.client.post(
            f"/studio/projects/{self.project.id}/movie-editor/edits/",
            data=json.dumps({"action": "detach", "timelineId": str(copy.id)}),
            content_type="application/json",
        )
        self.assertEqual(detached.status_code, 200, detached.content)
        self.assertFalse(copy.projects.filter(id=self.project.id).exists())

        project_page = self.client.get(f"/studio/projects/{self.project.id}/")
        self.assertNotContains(project_page, "Edit projects")
        editor_page = self.client.get(f"/studio/projects/{self.project.id}/movie-editor/")
        self.assertContains(editor_page, "Edit projects")
        self.assertContains(editor_page, "Second montage")
        self.assertContains(editor_page, str(second.id))
        self.assertContains(editor_page, str(copy.id))
        workspace_page = self.client.get(f"/studio/workspaces/{self.workspace.id}/")
        self.assertContains(workspace_page, "Edit projects")
        self.assertContains(workspace_page, str(copy.id))

        exported = self.client.get(
            f"/studio/projects/{self.project.id}/movie-editor/edits/{second.id}/export/",
        )
        self.assertEqual(exported.status_code, 200)
        self.assertEqual(exported.json()["format"], "lexamora-montage-project")

        imported = self.client.post(
            f"/studio/projects/{self.project.id}/movie-editor/edits/",
            {"file": SimpleUploadedFile(
                "imported-montage.json",
                json.dumps({"title": "Imported", "timeline": second.timeline}).encode("utf-8"),
                content_type="application/json",
            )},
        )
        self.assertEqual(imported.status_code, 201, imported.content)
        self.assertTrue(MovieTimeline.objects.get(id=imported.json()["id"]).projects.filter(id=self.project.id).exists())

    def test_movie_editor_round_trips_zip_bundle_with_workspace_media(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        from .models import Asset, MovieTimeline
        from .storage import create_asset

        asset = create_asset(
            user=self.owner,
            workspace=self.workspace,
            project=None,
            uploaded=SimpleUploadedFile("editor-source.mp4", b"video-data", content_type="video/mp4"),
            kind=Asset.Kind.OTHER,
            prevent_duplicate=False,
        )
        Asset.objects.filter(id=asset.id).update(
            duration_ms=5000,
            processing_status=Asset.ProcessingStatus.READY,
        )
        self.client.force_login(self.owner)
        self.client.get(f"/studio/projects/{self.project.id}/movie-editor/")
        timeline = MovieTimeline.objects.get(project=self.project)
        timeline.timeline = {
            "schemaVersion": 1,
            "mediaAssetIds": [str(asset.id)],
            "tracks": [{
                "id": "video-1", "kind": "VIDEO", "muted": False,
                "locked": False, "hidden": False, "height": 84,
                "clips": [{
                    "id": "clip-1", "assetId": str(asset.id),
                    "sourceAssetId": str(asset.id), "start": 0,
                    "sourceStart": 0, "duration": 5000,
                }],
            }],
        }
        timeline.save(update_fields=["timeline", "updated_at"])

        exported = self.client.get(
            f"/studio/projects/{self.project.id}/movie-editor/edits/{timeline.id}/export/?bundle=media",
        )
        self.assertEqual(exported.status_code, 200)
        archive_bytes = b"".join(exported.streaming_content)
        imported = self.client.post(
            f"/studio/projects/{self.project.id}/movie-editor/edits/",
            {"file": SimpleUploadedFile("edit-with-media.zip", archive_bytes, content_type="application/zip")},
        )
        self.assertEqual(imported.status_code, 201, imported.content)
        imported_timeline = MovieTimeline.objects.get(id=imported.json()["id"])
        imported_asset_id = imported_timeline.timeline["tracks"][0]["clips"][0]["assetId"]
        self.assertNotEqual(imported_asset_id, str(asset.id))
        imported_asset = Asset.objects.get(id=imported_asset_id)
        self.assertEqual(imported_asset.workspace, self.workspace)
        self.assertFalse(imported_asset.projects.exists())
        self.assertEqual(imported_timeline.timeline["mediaAssetIds"], [imported_asset_id])

    def test_movie_editor_accepts_media_on_any_organizational_track(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        from .models import Asset
        from .storage import create_asset

        asset = create_asset(
            user=self.owner,
            workspace=self.workspace,
            project=self.project,
            uploaded=SimpleUploadedFile("source.mp4", b"video", content_type="video/mp4"),
            kind=Asset.Kind.OTHER,
        )
        Asset.objects.filter(id=asset.id).update(
            duration_ms=5000,
            processing_status=Asset.ProcessingStatus.READY,
        )
        self.client.force_login(self.owner)
        response = self.client.post(
            f"/studio/projects/{self.project.id}/movie-editor/",
            data=json.dumps({
                "timeline": {
                    "schemaVersion": 1,
                    "tracks": [{
                        "id": "audio-1",
                        "kind": "AUDIO",
                        "clips": [{
                            "id": "clip-1",
                            "assetId": str(asset.id),
                            "start": 0,
                            "sourceStart": 0,
                            "duration": 1000,
                        }],
                    }],
                },
            }),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 200, response.content)

    def test_movie_editor_clamps_trim_to_source_duration(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        from .models import Asset
        from .storage import create_asset

        asset = create_asset(
            user=self.owner,
            workspace=self.workspace,
            project=self.project,
            uploaded=SimpleUploadedFile("trim.mp4", b"video", content_type="video/mp4"),
            kind=Asset.Kind.OTHER,
        )
        Asset.objects.filter(id=asset.id).update(
            duration_ms=5000,
            processing_status=Asset.ProcessingStatus.READY,
        )
        self.client.force_login(self.owner)
        response = self.client.post(
            f"/studio/projects/{self.project.id}/movie-editor/",
            data=json.dumps({
                "timeline": {
                    "schemaVersion": 1,
                    "tracks": [{
                        "id": "video-1",
                        "kind": "VIDEO",
                        "clips": [{
                            "id": "clip-1",
                            "assetId": str(asset.id),
                            "name": "Invalid trim",
                            "start": 0,
                            "sourceStart": 4500,
                            "duration": 1000,
                        }],
                    }],
                },
            }),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 200, response.content)
        self.assertTrue(response.json()["repaired"])
        self.assertEqual(response.json()["timeline"]["tracks"][0]["clips"][0]["duration"], 500)

    def test_movie_editor_upload_queues_media_and_reports_status(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        from .models import Asset

        self.client.force_login(self.owner)
        response = self.client.post(
            f"/studio/projects/{self.project.id}/movie-editor/media/",
            {"files": SimpleUploadedFile("rough-cut.mp4", b"test-video", content_type="video/mp4")},
        )
        self.assertEqual(response.status_code, 201, response.content)
        asset = Asset.objects.get(id=response.json()["items"][0]["id"])
        self.assertEqual(asset.processing_status, Asset.ProcessingStatus.QUEUED)
        self.assertTrue(asset.projects.filter(id=self.project.id).exists())

        status = self.client.get(f"/studio/projects/{self.project.id}/movie-editor/media/")
        self.assertEqual(status.status_code, 200)
        self.assertEqual(len(status.json()["items"]), 1)
        uploaded = next(item for item in status.json()["libraryItems"] if item["id"] == str(asset.id))
        self.assertEqual(uploaded["status"], "QUEUED")
        self.assertEqual(uploaded["kind"], "VIDEO")

    def test_movie_editor_saves_with_warning_for_media_still_processing(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        from .models import Asset
        from .storage import create_asset

        asset = create_asset(
            user=self.owner,
            workspace=self.workspace,
            project=self.project,
            uploaded=SimpleUploadedFile("processing.mp4", b"video", content_type="video/mp4"),
            kind=Asset.Kind.OTHER,
        )
        Asset.objects.filter(id=asset.id).update(
            processing_status=Asset.ProcessingStatus.PROCESSING,
        )
        self.client.force_login(self.owner)
        response = self.client.post(
            f"/studio/projects/{self.project.id}/movie-editor/",
            data=json.dumps({
                "title": "Processing media draft",
                "aspectRatio": "9:16",
                "resolution": "1080x1920",
                "fps": 25,
                "timeline": {
                    "tracks": [{
                        "id": "video-1",
                        "kind": "VIDEO",
                        "clips": [{
                            "id": "processing-clip",
                            "assetId": str(asset.id),
                            "name": "Processing clip",
                            "start": 0,
                            "sourceStart": 0,
                            "duration": 1000,
                        }],
                    }],
                },
            }),
            content_type="application/json",
        )

        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(len(response.json()["warnings"]), 1)
        self.assertIn("not ready", response.json()["warnings"][0])

    def test_direct_project_media_upload_attaches_video(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        from .models import Asset

        self.client.force_login(self.owner)
        response = self.client.post(
            f"/studio/projects/{self.project.id}/media/direct/",
            {"files": SimpleUploadedFile("direct-video.mp4", b"direct-video", content_type="video/mp4")},
        )
        self.assertEqual(response.status_code, 201, response.content)
        asset = Asset.objects.get(id=response.json()["items"][0]["id"])
        self.assertEqual(asset.workspace, self.workspace)
        self.assertTrue(asset.projects.filter(id=self.project.id).exists())
        self.assertEqual(asset.processing_status, Asset.ProcessingStatus.QUEUED)

    def test_direct_project_media_upload_attaches_audio(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        from .models import Asset

        self.client.force_login(self.owner)
        response = self.client.post(
            f"/studio/projects/{self.project.id}/media/direct/",
            {"files": SimpleUploadedFile("dialogue.mp3", b"test-audio", content_type="audio/mpeg")},
        )
        self.assertEqual(response.status_code, 201, response.content)
        asset = Asset.objects.get(id=response.json()["items"][0]["id"])
        self.assertEqual(asset.content_type, "audio/mpeg")
        self.assertEqual(asset.workspace, self.workspace)
        self.assertTrue(asset.projects.filter(id=self.project.id).exists())
        self.assertEqual(asset.processing_status, Asset.ProcessingStatus.QUEUED)

    def test_workspace_and_project_pages_expose_direct_media_upload(self):
        self.client.force_login(self.owner)
        workspace_response = self.client.get(f"/studio/workspaces/{self.workspace.id}/")
        project_response = self.client.get(f"/studio/projects/{self.project.id}/")

        self.assertEqual(workspace_response.status_code, 200)
        self.assertEqual(project_response.status_code, 200)
        self.assertContains(workspace_response, f"/studio/workspaces/{self.workspace.id}/media/direct/")
        self.assertContains(project_response, f"/studio/projects/{self.project.id}/media/direct/")
        self.assertContains(workspace_response, "Add Video")
        self.assertContains(project_response, "Add Video")
        self.assertContains(workspace_response, "Add Audio")
        self.assertContains(project_response, "Add Audio")

    def test_direct_project_upload_reuses_matching_workspace_media(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        from .models import Asset

        self.client.force_login(self.owner)
        workspace_response = self.client.post(
            f"/studio/workspaces/{self.workspace.id}/media/direct/",
            {"files": self.image_file("shared-photo.png")},
        )
        self.assertEqual(workspace_response.status_code, 201, workspace_response.content)
        asset_id = workspace_response.json()["items"][0]["id"]
        before = Asset.objects.filter(workspace=self.workspace).count()
        project_response = self.client.post(
            f"/studio/projects/{self.project.id}/media/direct/",
            {"files": self.image_file("shared-photo.png")},
        )
        self.assertEqual(project_response.status_code, 201, project_response.content)
        self.assertEqual(project_response.json()["items"][0]["id"], asset_id)
        self.assertEqual(Asset.objects.filter(workspace=self.workspace).count(), before)
        self.assertTrue(Asset.objects.get(id=asset_id).projects.filter(id=self.project.id).exists())

    def test_movie_editor_can_attach_accessible_workspace_media(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        from .models import Asset
        from .storage import create_asset

        asset = create_asset(
            user=self.owner, workspace=self.workspace,
            uploaded=SimpleUploadedFile("workspace-video.mp4", b"video", content_type="video/mp4"),
            kind=Asset.Kind.OTHER,
        )
        Asset.objects.filter(id=asset.id).update(
            duration_ms=4000, processing_status=Asset.ProcessingStatus.READY,
        )
        self.client.force_login(self.owner)
        before = self.client.get(f"/studio/projects/{self.project.id}/movie-editor/media/")
        self.assertEqual(before.status_code, 200)
        self.assertFalse(any(item["id"] == str(asset.id) for item in before.json()["items"]))
        self.assertTrue(any(item["id"] == str(asset.id) for item in before.json()["libraryItems"]))

        attached = self.client.post(
            f"/studio/projects/{self.project.id}/movie-editor/media/",
            data=json.dumps({"assetId": str(asset.id)}),
            content_type="application/json",
        )
        self.assertEqual(attached.status_code, 201, attached.content)
        self.assertTrue(attached.json()["attached"])
        self.assertTrue(asset.projects.filter(id=self.project.id).exists())

    def test_movie_editor_can_retry_failed_media(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        from .models import Asset
        from .storage import create_asset

        asset = create_asset(
            user=self.owner, workspace=self.workspace,
            uploaded=SimpleUploadedFile("retry.mp3", b"test-audio", content_type="audio/mpeg"),
            kind=Asset.Kind.OTHER, project=self.project,
        )
        Asset.objects.filter(id=asset.id).update(
            processing_status=Asset.ProcessingStatus.FAILED, processing_error="Broken source",
        )
        self.client.force_login(self.owner)
        response = self.client.post(
            f"/studio/projects/{self.project.id}/movie-editor/media/{asset.id}/retry/",
        )
        self.assertEqual(response.status_code, 202, response.content)
        asset.refresh_from_db()
        self.assertEqual(asset.processing_status, Asset.ProcessingStatus.QUEUED)
        self.assertEqual(asset.processing_error, "")

    def test_movie_editor_queues_cancels_and_retries_rough_cut_render(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        from .models import Asset, MovieRenderJob, MovieTimeline
        from .storage import create_asset

        asset = create_asset(
            user=self.owner, workspace=self.workspace,
            uploaded=SimpleUploadedFile("render-source.mp4", b"video", content_type="video/mp4"),
            kind=Asset.Kind.OTHER, project=self.project,
        )
        Asset.objects.filter(id=asset.id).update(
            duration_ms=8000,
            media_metadata={"video": {"codec": "h264"}, "audio": {"codec": "aac"}},
            processing_status=Asset.ProcessingStatus.READY,
        )
        self.client.force_login(self.owner)
        timeline_payload = {
            "schemaVersion": 1,
            "tracks": [{
                "id": "video-1", "kind": "VIDEO", "clips": [{
                    "id": "clip-1", "assetId": str(asset.id), "name": "Opening",
                    "start": 500, "sourceStart": 1000, "duration": 4000, "volume": 0.8,
                }],
            }],
        }
        saved = self.client.post(
            f"/studio/projects/{self.project.id}/movie-editor/",
            data=json.dumps({"title": "Episode render", "aspectRatio": "9:16", "fps": 25, "timeline": timeline_payload}),
            content_type="application/json",
        )
        self.assertEqual(saved.status_code, 200, saved.content)
        response = self.client.post(
            f"/studio/projects/{self.project.id}/movie-editor/renders/",
            data=json.dumps({
                "profile": "DRAFT_720", "audioProfile": "BALANCED", "targetLufs": -14,
            }), content_type="application/json",
        )
        self.assertEqual(response.status_code, 202, response.content)
        job = MovieRenderJob.objects.get(id=response.json()["id"])
        self.assertEqual((job.width, job.height), (720, 1280))
        self.assertEqual(job.duration_ms, 4500)
        self.assertEqual(job.audio_profile, MovieRenderJob.AudioProfile.BALANCED)
        self.assertEqual(job.target_lufs, -14)
        self.assertEqual(response.json()["audioProfileLabel"], "Balanced")
        self.assertEqual(job.snapshot, MovieTimeline.objects.get(project=self.project).timeline)

        clip_render = self.client.post(
            f"/studio/projects/{self.project.id}/movie-editor/renders/",
            data=json.dumps({
                "timelineId": str(job.timeline_id), "clipIds": ["clip-1"],
                "title": "Opening media", "profile": "DRAFT_720",
            }), content_type="application/json",
        )
        self.assertEqual(clip_render.status_code, 202, clip_render.content)
        clip_job = MovieRenderJob.objects.get(id=clip_render.json()["id"])
        self.assertEqual(clip_job.duration_ms, 4000)
        self.assertEqual(clip_job.snapshot["tracks"][0]["clips"][0]["start"], 0)

        cancelled = self.client.post(
            f"/studio/projects/{self.project.id}/movie-editor/renders/{job.id}/",
            data=json.dumps({"action": "cancel"}), content_type="application/json",
        )
        self.assertEqual(cancelled.status_code, 202, cancelled.content)
        job.refresh_from_db()
        self.assertEqual(job.status, MovieRenderJob.Status.CANCELLED)
        retried = self.client.post(
            f"/studio/projects/{self.project.id}/movie-editor/renders/{job.id}/",
            data=json.dumps({"action": "retry"}), content_type="application/json",
        )
        self.assertEqual(retried.status_code, 202, retried.content)
        job.refresh_from_db()
        self.assertEqual(job.status, MovieRenderJob.Status.QUEUED)
        self.assertFalse(job.cancel_requested)

    def test_movie_render_command_trims_composes_and_mixes_timeline(self):
        from types import SimpleNamespace
        from .movie_rendering import build_render_command

        video = SimpleNamespace(
            id="asset-1", proxy_file=SimpleNamespace(path="/tmp/video.mp4"),
            file=SimpleNamespace(path="/tmp/original.mp4"),
            processing_status="READY", media_metadata={"audio": {"codec": "aac"}},
        )
        job = SimpleNamespace(
            snapshot={"tracks": [{"kind": "VIDEO", "muted": False, "clips": [{
                "assetId": "asset-1", "start": 1200, "sourceStart": 500,
                "duration": 3000, "volume": 0.75,
                "scale": 1.35, "positionX": 60, "positionY": -25,
            }]}]},
            duration_ms=4200, width=1280, height=720, fps=25,
        )
        command = build_render_command(job, {"asset-1": video}, "/tmp/output.mp4")
        joined = " ".join(command)
        self.assertIn("trim=start=0.500:duration=3.000", joined)
        self.assertIn("overlay=eof_action=pass", joined)
        self.assertIn("ceil(iw*1.3500/2)*2", joined)
        self.assertIn("ceil(ih*1.3500/2)*2", joined)
        self.assertIn("(W-w)/2+60.000", joined)
        self.assertIn("(H-h)/2+25.000", joined)
        self.assertIn("volume='if(lt(t,3.000000),0.750000", joined)
        self.assertIn("highpass=f=80", joined)
        self.assertIn("afftdn=nf=-25:tn=1", joined)
        self.assertIn("acompressor=threshold=0.125", joined)
        self.assertIn("amix=inputs=1", joined)
        self.assertIn("loudnorm=I=-16:TP=-1.5:LRA=11", joined)
        self.assertIn("libx264", command)

        job.snapshot["tracks"][0]["clips"][0].update({
            "speed": 2, "speedMethod": "FRAME_BLEND", "fadeIn": 400, "fadeOut": 600,
            "volumeKeyframes": [{"time": 1500, "value": 0.25}],
        })
        adjusted = " ".join(build_render_command(job, {"asset-1": video}, "/tmp/adjusted.mp4"))
        self.assertIn("trim=start=0.500:duration=6.000", adjusted)
        self.assertIn("setpts=(PTS-STARTPTS)/2.000000", adjusted)
        self.assertIn("minterpolate=fps=25:mi_mode=blend", adjusted)
        self.assertIn("atempo=2.000000", adjusted)
        self.assertIn("afade=t=in:st=0:d=0.400", adjusted)
        self.assertIn("afade=t=out:st=2.400:d=0.600", adjusted)
        self.assertIn("0.250000", adjusted)

        job.audio_profile = "ORIGINAL"
        original = " ".join(build_render_command(job, {"asset-1": video}, "/tmp/original-output.mp4"))
        self.assertNotIn("afftdn=", original)
        self.assertNotIn("loudnorm=", original)

    def test_movie_render_accepts_video_on_an_organizational_audio_lane(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        from .models import Asset
        from .storage import create_asset

        asset = create_asset(
            user=self.owner, workspace=self.workspace,
            uploaded=SimpleUploadedFile("lane-video.mp4", b"video", content_type="video/mp4"),
            kind=Asset.Kind.OTHER, project=self.project,
        )
        Asset.objects.filter(id=asset.id).update(
            duration_ms=2000, media_metadata={"video": {"codec": "h264"}, "audio": {"codec": "aac"}},
            processing_status=Asset.ProcessingStatus.READY,
        )
        self.client.force_login(self.owner)
        saved = self.client.post(
            f"/studio/projects/{self.project.id}/movie-editor/",
            data=json.dumps({"timeline": {"schemaVersion": 1, "tracks": [{
                "id": "free-lane", "kind": "AUDIO", "clips": [{
                    "id": "video-clip", "assetId": str(asset.id), "name": "Video on a free lane",
                    "start": 0, "sourceStart": 0, "duration": 1500, "volume": 1,
                }],
            }]}}), content_type="application/json",
        )
        self.assertEqual(saved.status_code, 200, saved.content)
        queued = self.client.post(
            f"/studio/projects/{self.project.id}/movie-editor/renders/",
            data=json.dumps({"profile": "DRAFT_720"}), content_type="application/json",
        )
        self.assertEqual(queued.status_code, 202, queued.content)

    def test_movie_render_mixes_audio_from_all_overlapping_media(self):
        from types import SimpleNamespace
        from .movie_rendering import build_render_command, timeline_duration_ms

        lower = SimpleNamespace(
            id="lower", content_type="video/mp4", proxy_file=SimpleNamespace(path="/tmp/lower.mp4"),
            file=SimpleNamespace(path="/tmp/lower.mp4"), processing_status="READY",
            media_metadata={"video": {"codec": "h264"}, "audio": {"codec": "aac"}},
        )
        upper = SimpleNamespace(
            id="upper", content_type="video/mp4", proxy_file=SimpleNamespace(path="/tmp/upper.mp4"),
            file=SimpleNamespace(path="/tmp/upper.mp4"), processing_status="READY",
            media_metadata={"video": {"codec": "h264"}, "audio": {"codec": "aac"}},
        )
        audio = SimpleNamespace(
            id="music", content_type="audio/mpeg", proxy_file=SimpleNamespace(path="/tmp/music.mp3"),
            file=SimpleNamespace(path="/tmp/music.mp3"), processing_status="READY",
            media_metadata={"audio": {"codec": "mp3"}},
        )
        snapshot = {"tracks": [
            {"kind": "VIDEO", "muted": False, "clips": [{
                "assetId": "upper", "start": 1000, "sourceStart": 0, "duration": 2000, "volume": 1,
            }]},
            {"kind": "AUDIO", "muted": False, "clips": [{
                "assetId": "lower", "start": 0, "sourceStart": 0, "duration": 4000, "volume": 1,
            }, {"assetId": "music", "start": 0, "sourceStart": 0, "duration": 6000, "volume": 1}]},
        ]}
        job = SimpleNamespace(snapshot=snapshot, duration_ms=4000, width=1280, height=720, fps=25)
        render_command = build_render_command(
            job, {"lower": lower, "upper": upper, "music": audio}, "/tmp/visible.mp4",
        )
        command = " ".join(render_command)
        filter_graph = render_command[render_command.index("-filter_complex") + 1]
        self.assertEqual(render_command.count("-i"), 3)
        self.assertIn("[0:v]", filter_graph)
        self.assertIn("[1:v]", filter_graph)
        self.assertIn("[2:a]", filter_graph)
        self.assertNotIn("[4:v]", filter_graph)
        self.assertIn("atrim=start=0.000:duration=4.000", command)
        self.assertIn("atrim=start=0.000:duration=2.000", command)
        self.assertIn("amix=inputs=3", filter_graph)
        self.assertEqual(timeline_duration_ms(snapshot, {"lower": lower, "upper": upper, "music": audio}), 4000)

        snapshot["tracks"][0]["hidden"] = True
        hidden_command = build_render_command(
            job, {"lower": lower, "upper": upper, "music": audio}, "/tmp/hidden.mp4",
        )
        hidden_graph = hidden_command[hidden_command.index("-filter_complex") + 1]
        self.assertEqual(hidden_command.count("-i"), 2)
        self.assertIn("amix=inputs=2", hidden_graph)

    def test_movie_render_worker_persists_completed_mp4_asset(self):
        import io
        from pathlib import Path
        from django.core.files.uploadedfile import SimpleUploadedFile
        from .models import Asset, MovieRenderJob, MovieTimeline
        from .movie_rendering import process_movie_render
        from .storage import create_asset

        source = create_asset(
            user=self.owner, workspace=self.workspace,
            uploaded=SimpleUploadedFile("worker-source.mp4", b"source", content_type="video/mp4"),
            kind=Asset.Kind.OTHER, project=self.project,
        )
        Asset.objects.filter(id=source.id).update(
            duration_ms=2000, media_metadata={"video": {"codec": "h264"}},
            processing_status=Asset.ProcessingStatus.READY,
        )
        timeline = MovieTimeline.objects.create(
            project=self.project, title="Worker render", aspect_ratio="16:9",
            resolution="1280x720", fps=25, created_by=self.owner, updated_by=self.owner,
            timeline={"schemaVersion": 1, "tracks": [{
                "id": "video-1", "kind": "VIDEO", "muted": False, "clips": [{
                    "id": "clip-1", "assetId": str(source.id), "name": "Source",
                    "start": 0, "sourceStart": 0, "duration": 1000, "volume": 1,
                }],
            }]},
        )
        job = MovieRenderJob.objects.create(
            timeline=timeline, project=self.project, title="Worker render",
            profile=MovieRenderJob.Profile.DRAFT_720, aspect_ratio="16:9",
            width=1280, height=720, fps=25, duration_ms=1000,
            snapshot=timeline.timeline, requested_by=self.owner,
        )

        class FakeProcess:
            def __init__(self, command):
                Path(command[-1]).write_bytes(b"rendered-mp4")
                self.stdout = iter(["out_time_ms=1000000\n", "progress=end\n"])
                self.stderr = io.StringIO("")

            def wait(self, timeout=None):
                return 0

            def poll(self):
                return 0

            def kill(self):
                return None

        with patch("lexamora_studio.movie_rendering.subprocess.Popen", side_effect=lambda command, **_kwargs: FakeProcess(command)):
            process_movie_render(job.id)
        job.refresh_from_db()
        self.assertEqual(job.status, MovieRenderJob.Status.SUCCEEDED)
        self.assertEqual(job.progress, 100)
        self.assertIsNotNone(job.output_asset_id)
        self.assertEqual(job.output_asset.kind, Asset.Kind.EXPORT)
        self.assertEqual(job.output_asset.content_type, "video/mp4")
        self.assertEqual(job.output_asset.media_metadata["audio"]["cleanupProfile"], "CLEAN_SPEECH")
        self.assertEqual(job.output_asset.media_metadata["audio"]["targetLufs"], -16)
        self.assertFalse(job.output_asset.projects.filter(id=self.project.id).exists())
        self.client.force_login(self.owner)
        attached = self.client.post(
            f"/studio/projects/{self.project.id}/movie-editor/renders/{job.id}/",
            data=json.dumps({"action": "attach"}), content_type="application/json",
        )
        self.assertEqual(attached.status_code, 202, attached.content)
        self.assertTrue(attached.json()["attachedToProject"])
        self.assertTrue(job.output_asset.projects.filter(id=self.project.id).exists())
        if job.output_asset.file.storage.exists(job.output_asset.file.name):
            job.output_asset.file.storage.delete(job.output_asset.file.name)

    def test_media_metadata_extracts_video_audio_and_duration(self):
        from .media_processing import _metadata

        metadata, duration_ms, video, audio = _metadata({
            "format": {"format_name": "mov,mp4", "duration": "12.345", "bit_rate": "900000"},
            "streams": [
                {"codec_type": "video", "codec_name": "h264", "width": 1080, "height": 1920, "avg_frame_rate": "25/1", "pix_fmt": "yuv420p"},
                {"codec_type": "audio", "codec_name": "aac", "channels": 2, "sample_rate": "48000"},
            ],
        })
        self.assertEqual(duration_ms, 12345)
        self.assertEqual(metadata["video"]["frameRate"], 25.0)
        self.assertEqual(metadata["audio"]["sampleRate"], 48000)
        self.assertIsNotNone(video)
        self.assertIsNotNone(audio)

    def test_media_processor_persists_proxy_thumbnail_and_waveform(self):
        from pathlib import Path
        from django.core.files.uploadedfile import SimpleUploadedFile
        from .media_processing import process_media_asset
        from .models import Asset
        from .storage import create_asset

        asset = create_asset(
            user=self.owner, workspace=self.workspace,
            uploaded=SimpleUploadedFile("processable.mp4", b"fake-source", content_type="video/mp4"),
            kind=Asset.Kind.OTHER, project=self.project,
        )
        probe = {
            "format": {"format_name": "mov,mp4", "duration": "3.2", "bit_rate": "700000"},
            "streams": [
                {"codec_type": "video", "codec_name": "h264", "width": 720, "height": 1280, "avg_frame_rate": "25/1"},
                {"codec_type": "audio", "codec_name": "aac", "channels": 2, "sample_rate": "48000"},
            ],
        }

        def write_artifact(_source, target, *_args):
            Path(target).write_bytes(b"generated-artifact")

        try:
            with patch("lexamora_studio.media_processing._probe", return_value=probe), \
                    patch("lexamora_studio.media_processing._create_video_proxy", side_effect=write_artifact), \
                    patch("lexamora_studio.media_processing._create_thumbnail", side_effect=write_artifact), \
                    patch("lexamora_studio.media_processing._create_filmstrip", side_effect=write_artifact), \
                    patch("lexamora_studio.media_processing._create_waveform", side_effect=write_artifact):
                process_media_asset(asset.id)
            asset.refresh_from_db()
            self.assertEqual(asset.processing_status, Asset.ProcessingStatus.READY)
            self.assertEqual(asset.duration_ms, 3200)
            self.assertTrue(asset.proxy_file.name)
            self.assertTrue(asset.thumbnail.name)
            self.assertTrue(asset.filmstrip_file.name)
            self.assertTrue(asset.waveform_file.name)

            self.client.force_login(self.owner)
            for endpoint in ("proxy", "waveform", "filmstrip"):
                response = self.client.get(f"/api/v1/studio/assets/{asset.id}/{endpoint}")
                self.assertEqual(response.status_code, 200)
                list(response.streaming_content)
                response.close()
        finally:
            asset.refresh_from_db()
            for field in (asset.file, asset.proxy_file, asset.thumbnail, asset.filmstrip_file, asset.waveform_file):
                if field.name and field.storage.exists(field.name):
                    field.storage.delete(field.name)

    def test_movie_editor_2_is_available_to_project_users(self):
        self.client.force_login(self.owner)
        response = self.client.get(f"/studio/projects/{self.project.id}/movie-editor-2/")
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Movie Editor 2")
        self.assertContains(response, "studio/movie-editor-2/index.html")

    @patch("lexamora_studio.views.user_has_ai_access", return_value=True)
    def test_project_generation_can_queue_selected_model(self, _has_access):
        from .models import ImageGenerationJob

        self.client.force_login(self.owner)
        response = self.client.post(
            f"/studio/projects/{self.project.id}/image-generation/queue/",
            data=json.dumps({
                "prompt": "A cinematic test frame",
                "modelProfileId": str(self.image_model.id),
                "size": "1024x1536",
                "quality": "medium",
                "outputFormat": "png",
                "referenceAssetIds": [],
            }),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 202, response.content)
        job = ImageGenerationJob.objects.get(id=response.json()["jobId"])
        self.assertEqual(job.project, self.project)
        self.assertIsNone(job.prompt)
        self.assertEqual(job.model_profile, self.image_model)

    @patch("lexamora_studio.views.user_has_ai_access", return_value=True)
    def test_project_generation_accepts_multiple_jobs_without_waiting(self, _has_access):
        from .models import ImageGenerationJob

        self.client.force_login(self.owner)
        payload = {
            "prompt": "First independent frame",
            "modelProfileId": str(self.image_model.id),
            "size": "1024x1024",
            "quality": "low",
            "outputFormat": "png",
            "referenceAssetIds": [],
        }
        first = self.client.post(
            f"/studio/projects/{self.project.id}/image-generation/queue/",
            data=json.dumps(payload),
            content_type="application/json",
        )
        payload["prompt"] = "Second independent frame"
        second = self.client.post(
            f"/studio/projects/{self.project.id}/image-generation/queue/",
            data=json.dumps(payload),
            content_type="application/json",
        )
        self.assertEqual(first.status_code, 202)
        self.assertEqual(second.status_code, 202)
        self.assertNotEqual(first.json()["jobId"], second.json()["jobId"])
        self.assertEqual(
            ImageGenerationJob.objects.filter(project=self.project, status=ImageGenerationJob.Status.QUEUED).count(),
            2,
        )

    def test_generated_asset_can_be_starred_and_detached_from_generation_page(self):
        import tempfile
        from pathlib import Path
        from .models import Asset, ImageGenerationJob
        from .storage import create_asset

        self.client.force_login(self.owner)
        with tempfile.TemporaryDirectory() as directory:
            with self.settings(STUDIO_PRIVATE_MEDIA_ROOT=Path(directory)):
                generated = create_asset(
                    user=self.owner,
                    workspace=self.workspace,
                    project=self.project,
                    uploaded=self.image_file("generated-result.png"),
                    kind=Asset.Kind.GENERATION_OUTPUT,
                )
                generated.projects.add(self.project)
                job = ImageGenerationJob.objects.create(
                    workspace=self.workspace,
                    project=self.project,
                    requested_by=self.owner,
                    model_profile=self.image_model,
                    status=ImageGenerationJob.Status.SUCCESS,
                    request_prompt="Generated result",
                    options={"size": "1024x1024", "quality": "low", "output_format": "png"},
                    result_asset=generated,
                )
                listed = self.client.get(f"/studio/projects/{self.project.id}/image-generation/jobs/")
                self.assertEqual([item["jobId"] for item in listed.json()["jobs"]], [str(job.id)])
                self.assertTrue(listed.json()["jobs"][0]["downloadUrl"])

                starred = self.client.post(f"/studio/assets/{generated.id}/star/")
                self.assertEqual(starred.status_code, 200, starred.content)
                self.assertTrue(starred.json()["starred"])

                detached = self.client.post(
                    f"/studio/assets/{generated.id}/detach/project/{self.project.id}/",
                    HTTP_X_REQUESTED_WITH="XMLHttpRequest",
                )
                self.assertEqual(detached.status_code, 302, detached.content)
                generated.refresh_from_db()
                self.assertEqual(generated.kind, Asset.Kind.GENERATION_OUTPUT)
                self.assertTrue(generated.is_starred)
                self.assertFalse(generated.projects.filter(id=self.project.id).exists())
                listed = self.client.get(f"/studio/projects/{self.project.id}/image-generation/jobs/")
                self.assertEqual(listed.json()["jobs"], [])

    @patch("lexamora_studio.views.user_has_ai_access", return_value=True)
    def test_project_generation_page_renders_queue_workspace(self, _has_access):
        self.client.force_login(self.owner)
        response = self.client.get(
            f"/studio/projects/{self.project.id}/image-generation/"
        )
        self.assertEqual(response.status_code, 200, response.content)
        self.assertContains(response, "data-reference-slots")
        self.assertContains(response, "Generated here")
        self.assertContains(response, "data-project-generation-results")
        self.assertContains(response, "data-project-generation-send-direct")
        self.assertContains(response, "Generate Image")
        self.assertNotContains(response, "Review request")
        self.assertContains(response, "data-source-dialog")
        self.assertContains(response, "9:16 Story")
        self.assertContains(response, "data-project-prompt-action=\"translate\"")
        self.assertContains(response, "Final prompt")
        self.assertContains(response, "data-avatar-crop-dialog")
        self.assertContains(response, "data-generation-composer")
        self.assertContains(response, "data-project-media-type")
        self.assertContains(response, '<option value="VIDEO">Video</option>', html=True)
        self.assertContains(response, "data-project-generation-quantity")
        self.assertContains(response, "data-max-outputs")
        self.assertContains(response, "data-generation-key=\"video-size\"")
        self.assertContains(response, "Generated photos")
        self.assertContains(response, "Generated videos")
        self.assertContains(response, "data-video-preview-dialog")
        self.assertContains(response, 'data-library-type="VIDEO"')
        self.assertContains(response, "data-composer-undo")
        self.assertContains(response, "data-composer-redo")
        self.assertContains(response, "data-generation-reference-filter")
        self.assertNotContains(response, "data-project-prompt-target")

    @patch("lexamora_studio.views.run_text", return_value=("Improved cinematic prompt", "gpt-5.4-mini"))
    @patch("lexamora_studio.views.user_has_ai_access", return_value=True)
    def test_project_image_prompt_can_be_improved_in_preview(self, _has_access, _run_text):
        from .models import StudioTextModel

        StudioTextModel.objects.update_or_create(
            model_id="gpt-5.4-mini",
            defaults={"name": "GPT mini", "is_active": True, "is_default": True, "updated_by": self.admin},
        )
        self.client.force_login(self.owner)
        response = self.client.post(
            f"/studio/projects/{self.project.id}/image-generation/prompt-preview/",
            data=json.dumps({"prompt": "rough frame", "action": "improve", "textModel": "gpt-5.4-mini"}),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(response.json()["content"], "Improved cinematic prompt")

    def test_scene_prompt_expands_inline_and_can_be_deleted(self):
        from .ai_catalog import default_prompt_template
        from .models import Episode, Scene

        episode = Episode.objects.create(
            project=self.project, number=1, title="Pilot", created_by=self.owner, updated_by=self.owner,
        )
        scene = Scene.objects.create(
            episode=episode, number=1, position=0, title="Opening", created_by=self.owner, updated_by=self.owner,
        )
        prompt = Prompt.objects.create(
            scene=scene, ai_model=self.image_model, template=default_prompt_template(Prompt.Type.IMAGE),
            prompt_type=Prompt.Type.IMAGE, title="Street frame", content="A hero walks outside",
            created_by=self.owner, updated_by=self.owner,
        )
        self.client.force_login(self.owner)
        page = self.client.get(f"/studio/scenes/{scene.id}/")
        self.assertContains(page, f'<details class="scene-prompt-row" id="prompt-row-{prompt.id}">')
        self.assertNotContains(page, f'id="prompt-dialog-{prompt.id}"')
        self.assertContains(page, "scene-prompt-inline")
        self.assertContains(page, "A hero walks outside")
        deleted = self.client.post(f"/studio/prompts/{prompt.id}/delete/")
        self.assertRedirects(deleted, f"/studio/scenes/{scene.id}/#prompts")
        self.assertFalse(Prompt.objects.filter(id=prompt.id).exists())

    def test_external_image_import_rejects_private_network(self):
        from django.core.exceptions import ValidationError
        from .external_images import download_external_image

        with self.assertRaises(ValidationError):
            download_external_image("http://127.0.0.1/private.png")


class StudioTokenAssistantComicTests(TestCase):
    def setUp(self):
        import tempfile

        from .models import Episode, StudioTextModel, private_storage

        self.temp_media = tempfile.TemporaryDirectory()
        self.private_storage = private_storage
        self.previous_location = private_storage._location
        private_storage._location = self.temp_media.name
        private_storage.__dict__.pop("base_location", None)
        private_storage.__dict__.pop("location", None)
        users = get_user_model()
        self.owner = users.objects.create_user("quota-owner", password="strong-pass")
        self.workspace = create_workspace(user=self.owner, name="Quota Studio", slug="quota-studio")
        self.text_model, _ = StudioTextModel.objects.get_or_create(
            model_id="gpt-5.4-mini",
            defaults={"name": "GPT Test Mini", "is_active": True, "is_default": True},
        )
        self.project = Project.objects.create(
            workspace=self.workspace,
            project_type=Project.Type.SERIES,
            title="Token Project",
            documentation_language="RU",
            dialogue_language="PL",
            prompt_language="EN",
            default_translation_model=self.text_model,
            created_by=self.owner,
            updated_by=self.owner,
        )
        self.episode = Episode.objects.create(
            project=self.project,
            number=1,
            title="Pilot",
            created_by=self.owner,
            updated_by=self.owner,
        )
        self.client.force_login(self.owner)

    def tearDown(self):
        self.private_storage._location = self.previous_location
        self.private_storage.__dict__.pop("base_location", None)
        self.private_storage.__dict__.pop("location", None)
        self.temp_media.cleanup()

    def test_token_usage_defaults_to_million_and_paginates(self):
        from .models import AiUsageLog

        AiUsageLog.objects.bulk_create([
            AiUsageLog(
                workspace=self.workspace,
                user=self.owner,
                action="TRANSLATE" if index % 2 else "ASSIST",
                model="gpt-test-mini",
                status="SUCCESS",
                total_tokens=10,
            )
            for index in range(55)
        ])
        response = self.client.get("/studio/usage/me/?per_page=20&action=TRANSLATE")
        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload["summary"]["allowance"], 1_000_000)
        self.assertEqual(payload["summary"]["remaining"], 999_450)
        self.assertEqual(payload["count"], 27)
        self.assertEqual(len(payload["operations"]), 20)
        self.assertEqual(payload["pages"], 2)

    def test_workspace_and_project_avatar_removal_preserve_asset(self):
        from django.core.files.uploadedfile import SimpleUploadedFile

        from .models import Asset
        from .storage import create_asset

        asset = create_asset(
            user=self.owner,
            workspace=self.workspace,
            project=self.project,
            kind=Asset.Kind.OTHER,
            uploaded=SimpleUploadedFile("avatar.txt", b"avatar", content_type="text/plain"),
        )
        self.workspace.avatar_asset = asset
        self.workspace.save(update_fields=["avatar_asset", "updated_at"])
        self.project.cover_asset = asset
        self.project.save(update_fields=["cover_asset", "updated_at"])
        workspace_response = self.client.post(
            f"/studio/workspaces/{self.workspace.id}/avatar/remove/",
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )
        project_response = self.client.post(
            f"/studio/projects/{self.project.id}/cover/remove/",
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )
        self.assertEqual((workspace_response.status_code, project_response.status_code), (200, 200))
        self.workspace.refresh_from_db()
        self.project.refresh_from_db()
        self.assertIsNone(self.workspace.avatar_asset_id)
        self.assertIsNone(self.project.cover_asset_id)
        self.assertTrue(Asset.objects.filter(id=asset.id).exists())

    @patch("lexamora_studio.views.user_has_ai_access", return_value=True)
    @patch("lexamora_studio.views._run_logged_text", return_value=("translated", "gpt-5.4-mini"))
    def test_prompt_translation_preserves_named_dialogue_language(self, mocked_run, _mocked_access):
        response = self.client.post(
            f"/studio/projects/{self.project.id}/image-generation/prompt-preview/",
            data=json.dumps({
                "prompt": 'Character speaks in Polish: "I am healthy"',
                "action": "translate",
                "targetLanguage": "EN",
                "textModel": "gpt-5.4-mini",
            }),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 200)
        instruction = mocked_run.call_args.kwargs["text"]
        self.assertIn("translate that direct speech into the named language", instruction)
        self.assertIn("Character speaks in Polish", instruction)

    def test_assistant_context_merges_and_creates_downloadable_asset(self):
        first = self.client.post(
            f"/studio/projects/{self.project.id}/assistant/context/",
            {"content": "First context", "merge": "1"},
        )
        second = self.client.post(
            f"/studio/projects/{self.project.id}/assistant/context/",
            {"content": "Second context", "merge": "1"},
        )
        self.assertEqual((first.status_code, second.status_code), (200, 200))
        self.assertIn("First context", second.json()["content"])
        self.assertIn("Second context", second.json()["content"])
        self.assertTrue(second.json()["downloadUrl"])

    @patch("lexamora_studio.views._run_logged_text", return_value=("Panel plan", "gpt-5.4-mini"))
    def test_episode_comic_is_created_as_workspace_project_asset(self, _mocked_run):
        from .models import EpisodeComic

        response = self.client.post(
            f"/studio/episodes/{self.episode.id}/comic/generate/",
            {"model": "gpt-5.4-mini"},
        )
        self.assertRedirects(response, f"/studio/episodes/{self.episode.id}/edit/")
        comic = EpisodeComic.objects.get(episode=self.episode)
        self.assertEqual(comic.status, EpisodeComic.Status.SUCCESS)
        self.assertEqual(comic.asset.content_type, "application/pdf")
        self.assertTrue(comic.asset.projects.filter(id=self.project.id).exists())

    @patch("lexamora_studio.views._run_logged_text", return_value=("Continuity report", "gpt-5.4-mini"))
    def test_episode_consistency_review_is_saved_without_images(self, mocked_run):
        from .models import EpisodeConsistencyReview

        response = self.client.post(
            f"/studio/episodes/{self.episode.id}/consistency-reviews/",
            {"model": "gpt-5.4-mini"},
        )

        self.assertEqual(response.status_code, 302)
        self.assertTrue(response["Location"].endswith("#consistency-reviews"))
        review = EpisodeConsistencyReview.objects.get(episode=self.episode)
        self.assertEqual(review.content, "Continuity report")
        self.assertEqual(review.image_count, 0)
        self.assertEqual(mocked_run.call_args.kwargs["action"], "EPISODE_CONSISTENCY_REVIEW")
        self.assertIn(self.project.documentation_language, mocked_run.call_args.kwargs["text"])

    @patch("lexamora_studio.views._run_logged_multimodal_text", return_value=("Visual panel plan", "gpt-5.4-mini"))
    def test_episode_comic_sends_scene_images_to_multimodal_model(self, mocked_run):
        import io

        from PIL import Image
        from django.core.files.uploadedfile import SimpleUploadedFile

        from .models import Asset, Scene
        from .storage import create_asset

        scene = Scene.objects.create(
            episode=self.episode,
            number=1,
            position=1,
            title="Visual scene",
            description="A visible reference must guide the panel",
            created_by=self.owner,
            updated_by=self.owner,
        )
        image_buffer = io.BytesIO()
        Image.new("RGB", (64, 96), "green").save(image_buffer, format="PNG")
        create_asset(
            user=self.owner,
            workspace=self.workspace,
            project=self.project,
            scene=scene,
            kind=Asset.Kind.SCENE_IMAGE,
            uploaded=SimpleUploadedFile("scene-reference.png", image_buffer.getvalue(), content_type="image/png"),
        )

        response = self.client.post(
            f"/studio/episodes/{self.episode.id}/comic/generate/",
            {"model": "gpt-5.4-mini"},
        )

        self.assertRedirects(response, f"/studio/episodes/{self.episode.id}/edit/")
        references = mocked_run.call_args.kwargs["reference_images"]
        self.assertEqual(len(references), 1)
        self.assertEqual(references[0][0], "scene-reference.png")
        self.assertIn("reference image 1", mocked_run.call_args.kwargs["text"])
