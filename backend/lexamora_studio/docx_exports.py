import io

from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import transaction
from django.utils import timezone
from docx import Document
from docx.shared import Inches

from .models import Asset, DocxImport, ExportJob
from .revisions import audit
from .storage import create_asset

DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _clear_body(document):
    body = document._element.body
    for child in list(body):
        if not child.tag.endswith("}sectPr"):
            body.remove(child)


def _add_image(document, asset):
    if not asset.content_type.startswith("image/"):
        return
    try:
        with asset.file.open("rb") as stream:
            document.add_picture(stream, width=Inches(5.8))
        document.paragraphs[-1].alignment = 1
    except Exception:
        document.add_paragraph(f"[Image unavailable: {asset.original_filename}]")


@transaction.atomic
def generate_docx_export(*, project, user):
    source_import = project.source_imports.filter(status=DocxImport.Status.ACCEPTED).select_related("source_asset").first()
    if source_import:
        try:
            with source_import.source_asset.file.open("rb") as source:
                document = Document(source)
            _clear_body(document)
        except Exception:
            document = Document()
    else:
        document = Document()

    document.add_heading("LEXAMORA SERIES" if project.project_type == project.Type.SERIES else "LEXAMORA VIDEO", 0)
    document.add_heading("MASTER DOCUMENT", 0)
    document.add_heading(project.title, 1)
    metadata = document.add_table(rows=0, cols=2)
    for label, value in (
        ("Status", project.get_status_display()),
        ("Original language", project.original_language),
        ("Translation languages", ", ".join(project.translation_languages)),
        ("Rights holder", project.rights_holder),
        ("Publication", project.publication_info),
    ):
        cells = metadata.add_row().cells
        cells[0].text, cells[1].text = label, str(value or "")
    if project.concept:
        document.add_heading("Concept", 2)
        document.add_paragraph(project.concept)

    if project.characters.exists():
        document.add_heading("Characters", 1)
        for character in project.characters.all():
            document.add_heading(character.name, 2)
            document.add_paragraph(character.description)
            if character.visual_description:
                document.add_paragraph(character.visual_description)
            for asset in character.assets.all():
                _add_image(document, asset)

    for episode in project.episodes.all():
        document.add_heading(f"СЕРИЯ {episode.number}: {episode.title}", 1)
        document.add_heading("Краткое описание серии", 3)
        document.add_paragraph(episode.summary)
        for scene in episode.scenes.all():
            document.add_heading(f"Сцена {scene.number}. {scene.title}", 1)
            for heading, value in (
                ("Хук", scene.hook), ("Общее описание сцены", scene.description),
                ("Место", scene.location), ("Действия", scene.actions),
                ("Актёрская подача", scene.performance_notes),
            ):
                document.add_heading(heading, 3)
                document.add_paragraph(value or "")
            document.add_heading("Диалоги", 3)
            for line in scene.dialogue_lines.all():
                document.add_paragraph(f"{line.speaker}: “{line.text}”" if line.speaker else line.text)
            for asset in scene.assets.all():
                _add_image(document, asset)
            for prompt in scene.prompts.all():
                document.add_heading(prompt.title or f"{prompt.get_prompt_type_display()} prompt - {prompt.ai_model.name}", 2)
                table = document.add_table(rows=1, cols=2)
                table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Block", "Content"
                for block in prompt.blocks.all():
                    cells = table.add_row().cells
                    cells[0].text, cells[1].text = block.get_block_type_display(), block.content
                for asset in prompt.assets.all():
                    _add_image(document, asset)
        for track in episode.subtitle_tracks.all():
            document.add_heading(f"Subtitles {track.language} / {track.get_kind_display()}", 2)
            for line in track.lines.all():
                document.add_paragraph(line.text)

    linked_ids = set(project.assets.exclude(scene__isnull=False).exclude(character__isnull=False).values_list("id", flat=True))
    if linked_ids:
        document.add_heading("Project images", 1)
        for asset in project.assets.filter(id__in=linked_ids):
            _add_image(document, asset)

    output = io.BytesIO()
    document.save(output)
    output.seek(0)
    filename = f"{project.title[:120]}-master-document.docx"
    uploaded = SimpleUploadedFile(filename, output.read(), content_type=DOCX_TYPE)
    asset = create_asset(user=user, workspace=project.workspace, uploaded=uploaded, kind=Asset.Kind.EXPORT, project=project)
    job = ExportJob.objects.create(workspace=project.workspace, project=project, sections=["DOCX"], status=ExportJob.Status.SUCCESS, output_asset=asset, requested_by=user, started_at=timezone.now(), completed_at=timezone.now())
    audit(workspace=project.workspace, actor=user, action="DOCX_EXPORT_GENERATED", instance=asset, metadata={"exportId": str(job.id), "templateImportId": str(source_import.id) if source_import else None})
    return job
